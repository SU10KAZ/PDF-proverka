"""Отчёты «Сравнения версий»: Excel и HTML (план §7, этап 2).

Структура — предварительная (план §7, §15): согласуется с пользователем.
Отчёт строится только из сохранённых результатов прогона (``run.json``,
``result.json`` пар листов, ``spec.json``, ``permit.json``) — без повторного
сравнения. Каждый счётчик подписан, что он считает; непроверенные каналы и
«Нужна проверка» выводятся отдельно.
"""
from __future__ import annotations

import html
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import cv2
import fitz
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.writer.excel import save_workbook

from . import engine, paths, pdfio, profile, settings, store, tiles
from .engine import SHIFT_CLASSES
from .status import STATUS_LABELS

DARK, HEAD_FILL, NOTE_FILL = "263238", "F3F5F9", "FFF8E1"
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {"same": "E8F5E9", "same_zones_changed": "EEF6FF", "different": "FFF3E0", "needs_review": "FDECEA",
               "added": "E3F2FD", "removed": "FCE4EC", "unmatched": "FDECEA"}
CLASS_LABELS = {"text": "Текст", "text_moved": "Сдвиг надписи", "added": "Графика: добавлено (B)",
                "removed": "Графика: удалено (A)", "changed": "Графика: изменено", "color": "Изменение цвета",
                "light": "Светлая заливка / штриховка", "style": "Оформление линии", "block_moved": "Блок перемещён"}
REASONS = {"MATCH_LOW_CONFIDENCE": "сопоставление листов не уверенное", "DUPLICATE_SHEET": "дубль листа",
           "ALIGNMENT_UNRELIABLE": "совмещение ненадёжно", "ALIGNMENT_PARTIAL": "совмещена только часть листа",
           "PARTIAL_COVERAGE": "часть листа не проверена (предупреждение)", "TEXT_UNREADABLE_IN_DIFF_AREA":
           "надписи без текстового слоя в областях отличий", "NO_TEXT_LAYER": "нет текстового слоя",
           "RENDER_FAILED": "ошибка рендера", "SPEC_NOT_PARSED": "таблица СО не разобрана",
           "MATCH_UNRESOLVED": "лист не сопоставлен"}
CHANNELS = {"geometry": "графика", "color": "цвет и оформление", "text": "текст", "text_coverage": "читаемость надписей",
            "zones": "служебные зоны", "annotations": "PDF-пометки", "warnings": "предупреждения",
            "spec_positions": "позиции СО"}
MAX_CROPS = 300
CROP_DPI = 150


def _reasons(rs: list[str] | None) -> str:
    out = []
    for r in rs or []:
        if r.startswith("UNCHECKED_CHANNEL:"):
            out.append("не проверено: " + CHANNELS.get(r.split(":", 1)[1], r))
        else:
            out.append(REASONS.get(r, r))
    return "; ".join(dict.fromkeys(out))


def _fmt_box(b: list[float] | None) -> str:
    return ", ".join(str(round(v)) for v in b) if b else ""


# ------------------------------------------------------------------ данные

def collect(object_id: str, run: dict[str, Any]) -> dict[str, Any]:
    """Всё для отчёта по прогону: листы с результатами, СО, разрешение, паспорта версий."""
    from . import jobs     # поздний импорт: jobs импортирует много модулей
    sheets = []
    for it in run["items"]:
        res = None
        if it["kind"] == "compare" and it.get("status") == "done":
            try:
                res = jobs.sheet_result(object_id, run["run_id"], it["spid"])[0].get("result")
            except Exception:
                res = None
        sheets.append({"item": it, "result": res})
    catalog = store.load_catalog(object_id)

    def passport(side: str) -> dict[str, Any]:
        v = catalog["versions"].get(run[side]["version_id"], {})
        prof = profile.load(object_id, run[side]["sha256"])
        pages = prof["pages"]
        permits_path = paths.profile_dir(object_id, run[side]["sha256"]) / "permits_v2.json"
        permits = json.loads(permits_path.read_text(encoding="utf-8")) if permits_path.exists() else []
        return {"label": v.get("user_label") or v.get("file_name"), "file": v.get("file_name"), "sha256": run[side]["sha256"],
                "pages": len(pages), "spec_pages": sum(1 for p in pages if p.get("kind") == "spec"),
                "izm": max([int((p.get("stamp") or {}).get("izm_max") or 0) for p in pages] + [0]),
                "permits": ", ".join(f"{p['number']} (Изм. {', '.join(str(i) for i in sorted({r['izm'] for r in p['rows']}))})"
                                     for p in permits),
                "release": ", ".join(sorted(set(prof.get("summary", {}).get("release_dates") or [])))}
    return {"run": run, "sheets": sheets, "spec": jobs.run_spec(object_id, run["run_id"]),
            "permit": jobs.run_permit(object_id, run["run_id"]), "passport": {"a": passport("a"), "b": passport("b")},
            "built": datetime.now(timezone.utc).astimezone().strftime("%d.%m.%Y %H:%M")}


# ------------------------------------------------------------------- Excel

def _head(ws, row: int, columns: list[tuple[str, int]]) -> None:
    for col, (title, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        c = ws.cell(row, col, title)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=DARK)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    ws.freeze_panes = ws.cell(row + 1, 1)


def _row(ws, row: int, values: list[Any], fill: str | None = None) -> None:
    for col, v in enumerate(values, 1):
        c = ws.cell(row, col, v if v not in ("", None) else None)
        c.alignment = TOP
        c.border = BORDER
        c.font = Font(size=10)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)


def _sheet_label(it: dict[str, Any]) -> str:
    return it.get("title") or (f"стр. {it.get('a_page') or '—'} → {it.get('b_page') or '—'}")


def build_xlsx(data: dict[str, Any]) -> bytes:
    run, sheets = data["run"], data["sheets"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    pa, pb = data["passport"]["a"], data["passport"]["b"]
    counts: dict[str, int] = {}
    for s in sheets:
        key = s["item"].get("result_status") or s["item"]["status"]
        counts[key] = counts.get(key, 0) + 1
    rows = [("Сравнение версий", ""), ("Исходная версия A", f"{pa['label']} — {pa['file']}"), ("SHA-256 A", pa["sha256"]),
            ("Сравниваемая версия B", f"{pb['label']} — {pb['file']}"), ("SHA-256 B", pb["sha256"]),
            ("Охват", "все сопоставленные листы" if run.get("scope") == "all" else "выбранные пары листов"),
            ("Состояние прогона", run.get("state")), ("Дата отчёта", data["built"]), ("", ""),
            ("Листов (пар и непарных)", len(sheets))]
    rows += [(f"  {STATUS_LABELS.get(k, k)}", v) for k, v in sorted(counts.items(), key=lambda kv: -kv[1])]
    if data["spec"].get("available"):
        sp = data["spec"]
        rows += [("", ""), ("СО по позициям: изменено / добавлено / удалено",
                            " / ".join(str(sum(1 for c in sp["changes"] if c["kind"] == k)) for k in ("changed", "added", "removed"))),
                 ("СО: переехали между листами без изменений", sp.get("moved_between_pages"))]
    if data["permit"].get("available"):
        pm = data["permit"]
        rows += [("", ""), ("Разрешение: заявлено и найдено", len(pm["declared_and_found"])),
                 ("Разрешение: заявлено, отличий не найдено", len(pm["declared_no_diff"])),
                 ("Разрешение: отличия не заявлены", len(pm["diff_not_declared"]))]
    rows += [("", ""), ("Параметры", f"DPI {run['params']['dpi']}, допуск {run['params']['tol_pt']} pt, алгоритм {run['params'].get('algo')}")]
    for k, (a, b) in enumerate(rows, 1):
        ws.cell(k, 1, a).font = Font(bold=(k == 1), size=12 if k == 1 else 10)
        ws.cell(k, 2, b)
        ws.cell(k, 2).alignment = TOP

    # Листы
    ws = wb.create_sheet("Листы")
    cols = [("Лист (строка матрицы)", 46), ("Стр. A", 8), ("Лист A", 8), ("Стр. B", 8), ("Лист B", 8), ("Сопоставление", 14),
            ("Статус", 26), ("Причина «Нужна проверка»", 40), ("Отличий (основных)", 12), ("Оформление линий", 11),
            ("Изм. текста", 10), ("Служебных полей", 10), ("СО: изм. позиций", 10), ("Наложение (справочно)", 18),
            ("Скан (текст не сравнивается)", 12)]
    _head(ws, 1, cols)
    for i, s in enumerate(sheets, 2):
        it, res = s["item"], s["result"] or {}
        status = it.get("result_status") or it["status"]
        regions = it.get("regions") or 0
        style = it.get("style_regions") or 0
        texts = [c for c in ((res.get("text") or {}).get("changes") or []) if c["kind"] != "moved" and c.get("norm") != "spaces"]
        _row(ws, i, [_sheet_label(it), it.get("a_page"), (res.get("a") or {}).get("sheet_no"), it.get("b_page"),
                     (res.get("b") or {}).get("sheet_no"),
                     {"part": "часть листа"}.get(it.get("confidence"), it.get("confidence")),
                     STATUS_LABELS.get(status, status), _reasons(res.get("reasons") or it.get("reasons")),
                     max(0, regions - style) if it["kind"] == "compare" else None, style or None, len(texts) or None,
                     len(res.get("service") or []) or None, (it.get("spec") or {}).get("changes"),
                     STATUS_LABELS.get(it.get("overlay_status"), "") if it.get("overlay_status") else None,
                     ", ".join(x.upper() for x in ("a", "b") if (res.get(x) or {}).get("scan")) or None],
             STATUS_FILL.get(status))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(sheets) + 1)}"

    # Отличия
    ws = wb.create_sheet("Отличия")
    cols = [("Лист", 40), ("Стр. A", 7), ("Стр. B", 7), ("Класс", 22), ("Было (A)", 40), ("Стало (B)", 40),
            ("Рамка на A, pt", 22), ("Рамка на B, pt", 22), ("Примечание", 30)]
    _head(ws, 1, cols)
    r = 2
    for s in sheets:
        it, res = s["item"], s["result"]
        if not res:
            continue
        tch = {c["id"]: c for c in (res.get("text") or {}).get("changes") or []}
        for g in res.get("regions") or []:
            if g["class"] == "style" or g["class"] in SHIFT_CLASSES:
                continue      # сдвиг надписи / перенос блока — не изменение
            changes = [tch[i] for i in g.get("text_changes") or [] if i in tch and tch[i]["kind"] != "moved"]
            before = "; ".join(c["before"] for c in changes if c["before"]) or " ".join(g.get("text_before") or [])
            after = "; ".join(c["after"] for c in changes if c["after"]) or " ".join(g.get("text_after") or [])
            note = ""
            if g.get("text_unreadable"):
                note = "надпись без текстового слоя — сравнена только графически"
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), CLASS_LABELS.get(g["class"], g["class"]),
                         before, after, _fmt_box(g.get("bbox_a_pt")), _fmt_box(g.get("bbox_b_pt")), note])
            r += 1
        style = sum(1 for g in res.get("regions") or [] if g["class"] == "style")
        if style:
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), "Оформление линии", "", "", "", "",
                         f"{style} мест: толщина, тон или порядок наложения линий (одной строкой)"], NOTE_FILL)
            r += 1
    ws.auto_filter.ref = f"A1:I{max(2, r - 1)}"

    # Служебные зоны
    ws = wb.create_sheet("Служебные зоны")
    cols = [("Лист", 40), ("Стр. A", 7), ("Стр. B", 7), ("Поле", 34), ("Изменение", 14), ("Было (A)", 50), ("Стало (B)", 50)]
    _head(ws, 1, cols)
    r = 2
    change = {"changed": "изменено", "added": "появилось", "removed": "исчезло"}
    for s in sheets:
        it, res = s["item"], s["result"]
        if not res:
            continue
        for f in res.get("service") or []:
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), f["label"], change.get(f["change"], f["change"]),
                         f["before"], f["after"]])
            r += 1
        ann = res.get("annotations") or {}
        for a in ann.get("added") or []:
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), f"PDF-пометка {a.get('type')}", "появилось",
                         "", " ".join(str(a.get(k) or "") for k in ("subject", "author", "content")).strip()])
            r += 1
        for a in ann.get("removed") or []:
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), f"PDF-пометка {a.get('type')}", "исчезло",
                         " ".join(str(a.get(k) or "") for k in ("subject", "author", "content")).strip(), ""])
            r += 1
        for o in res.get("overlays") or []:
            if o["kind"] == "warning" and o["change"] != "same_text":
                _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), "Предупреждение / водяной знак",
                             change.get(o["change"], o["change"]), " | ".join(o.get("before") or []),
                             (" | ".join(o.get("after") or [])) + " (исчезновение не подтверждает согласования)"], NOTE_FILL)
                r += 1

    # Спецификация
    if data["spec"].get("available"):
        sp = data["spec"]
        ws = wb.create_sheet("Спецификация")
        cols = [("Вид", 14), ("Позиция", 60), ("Раздел", 34), ("Графа", 14), ("Было (A)", 24), ("Стало (B)", 24),
                ("Стр. A", 10), ("Стр. B", 10)]
        _head(ws, 1, cols)
        kinds = {"changed": "изменено", "added": "добавлено", "removed": "удалено", "ambiguous": "неоднозначно",
                 "renumbered": "перенумерация"}
        r = 2
        for c in sp["changes"]:
            fields = list(c["fields"].items()) or [("", ["", ""])]
            for f, (va, vb) in fields:
                _row(ws, r, [kinds.get(c["kind"], c["kind"]), c["label"], c["section"], f, va, vb,
                             ", ".join(map(str, c["a_pages"])), ", ".join(map(str, c["b_pages"]))],
                     NOTE_FILL if c["kind"] in ("ambiguous", "renumbered") else None)
                r += 1
        r += 1
        ws.cell(r, 1, "Итоги по марке (сумма по всем листам СО)").font = Font(bold=True)
        r += 1
        for t in sp.get("totals") or []:
            _row(ws, r, ["итог", t["label"], "", "кол.", t["a"], t["b"], "", f"Δ {t['delta']:+g}"])
            r += 1

    # Разрешение
    if data["permit"].get("available"):
        pm = data["permit"]
        ws = wb.create_sheet("Разрешение")
        cols = [("Сверка", 34), ("Лист", 12), ("Изм.", 8), ("Заявлено проектировщиком / найдено", 90)]
        _head(ws, 1, cols)
        r = 2
        for title, key in (("Заявлено и найдено", "declared_and_found"), ("Заявлено, отличий не найдено", "declared_no_diff"),
                           ("Нужна проверка", "needs_review")):
            for e in pm[key]:
                _row(ws, r, [title, e["ref"], ", ".join(map(str, e["izm"])), "; ".join(e["declared"]) +
                             (f" — {e['reason']}" if e.get("reason") else "")],
                     NOTE_FILL if key != "declared_and_found" else None)
                r += 1
        for e in pm["diff_not_declared"]:
            _row(ws, r, ["Отличия найдены, не заявлены", e["ref"], "", f"{e.get('title')} — {STATUS_LABELS.get(e['status'], e['status'])}"],
                 NOTE_FILL)
            r += 1
        r += 1
        ws.cell(r, 1, "Строки разрешений версии B (буквально)").font = Font(bold=True)
        r += 1
        for row in pm["rows"]:
            _row(ws, r, [row.get("permit"), row["sheet"], row["izm"], row["text"]])
            r += 1

    # Нужна проверка
    ws = wb.create_sheet("Нужна проверка")
    cols = [("Лист", 50), ("Стр. A", 8), ("Стр. B", 8), ("Причина", 80)]
    _head(ws, 1, cols)
    r = 2
    for s in sheets:
        it, res = s["item"], s["result"] or {}
        status = it.get("result_status") or it["status"]
        if status in ("needs_review", "unmatched") or it["status"] == "failed":
            reason = _reasons(res.get("reasons") or it.get("reasons")) or it.get("error") or STATUS_LABELS.get(status, status)
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), reason], STATUS_FILL.get("needs_review"))
            r += 1

    # Источники
    ws = wb.create_sheet("Источники")
    cols = [("Версия", 10), ("Метка", 34), ("Файл", 40), ("SHA-256", 66), ("Страниц", 9), ("Страниц СО", 10),
            ("Последнее «Изм.» по штампам", 14), ("Разрешения", 40), ("Дата выпуска (штамп «В производство»)", 24)]
    _head(ws, 1, cols)
    for i, side in enumerate(("a", "b"), 2):
        p = data["passport"][side]
        _row(ws, i, [side.upper(), p["label"], p["file"], p["sha256"], p["pages"], p["spec_pages"], p["izm"], p["permits"],
                     p["release"]])

    # Методика
    ws = wb.create_sheet("Методика")
    ws.column_dimensions["A"].width = 120
    lines = [
        "Сравнение детерминированное, без ИИ и без сети. Исходные PDF не изменяются (копии только для чтения).",
        f"Листы сопоставлены по штампам, тексту и растру; пара листов совмещается подобием (сдвиг, масштаб, поворот) "
        f"по опорным словам, без деформаций.",
        f"Наложение при {run['params']['dpi']} dpi, допуск {run['params']['tol_pt']} pt: серое — общее, красное — только в A, "
        f"синее — только в B. Отличия меньше допуска не скрываются (в результате листа — «в пределах допуска»).",
        f"Пороги: тёмные линии — серый < {settings.DARK_THRESHOLD}, светлые элементы — < {settings.LIGHT_THRESHOLD}; цвет — "
        f"разница цветности Lab > {settings.COLOR_DELTA}. Изменение толщины/тона линии — отдельный класс «оформление».",
        "Из основного наложения исключены: основная надпись и поля за рамкой (сравниваются отдельно по полям), отметка "
        "выпуска и QR, предупреждения (отделяются только с контролем, иначе площадь отмечается непроверенной), PDF-пометки.",
        "«Совпадает» выдаётся только при проверенных всех применимых каналах; иначе — «Нужна проверка» с причиной.",
        "Сдвиг той же надписи и перенос блока целиком изменением не считаются и в отличия не выводятся; изменения "
        "внутри перенесённого блока — обычные отличия.",
        "Спецификации сравниваются по позициям через страницы: перенос позиции на другой лист и перенумерация — не изменения.",
        "Разрешение сверяется по номерам листов, без смыслового анализа. Исчезновение предупреждения не подтверждает согласования.",
    ]
    if data["spec"].get("available"):
        lines.append(f"Версия разбора СО: {data['spec'].get('spec_version')}.")
    for k, line in enumerate(lines, 1):
        c = ws.cell(k, 1, line)
        c.alignment = TOP
    buf = io.BytesIO()
    save_workbook(wb, buf)
    return buf.getvalue()


def _pair_sheet_name(k: int, a_label: str, b_label: str, used: set[str]) -> str:
    """Имя листа перехода: «0-1» по номерам из подписей «(0) …», иначе по порядку «1-2»."""
    na, nb = (re.match(r"^\s*\((\d+)\)", x or "") for x in (a_label, b_label))
    name = f"{na.group(1)}-{nb.group(1)}" if na and nb else f"{k + 1}-{k + 2}"
    base, n = name, 2
    while name in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name)
    return name


def build_series_xlsx(series: dict[str, Any], runs: list[dict[str, Any]]) -> bytes:
    """Сводка переходов и по листу на каждый переход («0-1», «1-2», …)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка переходов"
    cols = [("Переход A → B", 46), ("Лист Excel", 10), ("Листов сравнено", 10), ("Совпадает", 10),
            ("Только служебные зоны", 12), ("Есть отличия", 10), ("Из них только оформление", 12), ("Нужна проверка", 10),
            ("Добавлено листов", 10), ("Удалено листов", 10), ("Вывод", 70)]
    _head(ws, 1, cols)
    runs_by_id = {run["run_id"]: run for run in runs}
    used: set[str] = set()
    names = [_pair_sheet_name(k, p.get("a_label") or "", p.get("b_label") or "", used) for k, p in enumerate(series["pairs"])]
    for i, (p, name) in enumerate(zip(series["pairs"], names), 2):
        sm = p.get("summary") or {}
        _row(ws, i, [f"{p['a_label']} → {p['b_label']}", name, sm.get("compared"), sm.get("same"), sm.get("zones_only"),
                     sm.get("different"), sm.get("style_only"), sm.get("needs_review"), sm.get("added"), sm.get("removed"),
                     "; ".join(sm.get("verdict") or [])])
        if p.get("run_id") in runs_by_id:
            link = ws.cell(i, 2)
            link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{name}'!A1")
            link.font = Font(color="1565C0", underline="single")
    cols = [("Лист", 46), ("Стр. A", 8), ("Стр. B", 8), ("Статус", 26), ("Отличий", 10), ("Причина", 50)]
    for p, name in zip(series["pairs"], names):
        run = runs_by_id.get(p.get("run_id"))
        if not run:
            continue
        ws = wb.create_sheet(name)
        ws.cell(1, 1, f"{p['a_label']} → {p['b_label']}").font = Font(bold=True)
        _head(ws, 2, cols)
        r = 3
        for it in run["items"]:
            status = it.get("result_status") or it["status"]
            _row(ws, r, [_sheet_label(it), it.get("a_page"), it.get("b_page"), STATUS_LABELS.get(status, status),
                         max(0, (it.get("regions") or 0) - (it.get("style_regions") or 0)) or None, _reasons(it.get("reasons"))],
                 STATUS_FILL.get(status))
            r += 1
        ws.auto_filter.ref = f"A2:F{max(3, r - 1)}"
        ws.freeze_panes = "A3"
    buf = io.BytesIO()
    save_workbook(wb, buf)
    return buf.getvalue()


# -------------------------------------------------------------------- HTML

CSS = """body{font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif;margin:0;color:#1c2733;background:#fff}
main{max-width:1200px;margin:0 auto;padding:16px}h1{font-size:22px}h2{font-size:18px;margin-top:28px}
table{border-collapse:collapse;width:100%;font-size:13px}th,td{border:1px solid #d0d7de;padding:4px 6px;text-align:left;vertical-align:top}
th{background:#263238;color:#fff}.st-same{background:#e8f5e9}.st-same_zones_changed{background:#eef6ff}
.st-different{background:#fff3e0}.st-needs_review,.st-unmatched{background:#fdecea}.st-added{background:#e3f2fd}.st-removed{background:#fce4ec}
.was{color:#d4303a}.now{color:#146cc9}.muted{color:#5b6b80;font-size:12px}img{max-width:100%;border:1px solid #d0d7de}
.crops{display:grid;grid-template-columns:repeat(3,1fr);gap:6px;margin:6px 0 14px}.crops div{font-size:11px;color:#5b6b80}
.legend span{display:inline-block;padding:1px 6px;margin-right:8px;border-radius:3px;font-size:12px}
@media(max-width:700px){.crops{grid-template-columns:1fr}}"""


def _png(arr: np.ndarray, max_w: int = 1400) -> bytes:
    h, w = arr.shape[:2]
    if w > max_w:
        arr = cv2.resize(arr, (max_w, int(h * max_w / w)), interpolation=cv2.INTER_AREA)
    ok, buf = cv2.imencode(".png", cv2.cvtColor(arr, cv2.COLOR_RGB2BGR) if arr.ndim == 3 else arr)
    return bytes(buf) if ok else b""


def _crop_page(pdf: Path, page: int, bbox_pt: list[float], pad: float = 12.0) -> np.ndarray | None:
    doc = pdfio.open_memory(pdfio.read_bytes(pdf))
    pg = doc[page - 1]
    clip = fitz.Rect(bbox_pt[0] - pad, bbox_pt[1] - pad, bbox_pt[2] + pad, bbox_pt[3] + pad) & pg.rect
    if clip.is_empty:
        return None
    arr, _x, _y = pdfio.render(pg, fitz.Matrix(CROP_DPI / 72, CROP_DPI / 72), clip=clip)
    return arr


def _crop_overlay(result_dir: Path, bbox_px: list[int], pad_px: int = 30) -> np.ndarray | None:
    data = tiles._load_masks(result_dir)
    h, w = data["shape"]
    x0, y0 = max(0, bbox_px[0] - pad_px), max(0, bbox_px[1] - pad_px)
    x1, y1 = min(w, bbox_px[2] + pad_px), min(h, bbox_px[3] + pad_px)
    if x1 <= x0 or y1 <= y0:
        return None
    layers = {name: tiles._unpack(packed, x0, y0, x1, y1) for name, packed in data["layers"].items()}
    return engine.compose(layers, tiles.DEFAULT_LAYERS)


def build_html_zip(object_id: str, data: dict[str, Any]) -> bytes:
    run = data["run"]
    pa, pb = data["passport"]["a"], data["passport"]["b"]
    pdf_a, pdf_b = paths.source_pdf(object_id, pa["sha256"]), paths.source_pdf(object_id, pb["sha256"])
    buf = io.BytesIO()
    zf = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
    esc = html.escape
    rows = []
    crops_left = MAX_CROPS
    for s in data["sheets"]:
        it, res = s["item"], s["result"]
        status = it.get("result_status") or it["status"]
        link = ""
        if res and status not in ("same",):
            name = f"sheet_{it['spid']}.html"
            link = f'<a href="{name}">открыть</a>'
            body = [f"<h1>{esc(_sheet_label(it))}</h1>",
                    f"<p>A: {esc(pa['label'] or '')}, стр. {it.get('a_page')} → B: {esc(pb['label'] or '')}, стр. {it.get('b_page')} · "
                    f"<b>{esc(STATUS_LABELS.get(status, status))}</b> {esc(_reasons(res.get('reasons')))}</p>",
                    '<p class="legend"><span style="background:#969ca3;color:#fff">общее</span>'
                    '<span style="background:#d4303a;color:#fff">только A</span><span style="background:#146cc9;color:#fff">только B</span>'
                    '<span style="background:#009688;color:#fff">оформление линий</span><span style="background:#288c3c;color:#fff">перемещено</span>'
                    '<span style="background:#904aaa;color:#fff">цвет</span></p>']
            rdir = Path(it["result_dir"])
            prev = rdir / "preview.png"
            if prev.exists():
                img = cv2.cvtColor(cv2.imread(str(prev)), cv2.COLOR_BGR2RGB)
                zf.writestr(f"img/{it['spid']}_preview.png", _png(img))
                body.append(f'<img src="img/{it["spid"]}_preview.png" alt="наложение">')
            tch = {c["id"]: c for c in (res.get("text") or {}).get("changes") or []}
            text_rows = [c for c in tch.values() if c.get("norm") != "spaces" and c["kind"] != "moved"]
            if text_rows:
                body.append("<h2>Текст «было → стало»</h2><table><tr><th>Вид</th><th>Было (A)</th><th>Стало (B)</th><th>A, pt</th><th>B, pt</th></tr>")
                kinds = {"changed": "изменено", "added": "добавлено", "removed": "удалено"}
                for c in text_rows[:400]:
                    body.append(f"<tr><td>{kinds.get(c['kind'], c['kind'])}</td><td class=was>{esc(c['before'])}</td>"
                                f"<td class=now>{esc(c['after'])}</td><td>{_fmt_box(c.get('a_bbox'))}</td><td>{_fmt_box(c.get('b_bbox'))}</td></tr>")
                body.append("</table>")
            if res.get("service"):
                body.append("<h2>Служебные зоны</h2><table><tr><th>Поле</th><th>Было (A)</th><th>Стало (B)</th></tr>")
                for f in res["service"]:
                    body.append(f"<tr><td>{esc(f['label'])}</td><td class=was>{esc(f['before'])}</td><td class=now>{esc(f['after'])}</td></tr>")
                body.append("</table>")
            if (res.get("spec_page") or {}).get("items"):
                body.append("<h2>Спецификация по позициям (на этом листе)</h2><table><tr><th>Вид</th><th>Позиция</th><th>Было → стало</th></tr>")
                for c in res["spec_page"]["items"]:
                    if c["kind"] == "renumbered":
                        continue
                    ch = "; ".join(f"{k}: {v[0] or '—'} → {v[1] or '—'}" for k, v in c["fields"].items())
                    body.append(f"<tr><td>{esc(c['kind'])}</td><td>{esc(c['label'])}</td><td>{esc(ch)}</td></tr>")
                body.append("</table>")
            regs = [g for g in res.get("regions") or [] if g["class"] != "style" and g["class"] not in SHIFT_CLASSES]
            if regs:
                body.append(f"<h2>Отличия: {len(regs)}</h2>")
                for k, g in enumerate(regs, 1):
                    changes = [tch[i] for i in g.get("text_changes") or [] if i in tch and tch[i]["kind"] != "moved"]
                    txt = "; ".join(f"{c['before'] or '—'} → {c['after'] or '—'}" for c in changes[:4])
                    body.append(f"<p><b>{k}. {esc(CLASS_LABELS.get(g['class'], g['class']))}</b> {esc(txt)} "
                                f"<span class=muted>A: {_fmt_box(g.get('bbox_a_pt'))} pt · B: {_fmt_box(g.get('bbox_b_pt'))} pt</span></p>")
                    if crops_left > 0:
                        crops_left -= 1
                        ca = _crop_page(pdf_a, it["a_page"], g["bbox_a_pt"]) if g.get("bbox_a_pt") else None
                        cb = _crop_page(pdf_b, it["b_page"], g["bbox_b_pt"]) if g.get("bbox_b_pt") else None
                        co = _crop_overlay(rdir, g["bbox_px"])
                        cells = []
                        for tag, arr in (("A", ca), ("B", cb), ("наложение", co)):
                            if arr is not None:
                                fn = f"img/{it['spid']}_{g['id']}_{'ov' if tag == 'наложение' else tag}.png"
                                zf.writestr(fn, _png(arr, 700))
                                cells.append(f'<div>{tag}<br><img src="{fn}" alt=""></div>')
                        body.append('<div class="crops">' + "".join(cells) + "</div>")
            zf.writestr(name, _page(_sheet_label(it), "\n".join(body) + '<p><a href="index.html">← к сводке</a></p>'))
        rows.append(f'<tr class="st-{status}"><td>{esc(_sheet_label(it))}</td><td>{it.get("a_page") or ""}</td>'
                    f'<td>{it.get("b_page") or ""}</td><td>{esc(STATUS_LABELS.get(status, status))}</td>'
                    f'<td>{esc(_reasons((res or {}).get("reasons") or it.get("reasons")))}</td><td>{link}</td></tr>')
    head = (f"<h1>Сравнение версий</h1><p>A: <b>{esc(pa['label'] or '')}</b> ({esc(pa['file'] or '')})<br>"
            f"B: <b>{esc(pb['label'] or '')}</b> ({esc(pb['file'] or '')})<br>"
            f"<span class=muted>DPI {run['params']['dpi']}, допуск {run['params']['tol_pt']} pt, алгоритм {run['params'].get('algo')}; "
            f"отчёт {data['built']}. Сравнение детерминированное, без ИИ.</span></p>")
    extra = ""
    if data["permit"].get("available"):
        pm = data["permit"]
        extra += ("<h2>Разрешение: заявлено ↔ найдено</h2><p>Заявлено и найдено: "
                  + esc(", ".join(e["ref"] for e in pm["declared_and_found"]) or "—")
                  + "<br>Заявлено, отличий не найдено: " + esc(", ".join(e["ref"] for e in pm["declared_no_diff"]) or "—")
                  + "<br>Отличия найдены, но не заявлены: " + esc(", ".join(e["ref"] for e in pm["diff_not_declared"]) or "—") + "</p>")
    index = (head + extra + "<h2>Листы</h2><table><tr><th>Лист</th><th>Стр. A</th><th>Стр. B</th><th>Статус</th>"
             "<th>Причина</th><th></th></tr>" + "".join(rows) + "</table>")
    zf.writestr("index.html", _page("Сравнение версий", index))
    zf.close()
    return buf.getvalue()


def _page(title: str, body: str) -> str:
    return (f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, '
            f'initial-scale=1"><title>{html.escape(title)}</title><style>{CSS}</style></head><body><main>{body}</main></body></html>')
