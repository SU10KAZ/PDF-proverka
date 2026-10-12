"""Инженерный разбор итоговых карточек по четырём корзинам: загрузка xlsx и хранение.

Разбор делается вне портала (по промпту «разбор ИТ для доп. работ ПД → РД») в
шаблоне, который выдаёт кнопка «Выгрузить отчёт», и загружается обратно. Здесь
только чтение листа «Карточки ИТ» (и «Кратко» со сводки, если его правили
вручную), проверка значений и сохранение рядом с парой — без вызова моделей.
Разбор привязан к прогону-источнику и прогону Консолидатора (+ sha256 итогового
результата): к другой итоговой сводке он не применяется.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .report_xlsx import BASKETS, CARD_COLUMNS, CONFIDENCE, ID_RE, SHEETS, SUMMARY_COLUMNS

SCHEMA = "projectchange-it-analysis/1"
DIR_NAME = "project_change_it_reports"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
_RUN_ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,80}$")
WRITE_LOCK = threading.RLock()


class AnalysisError(ValueError):
    """Файл разбора не принят; ``errors`` — построчные причины для инженера."""

    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


def _dir(session_id: str, pair_id: str, source_run_id: str, consolidator_run_id: str) -> Path:
    from backend.app.services.stage_comparison import paths

    for value in (source_run_id, consolidator_run_id):
        if not _RUN_ID_RE.fullmatch(value or ""):
            raise ValueError("Некорректный идентификатор прогона")
    return paths.pair_dir(session_id, pair_id) / DIR_NAME / source_run_id / consolidator_run_id


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def _header_map(row: tuple[Any, ...], columns) -> dict[str, int]:
    names = {_text(v).lower(): i for i, v in enumerate(row) if _text(v)}
    return {key: names[header.lower()] for key, header, _ in columns if header.lower() in names}


def parse(data: bytes, *, n_items: int) -> tuple[list[dict[str, str]], list[str]]:
    """Строки «Карточек ИТ» + предупреждения. Нарушение легенды — AnalysisError."""
    from openpyxl import load_workbook

    if len(data) > MAX_UPLOAD_BYTES:
        raise AnalysisError([f"Файл больше {MAX_UPLOAD_BYTES // (1024 * 1024)} МБ"])
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # повреждённый или не-xlsx файл
        raise AnalysisError([f"Файл не читается как xlsx: {exc.__class__.__name__}"]) from exc
    if SHEETS[1] not in wb.sheetnames:
        raise AnalysisError([f"Нет листа «{SHEETS[1]}»"])
    sheet = wb[SHEETS[1]].iter_rows(values_only=True)
    header = next(sheet, ())
    cols = _header_map(header, CARD_COLUMNS)
    missing = [h for key, h, _ in CARD_COLUMNS if key not in cols]
    if missing:
        raise AnalysisError([f"На листе «{SHEETS[1]}» нет колонок: {', '.join(missing)}"])
    rows, errors, warnings, seen = [], [], [], set()
    for line, values in enumerate(sheet, 2):
        row = {key: _text(values[i]) if i < len(values) else "" for key, i in cols.items()}
        if not any(row.values()):
            continue
        rid, where = row["id"], f"строка {line}"
        m = ID_RE.match(rid)
        if not m:
            errors.append(f"{where}: ID «{rid}» не в формате ИТ-001 / ИТ-001.2")
            continue
        if not 1 <= int(m.group(1)) <= n_items:
            errors.append(f"{where}: {rid} — в итоговой сводке {n_items} ИТ")
        if rid in seen:
            errors.append(f"{where}: {rid} повторяется")
        seen.add(rid)
        if row["basket"] and row["basket"] not in BASKETS:
            errors.append(f"{where}: {rid} — корзина «{row['basket']}» не из легенды ({'; '.join(BASKETS)})")
        if row["confidence"] and row["confidence"] not in CONFIDENCE:
            errors.append(f"{where}: {rid} — уверенность «{row['confidence']}» не из ({', '.join(CONFIDENCE)})")
        if not row["basket"] or not row["confidence"]:
            warnings.append(f"{rid}: не заполнены корзина и/или уверенность")
        if m.group(2) and f"ИТ-{m.group(1)}" not in {r["id"] for r in rows}:
            warnings.append(f"{rid}: хвост стоит раньше родителя ИТ-{m.group(1)} или родителя нет")
        rows.append(row)
    if errors:
        raise AnalysisError(errors)
    if not rows:
        raise AnalysisError([f"Лист «{SHEETS[1]}» пуст"])
    absent = [f"ИТ-{i:03d}" for i in range(1, n_items + 1) if f"ИТ-{i:03d}" not in seen]
    if absent:
        warnings.append(f"Нет карточек: {', '.join(absent)}")
    briefs = _briefs(wb)
    for row in rows:
        if briefs.get(row["id"]):
            row["brief"] = briefs[row["id"]]
    return rows, warnings


def _briefs(wb) -> dict[str, str]:
    """«Кратко» со сводки — если его правили вручную, при повторной выгрузке он сохраняется."""
    if SHEETS[2] not in wb.sheetnames:
        return {}
    out, cols = {}, None
    for values in wb[SHEETS[2]].iter_rows(values_only=True):
        if cols is None:
            found = _header_map(values, SUMMARY_COLUMNS)
            if "id" in found and "brief" in found:
                cols = found
            continue
        rid = _text(values[cols["id"]]) if cols["id"] < len(values) else ""
        if ID_RE.match(rid) and cols["brief"] < len(values):
            out[rid] = _text(values[cols["brief"]])
    return out


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, path)


SOURCE_MANUAL, SOURCE_MODEL = "manual", "model"


def _stamp(now: datetime) -> str:
    # микросекунды: две записи в одну секунду не перезаписывают ни файл, ни прежний разбор в history
    return now.strftime("%Y%m%dT%H%M%S%fZ")


def store(session_id: str, pair_id: str, view: dict[str, Any], rows: list[dict[str, str]], *,
          warnings: list[str], source: str, author: str, extra: dict[str, Any] | None = None) -> dict[str, Any]:
    """Записать разбор этой итоговой сводки; прежний уходит в history/."""
    with WRITE_LOCK:
        return _store(session_id, pair_id, view, rows, warnings=warnings, source=source, author=author, extra=extra)


def _store(session_id, pair_id, view, rows, *, warnings, source, author, extra):
    base = _dir(session_id, pair_id, view["source_run_id"], view["consolidator_run_id"])
    base.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc)
    current = base / "it_analysis.json"
    if current.exists():  # прежний разбор не теряется
        (base / "history").mkdir(exist_ok=True)
        shutil.copyfile(current, base / "history" / f"it_analysis.{_stamp(now)}.json")
    payload = {"schema": SCHEMA, "session_id": session_id, "pair_id": pair_id,
               "source_run_id": view["source_run_id"], "consolidator_run_id": view["consolidator_run_id"],
               "shadow_result_sha256": view["shadow_result_sha256"], "uploaded_at": now.isoformat(),
               "uploaded_by": author, "source": source, **(extra or {}), "warnings": warnings, "rows": rows}
    _atomic_json(current, payload)
    return status(session_id, pair_id, view)


def save(session_id: str, pair_id: str, view: dict[str, Any], data: bytes, *, filename: str,
         author: str) -> dict[str, Any]:
    """Ручной разбор инженера (кнопка «Загрузить разбор»)."""
    rows, warnings = parse(data, n_items=len(view.get("consolidated") or []))
    base = _dir(session_id, pair_id, view["source_run_id"], view["consolidator_run_id"])
    (base / "uploads").mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(data).hexdigest()
    (base / "uploads" / f"{_stamp(datetime.now(timezone.utc))}_{digest[:8]}.xlsx").write_bytes(data)
    return store(session_id, pair_id, view, rows, warnings=warnings, source=SOURCE_MANUAL, author=author,
                 extra={"filename": filename, "file_sha256": digest})


def load(session_id: str, pair_id: str, view: dict[str, Any]) -> dict[str, Any] | None:
    """Разбор этой итоговой сводки или None (нет, другой формат, другая сводка)."""
    path = _dir(session_id, pair_id, view["source_run_id"], view["consolidator_run_id"]) / "it_analysis.json"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if payload.get("schema") != SCHEMA or payload.get("shadow_result_sha256") != view.get("shadow_result_sha256"):
        return None
    return payload


def prune_except(session_id: str, pair_id: str, source_run_id: str, consolidator_run_id: str, *,
                 busy: set[tuple[str, str]] = frozenset()) -> list[str]:
    """Удалить разборы прочих итоговых сводок пары (их прогоны Консолидатора удалены);
    ``busy`` — сводки, по которым сейчас идёт разбор моделью, их не трогаем."""
    keep = _dir(session_id, pair_id, source_run_id, consolidator_run_id)
    root = keep.parent.parent
    removed = []
    with WRITE_LOCK:
        if not root.is_dir():
            return removed
        for src in sorted(p for p in root.iterdir() if p.is_dir()):
            for run in sorted(p for p in src.iterdir() if p.is_dir()):
                if run != keep and (src.name, run.name) not in busy:
                    shutil.rmtree(run)
                    removed.append(f"{src.name}/{run.name}")
            if src != keep.parent and not any(src.iterdir()):
                src.rmdir()
    return removed


def status(session_id: str, pair_id: str, view: dict[str, Any]) -> dict[str, Any]:
    payload = load(session_id, pair_id, view)
    if not payload:
        return {"schema": "projectchange-it-analysis-status/1", "available": False}
    rows = payload["rows"]
    return {"schema": "projectchange-it-analysis-status/1", "available": True,
            "source": payload.get("source") or SOURCE_MANUAL,  # разборы до появления поля — ручные
            "uploaded_at": payload["uploaded_at"], "uploaded_by": payload["uploaded_by"],
            "filename": payload.get("filename") or "", "model": payload.get("model") or "",
            "rows": len(rows), "filled": sum(1 for r in rows if r.get("basket") and r.get("confidence")),
            "failed_ids": payload.get("failed_ids") or [], "warnings": payload.get("warnings") or []}
