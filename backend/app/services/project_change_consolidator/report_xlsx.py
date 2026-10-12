"""Excel-отчёт по итоговым карточкам Консолидатора: «Итоговые · Карточки ИТ · Краткая сводка».

Лист «Итоговые» строится детерминированно из ``view.consolidated_view`` (то же, что
показывает портал: строка ИТ + свёрнутые подробности). Листы «Карточки ИТ» и
«Краткая сводка» — инженерный разбор по четырём корзинам: если разбор загружен
(``it_analysis``), он переносится как есть, иначе выдаётся шаблон для заполнения
(ID, название, было/стало и страницы источников заполнены, корзины пустые).
Модели здесь не вызываются.
"""
from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from itertools import zip_longest
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SHEETS = ("Итоговые", "Карточки ИТ", "Краткая сводка")
BASKET_FULL, BASKET_CHANGE, BASKET_OPT, BASKET_DETAIL = (
    "Полностью дополнительные работы", "Изменение проектного решения", "Оптимизация", "Детализация / не доп.")
BASKETS = (BASKET_FULL, BASKET_CHANGE, BASKET_OPT, BASKET_DETAIL)
CONFIDENCE = ("уверен", "сначала лист", "после АИ")
# (ключ, заголовок, ширина) — порядок и названия колонок как в эталонном разборе ИОС2.1.
CARD_COLUMNS = (
    ("id", "ID", 14), ("basket", "Корзина", 28), ("confidence", "Уверенность", 16), ("title", "Название", 40),
    ("resume", "Резюме для инженера", 46), ("old", "Было (ПД)", 40), ("new", "Стало (РД)", 40),
    ("customer", "Заказчику / почему не подаём", 48), ("engineer", "Инженеру: как проверить и считать", 48),
    ("complex", "Комплекс / смежники", 36), ("calc", "Что в расчёт", 32), ("related", "Связано с", 22),
    ("where", "Куда смотреть (шифры)", 36), ("note", "Примечание", 52),
)
SUMMARY_COLUMNS = (("id", "ID", 14), ("basket", "Корзина", 32), ("confidence", "Увер.", 14),
                   ("title", "Название", 48), ("brief", "Кратко", 56), ("note", "Примечание", 52))
ID_RE = re.compile(r"^ИТ-(\d{3})(?:\.(\d+))?$")

CHANNELS = {"ENGINEERING_CHANGE": "Инженерное изменение", "DOCUMENTARY_CHANGE": "Документальная правка",
            "REVIEW": "Нужна проверка инженера", "NOT_ASSESSED": "Не оценено Консолидатором"}
PARAM_STATUS = {"CONSISTENT": "одинаково", "LOCATION_VARIANT": "различается по местам", "CONFLICT": "противоречие"}
MSK = timezone(timedelta(hours=3))

INK, MUTED, HEAD_INK, TEAL, LINK = "142D4E", "486386", "435A76", "087F82", "008E92"
F_TITLE, F_HEAD, F_SECTION, F_DETAIL = "E5F3F2", "F3F5F9", "E5F3F2", "F6F7FA"
DARK, NOTE_FILL = "263238", "FFF8E1"
# Цвета таблицы характеристик — как на портале: «Было» жёлтая, «Стало» зелёная,
# «Характеристика» бирюзовая; тип источника — цветом метки портала.
F_OLD_HEAD, F_OLD, F_NEW_HEAD, F_NEW = "FCE7B5", "FFF7E3", "C9EEDF", "EEF9F4"
F_NAME, F_NAME_HEAD, F_STATUS, STATUS_INK = "DDF1EF", "BFE5E1", "FFEEC4", "AD6900"
SOURCE_INK = {"TEXT": "5B6B80", "TABLE": "2563EB", "GRAPHIC": "087F82"}
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(wrap_text=True, vertical="top")


def it_id(index: int) -> str:
    return f"ИТ-{index + 1:03d}"


def plural(n: int, one: str, few: str, many: str) -> str:
    if n % 10 == 1 and n % 100 != 11:
        return one
    return few if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14 else many


def first_sentence(text: str) -> str:
    """«Кратко» на сводке — первое предложение резюме (правило эталона, без разбора сокращений)."""
    return re.split(r"(?<=[.!?])\s+", (text or "").strip(), maxsplit=1)[0]


def short_code(document_code: str) -> str:
    """«СТ26-01-14-ИОС2.1» → «ИОС2.1» (марка в заголовке сводки и имени файла)."""
    code = (document_code or "").strip()
    return code.rsplit("-", 1)[-1] if "-" in code else code


def _fill(rgb: str | None) -> PatternFill:
    return PatternFill("solid", fgColor=rgb) if rgb else PatternFill()


def _value(v: str, unit: str) -> str:
    """Значение с единицей; единицу, уже написанную в значении, не повторяем («10 мм (по расчёту)»)."""
    if not v:
        return "—"
    parts = [u.strip() for u in unit.split(";") if u.strip()]
    return v if not parts or all(u in v for u in parts) else f"{v} {unit}"


def _sources(item: dict[str, Any]) -> list[str]:
    return list(dict.fromkeys(e.get("source_type") for e in item.get("evidence") or [] if e.get("source_type")))


def _pages(item: dict[str, Any]) -> str:
    pages = defaultdict(set)
    for e in item.get("evidence") or []:
        if e.get("side") in ("OLD", "NEW") and e.get("page"):
            pages[e["side"]].add(int(e["page"]))
    parts = [f"{label} стр. {', '.join(map(str, sorted(pages[side])))}"
             for side, label in (("OLD", "ПД"), ("NEW", "РД")) if pages[side]]
    return "; ".join(parts)


SOURCE_TYPES = ("TEXT", "TABLE", "GRAPHIC")


def parameter_sources(card: dict[str, Any], parameter: dict[str, Any],
                      originals: list[dict[str, Any]]) -> dict[str, list[dict[str, str]]]:
    """Шифр · источник · страницы по PDF для характеристики — как ``parameterSources`` портала.

    Сводная карточка берёт доказательства исходных карточек, на которые ссылаются
    значения характеристики; исходная — свои. Сборки несут шифр и страницу
    исходного PDF (``origin``), а не сборного документа.
    """
    members = {str(c.get("id")): c for c in originals or []}
    if card.get("kind") == "CONSOLIDATED":
        ids = list(dict.fromkeys(str(v.get("card_id")) for v in parameter.get("values") or []))
        source_cards = [members[i] for i in ids if i in members]
    else:
        source_cards = [card]
    result: dict[str, list[dict[str, str]]] = {}
    for side in ("OLD", "NEW"):
        groups: dict[tuple, dict[str, Any]] = {}
        for source in source_cards:
            for e in source.get("evidence") or []:
                if e.get("side") != side:
                    continue
                origin = e.get("origin") or None
                doc = (origin or {}).get("source") if origin else e.get("document")
                doc = doc or {}
                cipher = str((doc.get("document_code") if origin else doc.get("label")) or "").strip() or "—"
                raw_page = ((origin.get("page") or {}).get("source_page")) if origin else e.get("page")
                source_type = e.get("source_type") if e.get("source_type") in SOURCE_TYPES else "—"
                key = (cipher, doc.get("version_id") or doc.get("version"), doc.get("pdf_path") or doc.get("id"),
                       source_type)
                group = groups.setdefault(key, {"cipher": cipher, "source_type": source_type, "pages": set()})
                try:
                    page = int(raw_page)
                except (TypeError, ValueError):
                    continue
                if page > 0:
                    group["pages"].add(page)
        result[side] = [{"cipher": g["cipher"], "source_type": g["source_type"],
                         "pages": ", ".join(map(str, sorted(g["pages"]))) or "—"} for g in groups.values()]
    return result


def _row_height(texts: list[tuple[str, float]], minimum: float = 15.0) -> float:
    """Высота строки по самой «длинной» ячейке: строки текста с учётом переносов по ширине."""
    lines = 1
    for text, width in texts:
        chars = max(int(width * 1.15), 8)
        lines = max(lines, sum(max(1, -(-len(part) // chars)) for part in str(text or "").split("\n")))
    return min(409.0, max(minimum, lines * 13.0 + 6))


class _Final:
    """Лист «Итоговые»: строка ИТ (уровень 0) + скрытые подробности (уровень 1)."""

    # A–F — строка ИТ как на портале; G–J нужны таблице характеристик, где у «Было» и
    # «Стало» свои «Шифр · Источник · Страница по PDF» (9 столбцов, как на портале).
    WIDTHS = {"A": 12.7, "B": 70.7, "C": 50.0, "D": 50.0, "E": 16.7, "F": 26.7,
              "G": 50.0, "H": 30.0, "I": 12.0, "J": 16.0}
    LAST = 10
    PARAM_HEAD = ("Значение", "Шифр", "Источник", "Страница по PDF")

    def __init__(self, ws, base_url: str):
        self.ws, self.base_url, self.row = ws, base_url.rstrip("/"), 0
        for col, width in self.WIDTHS.items():
            ws.column_dimensions[col].width = width

    def _width(self, first: int, last: int) -> float:
        return sum(self.WIDTHS[get_column_letter(c)] for c in range(first, last + 1))

    def put(self, cells: list[tuple[int, int, Any, dict[str, Any]]], *, level: int = 0, height: float | None = None) -> int:
        """cells: (первая колонка, последняя колонка, значение, стиль) — диапазоны объединяются."""
        self.row += 1
        r, ws = self.row, self.ws
        sizing = []
        for first, last, value, style in cells:
            cell = ws.cell(r, first, value)
            cell.font = Font(bold=style.get("bold", False), size=style.get("size", 10), color=style.get("color", INK))
            cell.alignment = TOP
            if style.get("fill"):
                for c in range(first, last + 1):
                    ws.cell(r, c).fill = _fill(style["fill"])
            if style.get("grid"):
                for c in range(first, last + 1):
                    ws.cell(r, c).border = BORDER
            if style.get("link"):
                cell.hyperlink = style["link"]
            if last > first:
                ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
            sizing.append((value, self._width(first, last)))
        ws.row_dimensions[r].height = height or _row_height(sizing, 23.0 if level else 15.0)
        if level:
            ws.row_dimensions[r].outline_level = level
            ws.row_dimensions[r].hidden = True
        return r

    def section(self, text: str) -> None:
        self.put([(1, self.LAST, text, {"bold": True, "size": 11, "color": TEAL, "fill": F_SECTION})], level=1)

    def detail(self, text: str) -> None:
        self.put([(1, self.LAST, text, {"fill": F_DETAIL})], level=1)

    def head(self, left: str, mid: str | None, right: str, split: int = 2) -> None:
        style = {"bold": True, "size": 11, "color": HEAD_INK, "fill": F_HEAD}
        cells = [(1, split, left, style)]
        if mid is not None:
            cells.append((split + 1, split + 1, mid, style))
        cells.append((split + 1 if mid is None else split + 2, self.LAST, right, style))
        self.put(cells, level=1)

    def param_head(self) -> None:
        """Двухуровневая шапка: «Было (ПД)» / «Стало (РД)», под ними Значение · Шифр · Источник · Страница."""
        name = {"bold": True, "size": 11, "color": TEAL, "fill": F_NAME_HEAD, "grid": True}
        old = {"bold": True, "size": 11, "color": INK, "fill": F_OLD_HEAD, "grid": True}
        new = {"bold": True, "size": 11, "color": INK, "fill": F_NEW_HEAD, "grid": True}
        self.put([(1, 2, "Характеристика", name), (3, 6, "Было (ПД)", old), (7, 10, "Стало (РД)", new)], level=1)
        sub = lambda fill: {"bold": True, "size": 10, "color": HEAD_INK, "fill": fill, "grid": True}  # noqa: E731
        self.put([(1, 2, "", sub(F_NAME_HEAD))]
                 + [(c, c, h, sub(F_OLD_HEAD)) for c, h in enumerate(self.PARAM_HEAD, 3)]
                 + [(c, c, h, sub(F_NEW_HEAD)) for c, h in enumerate(self.PARAM_HEAD, 7)], level=1)

    def param_group(self, name: str, status: str) -> None:
        self.put([(1, 2, name, {"bold": True, "color": TEAL, "fill": F_NAME, "grid": True}),
                  (3, self.LAST, f"{status} · значения по каждому месту — ниже",
                   {"bold": True, "color": STATUS_INK, "fill": F_STATUS, "grid": True})], level=1)

    def param(self, name: str, old: str, new: str, sources: dict[str, list[dict[str, str]]] | None = None) -> None:
        cells = [(1, 2, name, {"bold": True, "color": TEAL, "fill": F_NAME, "grid": True})]
        for first, (side, value) in ((3, ("OLD", old)), (7, ("NEW", new))):
            fill = F_OLD if side == "OLD" else F_NEW
            groups = (sources or {}).get(side) or []
            column = lambda key: "\n".join(g[key] for g in groups) or "—"  # noqa: E731
            types = {g["source_type"] for g in groups}
            type_ink = SOURCE_INK.get(next(iter(types)), MUTED) if len(types) == 1 else MUTED
            cells += [(first, first, value, {"bold": True, "fill": fill, "grid": True}),
                      (first + 1, first + 1, column("cipher"), {"color": MUTED, "fill": fill, "grid": True}),
                      (first + 2, first + 2, column("source_type"), {"bold": True, "size": 9, "color": type_ink,
                                                                    "fill": fill, "grid": True}),
                      (first + 3, first + 3, column("pages"), {"color": MUTED, "fill": fill, "grid": True})]
        self.put(cells, level=1)

    def link(self, e: dict[str, Any] | None) -> tuple[str, dict[str, Any]]:
        if not e:
            return "", {"fill": F_DETAIL}
        style = {"color": LINK, "fill": F_DETAIL}
        if e.get("image_url"):
            style["link"] = (self.base_url + e["image_url"]) if e["image_url"].startswith("/") else e["image_url"]
        return f"Открыть {e.get('side')} · {e.get('source_type') or '—'} · стр. {e.get('page')}", style

    @staticmethod
    def fragment(e: dict[str, Any] | None) -> str:
        if not e:
            return ""
        head = " · ".join(x for x in (str((e.get("document") or {}).get("label") or ""), f"стр. {e.get('page')}",
                                       str(e.get("source_type") or "")) if x)
        return "\n".join(x for x in (head, e.get("short_explanation_ru") or "", e.get("quote") or "") if x)

    def write(self, view: dict[str, Any], *, document_code: str, object_label: str, exported: str) -> None:
        ws, items = self.ws, view.get("consolidated") or []
        pc_no = {str(c.get("id")): f"PC-{i + 1:03d}" for i, c in enumerate(view.get("original") or [])}
        n = len(items)
        self.put([(1, self.LAST, f"{document_code} · Итоговые", {"bold": True, "size": 16, "fill": F_TITLE})], height=30)
        self.put([(1, self.LAST, f"{object_label} · {n} {plural(n, 'изменение', 'изменения', 'изменений')} · Выгрузка {exported}",
                   {"bold": True, "size": 11})], height=24)
        self.put([(1, self.LAST, "Нажмите «+» слева от строки изменения, чтобы раскрыть подробности. Кнопки уровней 1 / 2 "
                         "сворачивают или раскрывают все группы. Только текст, без PDF и изображений.",
                   {"fill": F_DETAIL})], height=28)
        self.put([(1, self.LAST, f"Прогон: {view.get('source_run_id')} · Итоговая сводка: {view.get('consolidator_run_id')} "
                         "· Статусы сохранены как в веб-версии.", {"color": MUTED})], height=22)
        style = {"bold": True, "size": 11, "color": HEAD_INK, "fill": F_HEAD}
        self.put([(c, c, h, style) for c, h in enumerate(("ID", "Изменение", "Было", "Стало", "Источник", "Статус"), 1)],
                 height=26)
        originals = view.get("original") or []
        for i, item in enumerate(items):
            self._item(i, item, pc_no, originals)
        ws.freeze_panes = "C6"
        ws.print_title_rows = "1:5"
        ws.sheet_properties.outlinePr.summaryBelow = False

    def _item(self, i: int, c: dict[str, Any], pc_no: dict[str, str], originals: list[dict[str, Any]]) -> None:
        title = str(c.get("title") or c.get("summary") or "Изменение")
        lines = [title]
        if c.get("summary") and c["summary"] != title:
            lines.append(c["summary"])
        lines.append(f"Сводная · {len(c['lineage']['members'])} исходных" if c.get("kind") == "CONSOLIDATED"
                     else f"Исходная карточка {c.get('projectchange_id') or c.get('id')}")
        text = {"size": 11}
        self.put([(1, 1, it_id(i), {"color": MUTED}), (2, 2, "\n".join(lines), {**text, "bold": True}),
                  (3, 3, c.get("old_state") or "Состояние не установлено", text),
                  (4, 4, c.get("new_state") or "Состояние не установлено", text),
                  (5, 5, "\n".join(_sources(c)), {"color": LINK}),
                  (6, 6, CHANNELS.get(c.get("channel"), c.get("channel") or ""), {"color": "AD6900", "fill": "FFF0C5"})])
        params = c.get("parameters") or []
        if params:
            self.section(f"Изменившиеся характеристики · {len(params)}")
            self.param_head()
            for p in params:
                unit = p.get("unit") or ""
                sources = parameter_sources(c, p, originals)
                values = p.get("values") or []
                if c.get("kind") == "CONSOLIDATED" and p.get("status") != "CONSISTENT" and values:
                    # No single old/new: a group row, then one row per place with
                    # that place's values and its own member card's sources (as on the portal).
                    self.param_group(p.get("name") or "", PARAM_STATUS.get(p.get("status"), p.get("status") or ""))
                    for v in values:
                        vu = v.get("unit") or ""
                        self.param(f"По месту: {v.get('location') or v.get('card_id') or '—'}",
                                   _value(v.get("old_value") or "", vu), _value(v.get("new_value") or "", vu),
                                   parameter_sources(c, {"values": [v]}, originals))
                elif c.get("kind") == "CONSOLIDATED":
                    self.param(p.get("name") or "", _value(p.get("old") or "", unit), _value(p.get("new") or "", unit),
                               sources)
                else:
                    name = p.get("name") or ""
                    if p.get("location"):
                        name += f" · {p['location']}"
                    self.param(name, _value(p.get("old") or "", unit), _value(p.get("new") or "", unit), sources)
        conflicts = c.get("conflicts") or []
        if conflicts:
            self.section("Противоречия / недоказанность источников")
            for x in conflicts:
                self.detail(f"Конфликт: {x.get('statement') or ''}")
        if c.get("why_one_event"):
            self.detail(f"Основание: {c['why_one_event']}")
        if c.get("kind") == "CONSOLIDATED":
            self.section("Состав итоговой карточки")
            for m in c["lineage"]["members"]:
                body = " · ".join(x for x in (m.get("id"), m.get("region_id"), m.get("title")) if x)
                if m.get("summary"):
                    body += "\n" + m["summary"]
                self.put([(1, 1, pc_no.get(str(m.get("id")), ""), {"color": LINK, "fill": F_DETAIL}),
                          (2, self.LAST, body, {"fill": F_DETAIL})], level=1)
            if c["lineage"].get("distinguishing_check"):
                self.detail(f"Различия: {c['lineage']['distinguishing_check']}")
            for m in c.get("manifestations") or []:
                body = "; ".join(m.get("locations") or [])
                if m.get("note"):
                    body += "\n" + m["note"]
                self.detail(f"Проявление: {body}")
        evidence = c.get("evidence") or []
        if evidence:
            old = [e for e in evidence if e.get("side") == "OLD"]
            new = [e for e in evidence if e.get("side") == "NEW"]
            self.section(f"Фрагменты источников — текст · {len(evidence)}")
            word = lambda k: f"{k} {plural(k, 'фрагмент', 'фрагмента', 'фрагментов')}"  # noqa: E731
            self.head(f"OLD — Было · {word(len(old))}", None, f"NEW — Стало · {word(len(new))}", split=3)
            for a, b in zip_longest(old, new):
                (la, sa), (lb, sb) = self.link(a), self.link(b)
                self.put([(1, 3, la, sa), (4, self.LAST, lb, sb)], level=1)
                self.put([(1, 3, self.fragment(a), {"fill": F_DETAIL}),
                          (4, self.LAST, self.fragment(b), {"fill": F_DETAIL})],
                         level=1)


def _table_head(ws, row: int, columns) -> None:
    for col, (_, header, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True, size=11, color="FFFFFF")
        cell.fill = _fill(DARK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def _table_row(ws, row: int, columns, values: dict[str, str], bold: tuple[str, ...]) -> None:
    sizing = []
    for col, (key, _, width) in enumerate(columns, 1):
        value = values.get(key) or ""
        cell = ws.cell(row, col, value or None)
        cell.font = Font(bold=key in bold, size=10)
        cell.alignment = TOP
        cell.border = BORDER
        if key == "note":
            cell.fill = _fill(NOTE_FILL)
        sizing.append((value, width))
    ws.row_dimensions[row].height = _row_height(sizing, 36.0)


def _landscape(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def template_rows(view: dict[str, Any]) -> list[dict[str, str]]:
    """Строки шаблона «Карточек ИТ»: то, что известно из портала; разбор — пустой."""
    return [{"id": it_id(i), "title": str(c.get("title") or c.get("summary") or ""), "old": c.get("old_state") or "",
             "new": c.get("new_state") or "", "where": _pages(c)} for i, c in enumerate(view.get("consolidated") or [])]


def conditional_branches(rows: list[dict[str, str]]) -> set[str]:
    """Хвосты родителя с двумя и более ветками — взаимоисключающие варианты («сначала лист»)."""
    tails = defaultdict(list)
    for r in rows:
        m = ID_RE.match(r.get("id") or "")
        if m and m.group(2):
            tails[m.group(1)].append(r["id"])
    return {t for group in tails.values() if len(group) > 1 for t in group}


def _stat(rows: list[dict[str, str]], basket: str, branches: set[str]) -> str:
    total = sum(1 for r in rows if r.get("basket") == basket)
    cond = sum(1 for r in rows if r.get("basket") == basket and r.get("id") in branches)
    return f"{total} (в т.ч. условных веток: {cond})" if cond else str(total)


def _write_cards(ws, rows: list[dict[str, str]], *, filled: bool) -> None:
    _table_head(ws, 1, CARD_COLUMNS)
    for r, values in enumerate(rows, 2):
        _table_row(ws, r, CARD_COLUMNS, values, bold=("id", "basket"))
    last = max(len(rows) + 1, 2)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CARD_COLUMNS))}{last}"
    if not filled:  # шаблон: выпадающие списки, чтобы корзины и уверенность писались ровно как в легенде
        for col, values in (("B", BASKETS), ("C", CONFIDENCE)):
            dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
            dv.add(f"{col}2:{col}{max(last, 200)}")
            ws.add_data_validation(dv)
    _landscape(ws)


def _write_summary(ws, rows: list[dict[str, str]], *, document_code: str, n_items: int) -> None:
    width = len(SUMMARY_COLUMNS)
    for col, (_, _, w) in enumerate(SUMMARY_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.cell(1, 1, f"{short_code(document_code)} · краткая сводка ПД→РД (промежуточный разбор)").font = Font(
        bold=True, size=14, color=DARK)
    ws.cell(2, 1, f"{n_items} ИТ + хвосты .2/.3. Полный доп. = вырезать комплекс Δ. Изменение = дельта, не с нуля. "
                  "Оптимизация = минус. Детализация = не подаём (ссылка на ПД в карточке). Уверенность «сначала лист» "
                  "и «после АИ» — не в письмо заказчику до сверки листа.").font = Font(size=10)
    for r, h in ((1, 22), (2, 40)):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        ws.cell(r, 1).alignment = TOP
        ws.row_dimensions[r].height = h
    branches = conditional_branches(rows)
    counts = Counter(r.get("basket") for r in rows)
    stats = ("Полный доп.", _stat(rows, BASKET_FULL, branches), "Изменение", _stat(rows, BASKET_CHANGE, branches),
             "Оптимизация / детализация", f"{counts[BASKET_OPT]} / {counts[BASKET_DETAIL]}")
    for col, value in enumerate(stats, 1):
        cell = ws.cell(3, col, value)
        cell.font = Font(bold=True, size=10)
        cell.alignment = TOP
    _table_head(ws, 5, SUMMARY_COLUMNS)
    for r, values in enumerate(rows, 6):
        brief = values.get("brief") or first_sentence(values.get("resume") or "")
        _table_row(ws, r, SUMMARY_COLUMNS, {**values, "brief": brief}, bold=("id",))
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:{get_column_letter(width)}{max(len(rows) + 5, 6)}"
    _landscape(ws)


def build_report(view: dict[str, Any], *, document_code: str, object_label: str,
                 analysis: dict[str, Any] | None = None, base_url: str = "",
                 exported_at: datetime | None = None) -> bytes:
    """xlsx: «Итоговые» из портала + «Карточки ИТ» и «Краткая сводка» из разбора (или шаблон)."""
    exported = (exported_at or datetime.now(timezone.utc)).astimezone(MSK).strftime("%d.%m.%Y")
    wb = Workbook()
    final = wb.active
    final.title = SHEETS[0]
    _Final(final, base_url).write(view, document_code=document_code, object_label=object_label, exported=exported)
    _landscape(final)
    rows = list((analysis or {}).get("rows") or []) or template_rows(view)
    _write_cards(wb.create_sheet(SHEETS[1]), rows, filled=bool(analysis))
    _write_summary(wb.create_sheet(SHEETS[2]), rows, document_code=document_code,
                   n_items=len(view.get("consolidated") or []))
    wb.properties.creator = "AuditManager"
    wb.properties.title = f"{document_code} — итоговые изменения и разбор ИТ"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
