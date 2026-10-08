"""Excel-отчёт «Итоговые · Карточки ИТ · Краткая сводка» и загрузка разбора по корзинам."""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from backend.app.services.project_change_consolidator import it_analysis
from backend.app.services.project_change_consolidator import report_xlsx as R
from backend.tests.project_change_consolidator.test_consolidated_view import _base, _production_run, _shadow
from backend.tests.project_change_consolidator.test_e2e_production_boundary import env  # noqa: F401


def _ev(side, page, n, source="TEXT"):
    return {"id": f"e{side}{n}", "side": side, "page": page, "source_type": source, "quote": f"цитата {n}",
            "short_explanation_ru": f"{side}-состояние {n}", "document": {"label": "СТ-ИОС2.1"},
            "image_url": f"/api/crop/{side}{n}"}


def _view():
    merged = {
        "kind": "CONSOLIDATED", "id": "pcc:1", "channel": "REVIEW", "title": "Уклон", "summary": "Общее требование",
        "old_state": "без уклона", "new_state": "0,001", "why_one_event": "одно требование",
        "parameters": [{"name": "Уклон", "old": "0", "new": "0,001", "unit": "—",
                        "values": [{"card_id": "PC-A", "old_value": "0", "new_value": "0,001", "unit": "—",
                                    "location": "Корпус 1"}]},
                       {"name": "Толщина", "old": "не менее 10 мм (по расчёту)", "new": "20", "unit": "мм", "values": []}],
        "manifestations": [{"locations": ["Корпус 1", "Корпус 2"], "note": "п. 15"}],
        "evidence": [_ev("OLD", 10, 1), _ev("NEW", 4, 1, "TABLE"), _ev("NEW", 48, 2)],
        "conflicts": [{"source": "HINT", "statement": "Растр OLD отсутствует"}],
        "lineage": {"members": [{"id": "PC-B", "region_id": "R-2", "title": "Уклон 2", "summary": "s"}],
                    "distinguishing_check": "разные корпуса"},
    }
    single = {"kind": "SOURCE_CARD", "id": "pcs:2", "projectchange_id": "PC-A", "channel": "REVIEW",
              "title": "Отметки", "summary": "", "old_state": "", "new_state": "-0,770", "why_one_event": "",
              "parameters": [{"name": "Отметка 1 этажа", "old": "-0,620", "new": "-0,770", "unit": "м",
                              "location": "Корпус 3"}], "evidence": [_ev("OLD", 38, 3, "GRAPHIC")], "conflicts": []}
    return {"source_run_id": "src1", "consolidator_run_id": "cons1", "shadow_result_sha256": "b" * 64,
            "original": [{"id": "PC-A"}, {"id": "PC-B"}], "consolidated": [merged, single]}


def _book(data):
    return load_workbook(io.BytesIO(data))


def _build(view=None, analysis=None):
    return R.build_report(view or _view(), document_code="СТ26-01-14-ИОС2.1", object_label="256. Объект",
                          analysis=analysis, base_url="http://portal/")


def test_three_sheets_in_order_and_final_sheet_mirrors_the_portal():
    wb = _book(_build())
    assert wb.sheetnames == ["Итоговые", "Карточки ИТ", "Краткая сводка"]
    ws = wb["Итоговые"]
    assert ws["A1"].value == "СТ26-01-14-ИОС2.1 · Итоговые"
    assert ws["A2"].value.startswith("256. Объект · 2 изменения · Выгрузка ")
    assert "Прогон: src1 · Итоговая сводка: cons1" in ws["A4"].value
    assert [ws.cell(5, c).value for c in range(1, 7)] == ["ID", "Изменение", "Было", "Стало", "Источник", "Статус"]
    assert ws["A6"].value == "ИТ-001" and ws["B6"].value == "Уклон\nОбщее требование\nСводная · 1 исходных"
    assert ws["E6"].value == "TEXT\nTABLE" and ws["F6"].value == "Нужна проверка инженера"
    assert ws.freeze_panes == "C6" and ws.row_dimensions[6].outline_level == 0
    detail = {}
    for r in range(7, ws.max_row + 1):
        if ws.row_dimensions[r].outline_level == 0:
            break
        assert ws.row_dimensions[r].hidden
        detail[ws.cell(r, 1).value] = r
    assert "Изменившиеся характеристики · 2" in detail and "Состав итоговой карточки" in detail
    assert ws.cell(detail["По месту: Корпус 1"], 3).value == "0 —"
    assert ws.cell(detail["Толщина"], 3).value == "не менее 10 мм (по расчёту)"  # единица уже в значении
    assert ws.cell(detail["Толщина"], 4).value == "20 мм"
    assert "Конфликт: Растр OLD отсутствует" in detail and "Основание: одно требование" in detail
    assert ws.cell(detail["PC-002"], 2).value == "PC-B · R-2 · Уклон 2\ns"  # PC-NNN — номер на листе «Исходные»
    assert "Проявление: Корпус 1; Корпус 2\nп. 15" in detail and "Различия: разные корпуса" in detail
    head = detail["OLD — Было · 1 фрагмент"]
    assert ws.cell(head, 4).value == "NEW — Стало · 2 фрагмента"
    link = ws.cell(head + 1, 1)
    assert link.value == "Открыть OLD · TEXT · стр. 10" and link.hyperlink.target == "http://portal/api/crop/OLD1"
    assert ws.cell(head + 2, 1).value == "СТ-ИОС2.1 · стр. 10 · TEXT\nOLD-состояние 1\nцитата 1"
    assert ws.cell(head + 3, 1).value is None and ws.cell(head + 3, 4).value == "Открыть NEW · TEXT · стр. 48"
    second = next(r for r in range(7, ws.max_row + 1) if ws.cell(r, 1).value == "ИТ-002")
    assert ws.cell(second, 3).value == "Состояние не установлено"
    assert ws.cell(second + 3, 1).value == "Отметка 1 этажа · Корпус 3" and ws.cell(second + 3, 3).value == "-0,620 м"


def test_without_analysis_the_cards_are_a_template_with_legend_lists():
    wb = _book(_build())
    cards = wb["Карточки ИТ"]
    assert [cards.cell(1, c).value for c in range(1, 15)] == [h for _, h, _ in R.CARD_COLUMNS]
    assert [cards.cell(r, 1).value for r in (2, 3)] == ["ИТ-001", "ИТ-002"]
    assert cards["B2"].value is None and cards["D2"].value == "Уклон" and cards["F2"].value == "без уклона"
    assert cards["M2"].value == "ПД стр. 10; РД стр. 4, 48"
    lists = {str(dv.sqref).split(":")[0]: dv.formula1 for dv in cards.data_validations.dataValidation}
    assert lists["B2"] == '"' + ",".join(R.BASKETS) + '"' and lists["C2"] == '"' + ",".join(R.CONFIDENCE) + '"'
    summary = wb["Краткая сводка"]
    assert summary["A1"].value == "ИОС2.1 · краткая сводка ПД→РД (промежуточный разбор)"
    assert "Жёлт" not in summary["A2"].value
    assert [summary.cell(3, c).value for c in range(1, 7)] == [
        "Полный доп.", "0", "Изменение", "0", "Оптимизация / детализация", "0 / 0"]


def _filled(rows):
    wb = _book(_build())
    ws = wb["Карточки ИТ"]
    ws.delete_rows(2, ws.max_row)
    for r, values in enumerate(rows, 2):
        for c, (key, _, _) in enumerate(R.CARD_COLUMNS, 1):
            ws.cell(r, c, values.get(key))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


ROWS = [
    {"id": "ИТ-001", "basket": R.BASKET_DETAIL, "confidence": "сначала лист", "title": "Уклон",
     "resume": "Версии расходятся. Ветки ниже.", "note": "Примечание"},
    {"id": "ИТ-001.2", "basket": R.BASKET_FULL, "confidence": "сначала лист", "resume": "Новая группа."},
    {"id": "ИТ-001.3", "basket": R.BASKET_CHANGE, "confidence": "сначала лист", "resume": "Только Ду."},
    {"id": "ИТ-002", "basket": R.BASKET_FULL, "confidence": "уверен", "resume": "На 14 эт. корп. 14.1 больше."},
]


def test_filled_analysis_round_trip_without_basket_fills():
    rows, warnings = it_analysis.parse(_filled(ROWS), n_items=2)
    assert [r["id"] for r in rows] == ["ИТ-001", "ИТ-001.2", "ИТ-001.3", "ИТ-002"] and warnings == []
    wb = _book(_build(analysis={"rows": rows}))
    cards = wb["Карточки ИТ"]
    assert cards["B3"].value == R.BASKET_FULL and cards["B3"].fill.fill_type is None
    assert cards["C3"].fill.fill_type is None and cards["N2"].fill.fgColor.rgb.endswith(R.NOTE_FILL)
    assert not cards.data_validations.dataValidation
    summary = wb["Краткая сводка"]
    assert summary["B3"].value == "2 (в т.ч. условных веток: 1)" and summary["D3"].value == "1 (в т.ч. условных веток: 1)"
    assert summary["F3"].value == "0 / 1" and summary["B3"].fill.fill_type is None
    assert summary["E6"].value == "Версии расходятся." and summary["E9"].value == "На 14 эт."  # первое предложение


def test_manual_brief_on_the_summary_survives_reupload():
    rows, _ = it_analysis.parse(_filled(ROWS), n_items=2)
    wb = _book(_build(analysis={"rows": rows}))
    wb["Краткая сводка"]["E9"] = "Ручная формулировка"
    out = io.BytesIO()
    wb.save(out)
    again, _ = it_analysis.parse(out.getvalue(), n_items=2)
    assert again[-1]["brief"] == "Ручная формулировка"
    assert _book(_build(analysis={"rows": again}))["Краткая сводка"]["E9"].value == "Ручная формулировка"


def test_legend_violations_are_rejected_and_gaps_are_warned():
    bad = [{**ROWS[0], "basket": "Доп. работы"}, {"id": "ИТ-7", "basket": R.BASKET_FULL},
           {"id": "ИТ-003", "basket": R.BASKET_FULL, "confidence": "наверное"}, {**ROWS[3]}, {**ROWS[3]}]
    with pytest.raises(it_analysis.AnalysisError) as exc:
        it_analysis.parse(_filled(bad), n_items=2)
    text = "\n".join(exc.value.errors)
    for needle in ("корзина «Доп. работы»", "ID «ИТ-7»", "ИТ-003 — в итоговой сводке 2 ИТ", "«наверное»", "ИТ-002 повторяется"):
        assert needle in text
    rows, warnings = it_analysis.parse(_filled([{"id": "ИТ-002", "basket": R.BASKET_OPT}]), n_items=2)
    assert "ИТ-002: не заполнены корзина и/или уверенность" in warnings and "Нет карточек: ИТ-001" in warnings
    with pytest.raises(it_analysis.AnalysisError):
        it_analysis.parse(b"not an xlsx", n_items=2)


def test_report_and_analysis_upload_through_the_api(env):
    pytest.importorskip("jsonschema")  # теневой прогон Консолидатора (requirements-projectchange-v3-runtime)
    client = env["client"]
    run_id = _production_run(env)
    manifest = _shadow(env, run_id)
    url = _base(env) + f"/consolidated/{run_id}/{manifest['consolidator_run_id']}"
    assert client.get(url + "/it-analysis").json() == {"schema": "projectchange-it-analysis-status/1", "available": False}

    template = client.get(url + "/report.xlsx")
    assert template.status_code == 200 and "_%D1%88%D0%B0%D0%B1%D0%BB%D0%BE%D0%BD.xlsx" in template.headers["content-disposition"]
    wb = _book(template.content)
    assert wb.sheetnames == list(R.SHEETS)
    n = sum(1 for r in range(2, wb["Карточки ИТ"].max_row + 1) if wb["Карточки ИТ"].cell(r, 1).value)
    assert n == len(client.get(url).json()["consolidated"])

    ws = wb["Карточки ИТ"]
    for r in range(2, n + 2):
        ws.cell(r, 2, R.BASKET_CHANGE)
        ws.cell(r, 3, "уверен")
    out = io.BytesIO()
    wb.save(out)
    files = {"file": ("разбор.xlsx", out.getvalue(), "application/octet-stream")}
    status = client.post(url + "/it-analysis", files=files).json()
    assert status["available"] and status["rows"] == n and status["filled"] == n
    client.post(url + "/it-analysis", files=files)  # повторная загрузка: прежний разбор уходит в history

    report = client.get(url + "/report.xlsx")
    assert "%D1%87%D0%B5%D1%82%D1%8B%D1%80%D0%B5_%D0%BA%D0%BE%D1%80%D0%B7%D0%B8%D0%BD%D1%8B" in report.headers["content-disposition"]
    assert _book(report.content)["Краткая сводка"]["B3"].value == "0" and \
        _book(report.content)["Карточки ИТ"]["B2"].value == R.BASKET_CHANGE

    bad = client.post(url + "/it-analysis", files={"file": ("x.xlsx", b"junk", "application/octet-stream")})
    assert bad.status_code == 422 and bad.json()["detail"]["errors"]
    from backend.app.services.stage_comparison import paths
    base = paths.pair_dir(env["session_id"], _base(env).rsplit("/", 1)[-1]) / it_analysis.DIR_NAME / run_id \
        / manifest["consolidator_run_id"]
    assert len(list((base / "history").glob("*.json"))) == 1 and len(list((base / "uploads").glob("*.xlsx"))) == 2
    assert json.loads((base / "it_analysis.json").read_text(encoding="utf-8"))["schema"] == it_analysis.SCHEMA
    assert client.get(url.replace(run_id, "..%2Fx") + "/report.xlsx").status_code == 404
