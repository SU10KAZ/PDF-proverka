#!/usr/bin/env python3
"""Build the isolated Critical 8 ChatGPT review package from frozen r008 data.

This script is deliberately local-only.  It performs no model or network calls
and adds no review judgments.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RESEARCH_ROOT = Path(
    "/home/coder/auditmanager/corpus-audits/20260915_openrouter_leg_replacement"
)
SOURCE_PREP = RESEARCH_ROOT / "revisions/r008_known24_source_review_prep"
PHASE2 = RESEARCH_ROOT / "revisions/r008_phase2_codex_retention_candidate"
OUTPUT = RESEARCH_ROOT / "r008_critical8_chatgpt_review_package"


def read_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value.rstrip() + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_record(path: Path, role: str, known_id: str | None = None) -> dict[str, Any]:
    record: dict[str, Any] = {
        "bytes": path.stat().st_size,
        "evidence_role": role,
        "sha256": sha256(path),
    }
    if known_id:
        record["known_id"] = known_id
    return record


def safe_name(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._-]+", "_", value).strip("_.")
    return value or "source"


def configure_sheet(ws, widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def write_workbook(path: Path, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Critical 8"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    configure_sheet(sheet, widths)
    workbook.save(path)


def extract_pages(source: Path, pages: list[int], destination: Path) -> None:
    if not pages:
        raise ValueError(f"No physical pages supplied for {source}")
    work = destination.parent / ".pdf_parts"
    work.mkdir(parents=True, exist_ok=True)
    parts: list[Path] = []
    for index, page in enumerate(pages, 1):
        pattern = work / f"part-{index}-%d.pdf"
        subprocess.run(
            ["pdfseparate", "-f", str(page), "-l", str(page), str(source), str(pattern)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
        generated = work / f"part-{index}-{page}.pdf"
        if not generated.is_file():
            raise FileNotFoundError(generated)
        parts.append(generated)
    if len(parts) == 1:
        shutil.move(parts[0], destination)
    else:
        subprocess.run(
            ["pdfunite", *map(str, parts), str(destination)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
    shutil.rmtree(work)


def zip_tree(source: Path, destination: Path, root_name: str) -> None:
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                archive.write(path, Path(root_name) / path.relative_to(source))


def zip_master(destination: Path, files: list[Path]) -> None:
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as archive:
        for path in files:
            archive.write(path, path.name)


def json_file_inventory(root: Path) -> list[dict[str, Any]]:
    return [
        {
            "path": str(path.relative_to(root)),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in sorted(root.rglob("*"))
        if path.is_file()
    ]


def build() -> dict[str, Any]:
    if OUTPUT.exists():
        raise RuntimeError(f"Refusing to overwrite existing output: {OUTPUT}")

    review_manifest = read_json(SOURCE_PREP / "CRITICAL8_REVIEW_MANIFEST.json")
    phase2_manifest = read_json(PHASE2 / "CRITICAL8_PHASE2_MANIFEST.json")
    phase2_results = read_json(PHASE2 / "CODEX_PHASE2_RESULTS.json")
    review_records = review_manifest["records"]
    phase2_by_known = {item["known_id"]: item for item in phase2_manifest["records"]}
    result_by_case = {item["case_id"]: item for item in phase2_results["records"]}
    review_ids = {item["known_id"] for item in review_records}
    if len(review_records) != 8 or review_ids != set(phase2_by_known):
        raise RuntimeError("Frozen Critical 8 manifests do not agree exactly")

    step1 = OUTPUT / "STEP1_SOURCE_FIRST"
    step2 = OUTPUT / "STEP2_CODEX_COMPARISON"
    original_pdfs = step1 / "original_pdfs"
    original_pdfs.mkdir(parents=True)
    (step1 / "cases").mkdir()
    (step2 / "cases").mkdir(parents=True)

    write_text(
        step1 / "00_INSTRUCTIONS.txt",
        """ЗАДАЧА

Независимо проверить 8 исторических критических замечаний по исходным документам.
Не считать historical claim автоматически правильным.

Для каждого case определить один source_verdict:
SUPPORTED
PARTIALLY_SUPPORTED
NOT_SUPPORTED
INSUFFICIENT_DATA

Для каждого case указать:
- OLD evidence и NEW evidence, только если такое разделение действительно есть в источниках;
- физические страницы исходного PDF;
- использованные модальности TEXT / TABLE / GRAPHIC;
- короткое объяснение;
- что остаётся недоказанным.

Особенно проверить:
- тот ли engineering subject;
- тот ли scope;
- сопоставимы ли OLD и NEW;
- не выдано ли отсутствие текста за отсутствие инженерного решения;
- нет ли source conflict.

PDF является authoritative source. EVIDENCE.json содержит машинно извлечённые
фрагменты без ручного исправления; crop/растр является вспомогательным материалом.

На этом шаге ответы Codex НЕ доступны. Не запрашивайте STEP2 до фиксации всех
восьми source-first выводов. Запишите выводы в пустые колонки source_verdict и
comment файла CRITICAL8_OVERVIEW.xlsx либо верните отдельную таблицу.
""",
    )
    write_text(
        step2 / "00_INSTRUCTIONS.txt",
        """ЗАДАЧА

Используй уже зафиксированные source-first verdicts из STEP1. Не меняй их после
просмотра результатов Codex.

Теперь сравни результат Codex с source-first выводом. Для каждого case определи:
RETAINED_STRONG
RETAINED_PARTIAL
CODEX_MISS
HISTORICAL_NOT_SUPPORTED
PACKAGING_GAP
AMBIGUOUS

Проверить совпадение engineering subject, scope, issue, parameter/state и
evidence. Не считать сходство текста достаточным доказательством.

Колонки retention_verdict и review_comment в CODEX_CRITICAL8_OVERVIEW.xlsx
намеренно пусты. Никакие результаты automatic matcher и ожидаемые retention
verdicts в этот пакет не включены.
""",
    )

    pdf_destinations: dict[str, str] = {}
    critical_manifest_records: list[dict[str, Any]] = []
    overview_rows: list[list[Any]] = []
    codex_rows: list[list[Any]] = []
    referenced_files: list[tuple[Path, str, str]] = []

    for frozen in review_records:
        known_id = frozen["known_id"]
        case_id = frozen["case_id"]
        package = Path(frozen["package_path"])
        case_meta = read_json(package / "CASE.json")
        historical = read_json(package / "HISTORICAL.json")
        codex_prep = read_json(package / "CODEX.json")
        phase2 = phase2_by_known[known_id]
        result = result_by_case[case_id]

        pages = sorted({int(page) for page in historical.get("pages", [])})
        source_pdf_assets = [
            asset for asset in historical["source_assets"] if asset.get("kind") == "source_pdf"
        ]
        crop_assets = [
            asset for asset in historical["source_assets"] if asset.get("kind") != "source_pdf"
        ]
        if len(source_pdf_assets) != 1:
            raise RuntimeError(f"Expected one authoritative PDF for {known_id}")
        source_asset = source_pdf_assets[0]
        source_pdf = Path(source_asset["resolved_path"])
        if not source_pdf.is_file() or source_pdf.read_bytes()[:5] != b"%PDF-":
            raise RuntimeError(f"Invalid source PDF for {known_id}: {source_pdf}")
        if sha256(source_pdf) != source_asset["sha256"]:
            raise RuntimeError(f"Source PDF hash mismatch for {known_id}")

        source_document = historical["source_documents"][0]
        pdf_hash = source_asset["sha256"]
        if pdf_hash not in pdf_destinations:
            pdf_name = f"{safe_name(source_document)}__{pdf_hash[:12]}.pdf"
            destination = original_pdfs / pdf_name
            shutil.copy2(source_pdf, destination)
            pdf_destinations[pdf_hash] = f"original_pdfs/{pdf_name}"
            referenced_files.append((destination, "full_authoritative_source_pdf", "SHARED"))
        shared_pdf_rel = pdf_destinations[pdf_hash]

        case_dir = step1 / "cases" / known_id
        case_dir.mkdir(parents=True)
        source_pages_pdf = case_dir / "SOURCE_PAGES.pdf"
        extract_pages(source_pdf, pages, source_pages_pdf)

        crop_records: list[dict[str, Any]] = []
        for number, asset in enumerate(crop_assets, 1):
            source_crop = Path(asset["resolved_path"])
            if not source_crop.is_file() or sha256(source_crop) != asset["sha256"]:
                raise RuntimeError(f"Crop missing or changed for {known_id}: {source_crop}")
            crop_name = f"HISTORICAL_CROP_{number}{source_crop.suffix.lower()}"
            crop_dest = case_dir / crop_name
            shutil.copy2(source_crop, crop_dest)
            crop_records.append(
                {
                    "path": crop_name,
                    "original_source_file": str(source_crop),
                    "original_physical_pdf_pages": pages,
                    **file_record(crop_dest, "auxiliary_historical_crop", known_id),
                }
            )

        page_map = {
            "known_id": known_id,
            "authoritative_source": shared_pdf_rel,
            "packet_pages": [
                {
                    "packet_page": index,
                    "original_file": source_document,
                    "original_full_pdf_sha256": pdf_hash,
                    "original_physical_pdf_page": page,
                }
                for index, page in enumerate(pages, 1)
            ],
        }
        write_json(case_dir / "PAGE_MAP.json", page_map)

        source_references = [
            {
                "source_document": source_document,
                "source_version": phase2["canonical_packages"][0]["identity"].get("version"),
                "physical_pdf_pages": pages,
                "full_original_pdf": f"../../{shared_pdf_rel}",
                "full_original_pdf_sha256": pdf_hash,
                "extracted_source_pages": "SOURCE_PAGES.pdf",
                "page_map": "PAGE_MAP.json",
            }
        ]
        historical_claim = {
            "known_id": known_id,
            "historical_issue_text": historical["issue_text"],
            "claim_to_test": historical["issue_text"],
            "engineering_subject": historical.get("engineering_subject"),
            "issue_type": historical.get("issue_type", []),
            "context": {
                "old": None,
                "new": None,
                "single_version_source": phase2["canonical_packages"][0]["identity"].get("version"),
                "note": "Frozen metadata identifies one source version; no OLD/NEW split is asserted.",
            },
            "source_references": source_references,
        }
        write_json(case_dir / "HISTORICAL_CLAIM.json", historical_claim)

        evidence_items = [
            {
                "evidence_id": f"HIST-EVID-{index:02d}",
                "source_file": source_document,
                "physical_pdf_pages": pages,
                "evidence_type": sorted(historical.get("evidence_types", [])),
                "modalities": sorted(historical.get("evidence_modalities", [])),
                "extracted_text_or_table": fragment,
                "machine_extracted": True,
                "authoritative_source": "SOURCE_PAGES.pdf and linked full original PDF",
            }
            for index, fragment in enumerate(historical.get("evidence_fragments", []), 1)
        ]
        write_json(
            case_dir / "EVIDENCE.json",
            {
                "known_id": known_id,
                "machine_extracted": True,
                "manual_ocr_corrections": False,
                "pdf_is_authoritative": True,
                "items": evidence_items,
            },
        )

        source_index = {
            "known_id": known_id,
            "source_project": phase2["project"],
            "source_document": source_document,
            "canonical_input_ids": phase2["canonical_bundle_ids"],
            "files": [
                {
                    "path": "SOURCE_PAGES.pdf",
                    "original_source_file": str(source_pdf),
                    "original_physical_pdf_pages": pages,
                    **file_record(source_pages_pdf, "authoritative_relevant_pdf_pages", known_id),
                },
                {
                    "path": f"../../{shared_pdf_rel}",
                    "original_source_file": str(source_pdf),
                    "original_physical_pdf_pages": pages,
                    **file_record(original_pdfs / Path(shared_pdf_rel).name, "full_authoritative_source_pdf", known_id),
                },
                *crop_records,
            ],
        }
        write_json(case_dir / "SOURCE_INDEX.json", source_index)
        write_text(
            case_dir / "CASE_README.txt",
            f"""CASE: {known_id}
PHASE 2 CASE ID (traceability only): {case_id}
SOURCE DOCUMENT: {source_document}
PHYSICAL PDF PAGE(S): {', '.join(map(str, pages))}

Review HISTORICAL_CLAIM.json as an unverified hypothesis. Use SOURCE_PAGES.pdf
and PAGE_MAP.json first. The full authoritative PDF is linked in SOURCE_INDEX.json.
HISTORICAL_CROP_*.png is auxiliary only. EVIDENCE.json is machine-extracted and
must not override the PDF.
""",
        )

        overview_rows.append(
            [
                known_id,
                historical["issue_text"],
                source_document,
                "",
                "",
                ", ".join(historical.get("evidence_modalities", [])),
                f"cases/{known_id}",
                "",
                "",
            ]
        )

        input_dir = Path(phase2["candidate_input"])
        input_files = []
        for input_name in ("input.json", "overview.json", "task.json", "PACKAGE_HASHES.json"):
            input_path = input_dir / input_name
            if not input_path.is_file():
                raise FileNotFoundError(input_path)
            input_files.append(
                {
                    "file": input_name,
                    "original_path": str(input_path),
                    "bytes": input_path.stat().st_size,
                    "sha256": sha256(input_path),
                }
            )
        input_assets = []
        for asset in sorted((input_dir / "assets").iterdir()):
            if asset.is_file():
                input_assets.append(
                    {
                        "file": asset.name,
                        "original_path": str(asset),
                        "bytes": asset.stat().st_size,
                        "sha256": sha256(asset),
                        "kind": asset.suffix.lower().lstrip("."),
                    }
                )

        step2_case = step2 / "cases" / known_id
        step2_case.mkdir(parents=True)
        raw_result = PHASE2 / "codex_raw" / case_id / "result.json"
        if not raw_result.is_file() or read_json(raw_result) != result:
            raise RuntimeError(f"Frozen raw and normalized Phase 2 results differ for {case_id}")
        shutil.copy2(raw_result, step2_case / "CODEX_RAW_RESPONSE.json")
        write_json(step2_case / "CODEX_NORMALIZED_FINDINGS.json", result)

        evidence_refs = []
        for finding in result.get("findings", []):
            for ref in finding.get("evidence_refs", []):
                evidence_refs.append(
                    {
                        "finding_id": finding.get("finding_id"),
                        "source_sha256": ref.get("source_sha256"),
                        "physical_pdf_page": ref.get("page_1based"),
                        "locator": ref.get("locator"),
                        "role": ref.get("role"),
                        "observation": ref.get("observation"),
                    }
                )
        write_json(
            step2_case / "CODEX_EVIDENCE_INDEX.json",
            {
                "known_id": known_id,
                "codex_case_id": case_id,
                "frozen_phase2_source": str(raw_result),
                "evidence_refs": evidence_refs,
            },
        )
        write_json(
            step2_case / "CODEX_INPUT_INDEX.json",
            {
                "known_id": known_id,
                "codex_case_id": case_id,
                "candidate_input_directory": str(input_dir),
                "candidate_input_sha256": phase2["input_sha256"],
                "candidate_package_sha256": phase2["package_sha256"],
                "canonical_input_ids": phase2["canonical_bundle_ids"],
                "input_files": input_files,
                "input_assets": input_assets,
                "step1_case_reference": f"STEP1_SOURCE_FIRST/cases/{known_id}",
            },
        )

        findings = result.get("findings", [])
        finding_ids = [finding.get("finding_id") for finding in findings]
        descriptions = [finding.get("description") or finding.get("defect") for finding in findings]
        subjects = []
        codex_pages = set()
        for finding in findings:
            element = finding.get("engineering_element") or {}
            subject = " / ".join(str(element.get(k)) for k in ("type", "tag") if element.get(k))
            if subject:
                subjects.append(subject)
            for ref in finding.get("evidence_refs", []):
                if ref.get("page_1based") is not None:
                    codex_pages.add(int(ref["page_1based"]))
        codex_rows.append(
            [
                known_id,
                case_id,
                ", ".join(filter(None, finding_ids)),
                "\n".join(filter(None, descriptions)),
                "\n".join(subjects),
                ", ".join(map(str, sorted(codex_pages))),
                ", ".join(codex_prep.get("evidence_modalities", [])),
                "",
                "",
            ]
        )

        critical_manifest_records.append(
            {
                "known_id": known_id,
                "source_project": phase2["project"],
                "source_document": source_document,
                "historical_issue_text": historical["issue_text"],
                "engineering_subject": historical.get("engineering_subject"),
                "canonical_input_ids": phase2["canonical_bundle_ids"],
                "source_files": [
                    f"cases/{known_id}/SOURCE_PAGES.pdf",
                    shared_pdf_rel,
                    *[f"cases/{known_id}/{item['path']}" for item in crop_records],
                ],
                "relevant_old_pages": [],
                "relevant_new_pages": [],
                "relevant_single_version_pages": pages,
                "old_new_status": "NOT_PRESENT_IN_FROZEN_METADATA",
                "evidence_modalities": historical.get("evidence_modalities", []),
                "phase2_codex_case_id": case_id,
            }
        )

    write_json(
        step1 / "CRITICAL8_MANIFEST.json",
        {
            "case_count": 8,
            "source": str(SOURCE_PREP / "CRITICAL8_REVIEW_MANIFEST.json"),
            "records": critical_manifest_records,
            "reviewer_verdicts_prefilled": 0,
        },
    )
    write_workbook(
        step1 / "CRITICAL8_OVERVIEW.xlsx",
        [
            "known_id",
            "historical_claim",
            "source_documents",
            "old_pages",
            "new_pages",
            "evidence_types",
            "case_folder",
            "source_verdict",
            "comment",
        ],
        overview_rows,
        [34, 70, 34, 14, 14, 24, 42, 24, 50],
    )
    write_workbook(
        step2 / "CODEX_CRITICAL8_OVERVIEW.xlsx",
        [
            "known_id",
            "codex_case_id",
            "codex_finding_ids",
            "codex_summary",
            "codex_subject",
            "codex_pages",
            "codex_evidence_types",
            "retention_verdict",
            "review_comment",
        ],
        codex_rows,
        [34, 30, 28, 80, 48, 18, 24, 24, 50],
    )

    readme = OUTPUT / "README_FIRST.txt"
    write_text(
        readme,
        """CRITICAL 8 — CHATGPT REVIEW PACKAGE

1. Не отправляйте весь MASTER ChatGPT сразу.
2. Сначала распакуйте MASTER.
3. В новый чистый ChatGPT chat загрузите только STEP1_SOURCE_FIRST.zip.
4. Попросите выполнить source-first проверку всех 8 cases по инструкции внутри.
5. Сохраните ответ и зафиксируйте все восемь source-first verdicts.
6. Только после этого загрузите STEP2_CODEX_COMPARISON.zip в тот же chat.
7. Попросите сравнить уже зафиксированные выводы с результатами Codex.

Причина разделения: не допустить влияния Codex answer на независимый
source-first review.
""",
    )

    # Ensure reviewer columns are truly blank before archiving.
    step1_book = load_workbook(step1 / "CRITICAL8_OVERVIEW.xlsx", data_only=False)
    step2_book = load_workbook(step2 / "CODEX_CRITICAL8_OVERVIEW.xlsx", data_only=False)
    for row in range(2, 10):
        if step1_book.active.cell(row, 8).value or step1_book.active.cell(row, 9).value:
            raise RuntimeError("STEP1 reviewer field was prefilled")
        if step2_book.active.cell(row, 8).value or step2_book.active.cell(row, 9).value:
            raise RuntimeError("STEP2 reviewer field was prefilled")

    # STEP1 case data must contain neither model answers nor matcher/retention labels.
    forbidden = re.compile(
        r"automatic[_ -]?retention|retention[_ -]?verdict|matcher[_ -]?diagnosis|"
        r"RETAINED_STRONG|RETAINED_PARTIAL|CODEX_MISS|HISTORICAL_NOT_SUPPORTED|"
        r"PACKAGING_GAP|AMBIGUOUS",
        re.IGNORECASE,
    )
    leakage_hits = []
    for path in sorted((step1 / "cases").rglob("*")):
        if path.suffix.lower() in {".json", ".txt"} and path.is_file():
            if forbidden.search(path.read_text(encoding="utf-8")):
                leakage_hits.append(str(path.relative_to(step1)))
    if leakage_hits:
        raise RuntimeError(f"Verdict leakage into STEP1 cases: {leakage_hits}")

    step1_zip = OUTPUT / "STEP1_SOURCE_FIRST.zip"
    step2_zip = OUTPUT / "STEP2_CODEX_COMPARISON.zip"
    zip_tree(step1, step1_zip, "STEP1_SOURCE_FIRST")
    zip_tree(step2, step2_zip, "STEP2_CODEX_COMPARISON")

    package_manifest = {
        "package": "CRITICAL8_CHATGPT_REVIEW_MASTER",
        "revision": "r008_critical8_chatgpt_review_package",
        "created_from_frozen_artifacts_only": True,
        "model_calls": 0,
        "new_engineering_judgments": 0,
        "critical_cases": 8,
        "critical_known_ids": [item["known_id"] for item in critical_manifest_records],
        "child_archives": {
            step1_zip.name: file_record(step1_zip, "source_first_review_archive"),
            step2_zip.name: file_record(step2_zip, "codex_comparison_archive"),
        },
        "step1_files": json_file_inventory(step1),
        "step2_files": json_file_inventory(step2),
        "validation": {
            "critical_cases": "8/8",
            "step1_cases_complete": "8/8",
            "step2_cases_complete": "8/8",
            "source_pdf_evidence": "8/8",
            "codex_phase2_outputs": "8/8",
            "codex_verdict_leakage_into_step1": 0,
            "reviewer_verdicts_prefilled": 0,
            "broken_file_references": 0,
        },
        "checksum_scope_note": (
            "CHECKSUMS_SHA256.txt covers every package file existing before the master ZIP, "
            "except itself (self-hashing is impossible). The master ZIP is excluded by design."
        ),
    }
    manifest_path = OUTPUT / "MANIFEST.json"
    write_json(manifest_path, package_manifest)

    # Validate all ZIP entries and the expected isolated top-level roots.
    for archive_path, expected_root in (
        (step1_zip, "STEP1_SOURCE_FIRST/"),
        (step2_zip, "STEP2_CODEX_COMPARISON/"),
    ):
        with zipfile.ZipFile(archive_path) as archive:
            bad = archive.testzip()
            if bad:
                raise RuntimeError(f"Broken ZIP member: {archive_path.name}:{bad}")
            if any(not name.startswith(expected_root) for name in archive.namelist()):
                raise RuntimeError(f"Archive isolation failure: {archive_path}")

    checksums_path = OUTPUT / "CHECKSUMS_SHA256.txt"
    checksum_lines = []
    for path in sorted(OUTPUT.rglob("*")):
        if not path.is_file() or path.name in {checksums_path.name, "CRITICAL8_CHATGPT_REVIEW_MASTER.zip"}:
            continue
        checksum_lines.append(f"{sha256(path)}  {path.relative_to(OUTPUT)}")
    write_text(checksums_path, "\n".join(checksum_lines))

    master_zip = OUTPUT / "CRITICAL8_CHATGPT_REVIEW_MASTER.zip"
    zip_master(
        master_zip,
        [readme, step1_zip, step2_zip, manifest_path, checksums_path],
    )
    with zipfile.ZipFile(master_zip) as archive:
        if archive.testzip():
            raise RuntimeError("Master ZIP integrity failure")
        expected = {
            "README_FIRST.txt",
            "STEP1_SOURCE_FIRST.zip",
            "STEP2_CODEX_COMPARISON.zip",
            "MANIFEST.json",
            "CHECKSUMS_SHA256.txt",
        }
        if set(archive.namelist()) != expected:
            raise RuntimeError("Master ZIP has unexpected contents")

    return {
        "status": "READY",
        "step1_zip": str(step1_zip),
        "step2_zip": str(step2_zip),
        "master_zip": str(master_zip),
        "step1_bytes": step1_zip.stat().st_size,
        "step2_bytes": step2_zip.stat().st_size,
        "master_bytes": master_zip.stat().st_size,
        "master_sha256": sha256(master_zip),
        "checksummed_files": len(checksum_lines),
    }


if __name__ == "__main__":
    try:
        print(json.dumps(build(), ensure_ascii=False, indent=2))
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
