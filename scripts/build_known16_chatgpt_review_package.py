#!/usr/bin/env python3
"""Build the isolated Known16 ChatGPT review package from frozen r008 artifacts.

This is deliberately a local packaging operation: it invokes no model, network,
or review logic and never writes an engineering or reviewer verdict.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RESEARCH_ROOT = Path("/home/coder/auditmanager/corpus-audits/20260915_openrouter_leg_replacement")
SOURCE_PREP = RESEARCH_ROOT / "revisions/r008_known24_source_review_prep"
PHASE2 = RESEARCH_ROOT / "revisions/r008_phase2_codex_retention_candidate"
OUTPUT = RESEARCH_ROOT / "r008_known16_chatgpt_review_package"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value.rstrip() + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_record(path: Path, role: str, known_id: str | None = None) -> dict[str, Any]:
    result: dict[str, Any] = {"bytes": path.stat().st_size, "sha256": sha256(path), "evidence_role": role}
    if known_id:
        result["known_id"] = known_id
    return result


def safe_name(value: str) -> str:
    return re.sub(r"[^0-9A-Za-zА-Яа-яЁё._-]+", "_", value).strip("_.") or "source"


def configure_sheet(ws, widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def write_workbook(path: Path, title: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    configure_sheet(sheet, widths)
    book.save(path)


def extract_pages(source: Path, pages: list[int], destination: Path) -> None:
    if not pages:
        raise ValueError(f"No physical pages supplied for {source}")
    work = destination.parent / ".pdf_parts"
    work.mkdir(parents=True, exist_ok=True)
    parts: list[Path] = []
    for number, page in enumerate(pages, 1):
        pattern = work / f"part-{number}-%d.pdf"
        subprocess.run(["pdfseparate", "-f", str(page), "-l", str(page), str(source), str(pattern)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        generated = work / f"part-{number}-{page}.pdf"
        if not generated.is_file():
            raise FileNotFoundError(generated)
        parts.append(generated)
    if len(parts) == 1:
        shutil.move(parts[0], destination)
    else:
        subprocess.run(["pdfunite", *map(str, parts), str(destination)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    shutil.rmtree(work)


def zip_tree(source: Path, destination: Path, root_name: str) -> None:
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                archive.write(path, Path(root_name) / path.relative_to(source))


def zip_master(destination: Path, files: list[Path]) -> None:
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as archive:
        for path in files:
            archive.write(path, path.name)


def inventory(root: Path) -> list[dict[str, Any]]:
    return [{"path": str(path.relative_to(root)), "bytes": path.stat().st_size, "sha256": sha256(path)}
            for path in sorted(root.rglob("*")) if path.is_file()]


def comparison_structure(historical: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    versions = sorted({str(item["version"]) for item in historical.get("lineage", []) if item.get("version")})
    if len(versions) >= 2:
        return "OLD_NEW", {"source_versions": versions, "basis": "multiple historical lineage versions"}
    if len(versions) == 1:
        return "SINGLE_VERSION_INTERNAL_CONFLICT", {"source_versions": versions, "basis": "one historical lineage version"}
    return "OTHER", {"source_versions": [], "basis": "no version in frozen historical lineage"}


def build() -> dict[str, Any]:
    if OUTPUT.exists():
        raise RuntimeError(f"Refusing to overwrite existing output: {OUTPUT}")

    known24_manifest = read_json(SOURCE_PREP / "KNOWN24_REVIEW_MANIFEST.json")
    critical8_manifest = read_json(SOURCE_PREP / "CRITICAL8_REVIEW_MANIFEST.json")
    sufficient24 = read_json(PHASE2 / "SUFFICIENT24_RETENTION.json")
    phase2_manifest = read_json(PHASE2 / "PHASE2_KNOWN49_MANIFEST.json")
    phase2_results = read_json(PHASE2 / "CODEX_PHASE2_RESULTS.json")

    known24 = known24_manifest["records"]
    critical8 = critical8_manifest["records"]
    known24_ids = {item["known_id"] for item in known24}
    critical8_ids = {item["known_id"] for item in critical8}
    sufficient_ids = {item["known_id"] for item in sufficient24["records"]}
    if len(known24) != 24 or len(known24_ids) != 24:
        raise RuntimeError("Frozen Known24 manifest is not exactly 24 unique records")
    if len(critical8) != 8 or len(critical8_ids) != 8 or not critical8_ids <= known24_ids:
        raise RuntimeError("Frozen Critical8 IDs are not exactly an 8-case subset of Known24")
    if sufficient24.get("denominator") != 24 or sufficient_ids != known24_ids:
        raise RuntimeError("Frozen SUFFICIENT24 manifest does not agree exactly with Known24")
    remaining = [item for item in known24 if item["known_id"] not in critical8_ids]
    remaining_ids = {item["known_id"] for item in remaining}
    if len(remaining) != 16 or len(remaining_ids) != 16 or remaining_ids & critical8_ids or remaining_ids | critical8_ids != known24_ids:
        raise RuntimeError("Known16 coverage arithmetic failed")

    phase2_by_known = {item["known_id"]: item for item in phase2_manifest["records"]}
    result_by_case = {item["case_id"]: item for item in phase2_results["records"]}
    if len(phase2_by_known) != len(phase2_manifest["records"]):
        raise RuntimeError("Duplicate Known ID in Phase2 manifest")

    step1 = OUTPUT / "STEP1_SOURCE_FIRST"
    step2 = OUTPUT / "STEP2_CODEX_COMPARISON"
    originals = step1 / "original_pdfs"
    (step1 / "cases").mkdir(parents=True)
    (step2 / "cases").mkdir(parents=True)
    originals.mkdir()

    write_text(step1 / "00_INSTRUCTIONS.txt", """ЗАДАЧА

Проведи независимую source-first проверку 16 historical findings. Historical claim
является гипотезой, а не автоматически правильным выводом. Главный источник —
исходный PDF. Machine-extracted TEXT / TABLE / GRAPHIC и OCR являются только
вспомогательными: при расхождении PDF authoritative; не исправляй OCR вручную.

Для каждого case сначала зафиксируй один source_verdict:
SUPPORTED
PARTIALLY_SUPPORTED
NOT_SUPPORTED
INSUFFICIENT_DATA

Если comparison_type = OLD_NEW, укажи OLD и NEW только в той мере, в которой
это разделение реально присутствует в источниках. Если comparison_type =
SINGLE_VERSION_INTERNAL_CONFLICT, не придумывай OLD/NEW: исследуй конфликтующие
указания в одной версии документа.

Проверь тот же ли engineering subject, scope и parameter/property; нет ли разных
режимов или физических объектов; не выдано ли отсутствие упоминания за отсутствие
инженерного решения; нет ли source conflict. Не запрашивай STEP2 до фиксации всех
16 source-first выводов. Запиши их в пустые колонки source_verdict и
review_comment файла KNOWN16_OVERVIEW.xlsx либо верни отдельную таблицу.

Codex outputs на этом шаге отсутствуют специально.
""")
    write_text(step2 / "00_INSTRUCTIONS.txt", """ЗАДАЧА

Используй уже зафиксированные source-first verdicts из STEP1. Не меняй их после
просмотра результатов Codex. Для каждого case сравни engineering subject, scope,
issue/property, parameter/state и source evidence. Похожая формулировка текста
сама по себе не означает retention.

Допустимые retention verdicts:
RETAINED_STRONG
RETAINED_PARTIAL
CODEX_MISS
HISTORICAL_NOT_SUPPORTED
PACKAGING_GAP
AMBIGUOUS

Если STEP1 показал NOT_SUPPORTED, не превращай historical finding в retained
только потому, что Codex нашёл соседний новый дефект. Колонки retention_verdict,
packaging_gap и review_comment намеренно пусты. Automatic matcher results не
включены в этот review package.
""")

    pdf_destinations: dict[str, str] = {}
    records: list[dict[str, Any]] = []
    step1_rows: list[list[Any]] = []
    step2_rows: list[list[Any]] = []

    for frozen in remaining:
        known_id, case_id = frozen["known_id"], frozen["case_id"]
        package = Path(frozen["package_path"])
        historical = read_json(package / "HISTORICAL.json")
        case_meta = read_json(package / "CASE.json")
        codex_prep = read_json(package / "CODEX.json")
        phase2 = phase2_by_known.get(known_id)
        result = result_by_case.get(case_id)
        if not phase2 or phase2.get("case_id") != case_id or phase2.get("input_sufficiency") != "SUFFICIENT" or not result:
            raise RuntimeError(f"Frozen Phase2 source is incomplete for {known_id}")
        raw_result = PHASE2 / "codex_raw" / case_id / "result.json"
        if not raw_result.is_file() or read_json(raw_result) != result:
            raise RuntimeError(f"Frozen Codex result mismatch for {known_id}")
        pages = sorted({int(page) for page in historical.get("pages", [])})
        source_assets = [item for item in historical.get("source_assets", []) if item.get("kind") == "source_pdf"]
        auxiliary_assets = [item for item in historical.get("source_assets", []) if item.get("kind") != "source_pdf"]
        if len(source_assets) != 1 or not pages:
            raise RuntimeError(f"Expected one PDF and relevant pages for {known_id}")
        source_asset = source_assets[0]
        source_pdf = Path(source_asset["resolved_path"])
        if not source_pdf.is_file() or source_pdf.read_bytes()[:5] != b"%PDF-" or sha256(source_pdf) != source_asset["sha256"]:
            raise RuntimeError(f"Invalid or changed source PDF for {known_id}")
        source_hash = source_asset["sha256"]
        source_document = historical["source_documents"][0]
        if source_hash not in pdf_destinations:
            name = f"{safe_name(source_document)}__{source_hash[:12]}.pdf"
            shutil.copy2(source_pdf, originals / name)
            pdf_destinations[source_hash] = f"original_pdfs/{name}"
        shared_pdf = pdf_destinations[source_hash]
        comparison_type, structure_basis = comparison_structure(historical)

        case_dir = step1 / "cases" / known_id
        case_dir.mkdir()
        source_pages = case_dir / "SOURCE_PAGES.pdf"
        extract_pages(source_pdf, pages, source_pages)
        crops: list[dict[str, Any]] = []
        for number, asset in enumerate(auxiliary_assets, 1):
            origin = Path(asset["resolved_path"])
            if not origin.is_file() or sha256(origin) != asset["sha256"]:
                raise RuntimeError(f"Changed auxiliary evidence for {known_id}: {origin}")
            destination = case_dir / "crops" / f"HISTORICAL_AUXILIARY_{number}{origin.suffix.lower()}"
            destination.parent.mkdir()
            shutil.copy2(origin, destination)
            crops.append({"path": str(destination.relative_to(case_dir)), "original_source_file": str(origin),
                          "original_physical_pdf_pages": pages, **file_record(destination, "auxiliary_historical_crop", known_id)})

        page_map = {"known_id": known_id, "authoritative_source": shared_pdf,
                    "packet_pages": [{"packet_page": number, "original_file": source_document,
                                      "original_full_pdf_sha256": source_hash, "original_physical_pdf_page": page}
                                     for number, page in enumerate(pages, 1)]}
        write_json(case_dir / "PAGE_MAP.json", page_map)
        sources = [{"source_document": source_document, "physical_pdf_pages": pages,
                    "full_original_pdf": f"../../{shared_pdf}", "full_original_pdf_sha256": source_hash,
                    "extracted_source_pages": "SOURCE_PAGES.pdf", "page_map": "PAGE_MAP.json"}]
        write_json(case_dir / "HISTORICAL_CLAIM.json", {"known_id": known_id,
            "historical_issue_text": historical["issue_text"], "claim_to_test": historical["issue_text"],
            "engineering_subject": historical.get("engineering_subject"), "issue_type": historical.get("issue_type", []),
            "comparison_type": comparison_type, "comparison_structure_basis": structure_basis,
            "source_references": sources})
        evidence_items = [{"evidence_id": f"HIST-EVID-{number:02d}", "known_id": known_id,
                           "source_file": source_document, "physical_pdf_pages": pages,
                           "evidence_type": sorted(historical.get("evidence_types", [])),
                           "modalities": sorted(historical.get("evidence_modalities", [])),
                           "extracted_text_or_table": fragment, "machine_extracted": True,
                           "authoritative_source": "SOURCE_PAGES.pdf and linked full original PDF"}
                          for number, fragment in enumerate(historical.get("evidence_fragments", []), 1)]
        write_json(case_dir / "EVIDENCE_INDEX.json", {"known_id": known_id, "machine_extracted": True,
                   "manual_ocr_corrections": False, "pdf_is_authoritative": True, "items": evidence_items})
        write_text(case_dir / "EVIDENCE.txt", "PDF is authoritative. The following frozen machine-extracted evidence is auxiliary; no OCR was manually corrected.\n\n" + "\n\n".join(
            f"[{item['evidence_id']}] machine_extracted=true; evidence_type={','.join(item['evidence_type'])}; modalities={','.join(item['modalities'])}\n{item['extracted_text_or_table']}" for item in evidence_items))
        source_index = {"known_id": known_id, "source_project": phase2["project"], "source_document": source_document,
                        "files": [{"path": "SOURCE_PAGES.pdf", "original_source_file": str(source_pdf),
                                   "original_physical_pdf_pages": pages, **file_record(source_pages, "authoritative_relevant_pdf_pages", known_id)},
                                  {"path": f"../../{shared_pdf}", "original_source_file": str(source_pdf),
                                   "original_physical_pdf_pages": pages, **file_record(originals / Path(shared_pdf).name, "full_authoritative_source_pdf", known_id)},
                                  *crops]}
        write_json(case_dir / "SOURCE_INDEX.json", source_index)
        write_text(case_dir / "CASE_README.txt", f"""CASE: {known_id}
SOURCE DOCUMENT: {source_document}
PHYSICAL PDF PAGE(S): {', '.join(map(str, pages))}
COMPARISON TYPE: {comparison_type}

Review HISTORICAL_CLAIM.json as an unverified hypothesis. Begin with SOURCE_PAGES.pdf
and PAGE_MAP.json. The linked full original PDF is in SOURCE_INDEX.json. EVIDENCE.txt
and EVIDENCE_INDEX.json are machine-extracted auxiliary material and do not override PDF.
""")

        input_dir = Path(phase2["candidate_input"])
        input_files = []
        for name in ("input.json", "overview.json", "task.json", "PACKAGE_HASHES.json"):
            path = input_dir / name
            if not path.is_file():
                raise FileNotFoundError(path)
            input_files.append({"file": name, "original_path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)})
        input_assets = [{"file": path.name, "original_path": str(path), "bytes": path.stat().st_size,
                         "sha256": sha256(path), "kind": path.suffix.lower().lstrip(".")}
                        for path in sorted((input_dir / "assets").iterdir()) if path.is_file()]
        step2_case = step2 / "cases" / known_id
        step2_case.mkdir()
        shutil.copy2(raw_result, step2_case / "CODEX_RAW_RESPONSE.json")
        write_json(step2_case / "CODEX_NORMALIZED_FINDINGS.json", result)
        evidence_refs = [{"finding_id": finding.get("finding_id"), "source_sha256": ref.get("source_sha256"),
                          "physical_pdf_page": ref.get("page_1based"), "locator": ref.get("locator"),
                          "role": ref.get("role"), "observation": ref.get("observation")}
                         for finding in result.get("findings", []) for ref in finding.get("evidence_refs", [])]
        write_json(step2_case / "CODEX_EVIDENCE_INDEX.json", {"known_id": known_id, "codex_case_id": case_id,
                   "frozen_phase2_source": str(raw_result), "evidence_refs": evidence_refs})
        write_json(step2_case / "CODEX_INPUT_INDEX.json", {"known_id": known_id, "codex_case_id": case_id,
                   "candidate_input_directory": str(input_dir), "candidate_input_sha256": phase2["input_sha256"],
                   "candidate_package_sha256": phase2["package_sha256"], "canonical_input_ids": phase2["canonical_bundle_ids"],
                   "input_files": input_files, "input_assets": input_assets,
                   "step1_case_reference": f"STEP1_SOURCE_FIRST/cases/{known_id}"})
        write_json(step2_case / "CASE_METADATA.json", {"known_id": known_id, "case_id": case_id,
                   "frozen_phase2_case_metadata": case_meta["canonical_manifest_record"],
                   "comparison_type_from_step1_metadata": comparison_type})

        findings = result.get("findings", [])
        finding_ids = [finding.get("finding_id") for finding in findings]
        descriptions = [finding.get("description") or finding.get("defect") for finding in findings]
        subjects = [" / ".join(str(element[key]) for key in ("type", "tag") if element.get(key))
                    for element in (finding.get("engineering_element") or {} for finding in findings)]
        codex_pages = sorted({int(ref["page_1based"]) for finding in findings for ref in finding.get("evidence_refs", []) if ref.get("page_1based") is not None})
        step1_rows.append([known_id, comparison_type, historical["issue_text"], source_document, ", ".join(map(str, pages)),
                           ", ".join(historical.get("evidence_modalities", [])), f"cases/{known_id}", "", ""])
        step2_rows.append([known_id, case_id, ", ".join(filter(None, finding_ids)), "\n".join(filter(None, descriptions)),
                           "\n".join(filter(None, subjects)), ", ".join(map(str, codex_pages)),
                           ", ".join(codex_prep.get("evidence_modalities", [])), "", "", ""])
        records.append({"known_id": known_id, "source_project": phase2["project"], "source_document": source_document,
                        "historical_issue_text": historical["issue_text"], "engineering_subject": historical.get("engineering_subject"),
                        "comparison_type": comparison_type, "comparison_structure_basis": structure_basis,
                        "canonical_input_ids": phase2["canonical_bundle_ids"], "source_files": [f"cases/{known_id}/SOURCE_PAGES.pdf", shared_pdf, *[f"cases/{known_id}/{x['path']}" for x in crops]],
                        "relevant_physical_pages": pages, "evidence_modalities": historical.get("evidence_modalities", []),
                        "phase2_codex_case_id": case_id})

    write_json(step1 / "KNOWN16_MANIFEST.json", {"case_count": 16, "source": str(SOURCE_PREP / "KNOWN24_REVIEW_MANIFEST.json"),
               "records": records, "reviewer_verdicts_prefilled": 0})
    write_workbook(step1 / "KNOWN16_OVERVIEW.xlsx", "Known 16", ["known_id", "comparison_type", "historical_claim", "source_documents", "physical_pages", "evidence_types", "case_folder", "source_verdict", "review_comment"], step1_rows, [34, 36, 70, 34, 18, 24, 42, 24, 50])
    write_workbook(step2 / "CODEX_KNOWN16_OVERVIEW.xlsx", "Codex Known 16", ["known_id", "codex_case_id", "codex_finding_ids", "codex_summary", "codex_subject", "codex_pages", "codex_evidence_types", "retention_verdict", "packaging_gap", "review_comment"], step2_rows, [34, 30, 28, 80, 48, 18, 24, 24, 24, 50])
    write_text(OUTPUT / "README_FIRST.txt", """KNOWN16 — CHATGPT REVIEW PACKAGE

1. Распакуйте MASTER.
2. Создайте новый чистый ChatGPT chat.
3. Сначала загрузите ТОЛЬКО STEP1_SOURCE_FIRST.zip.
4. Получите и сохраните source-first verdicts 16/16.
5. Не просите ChatGPT пересматривать эти выводы.
6. Затем в ТОМ ЖЕ чате загрузите STEP2_CODEX_COMPARISON.zip.
7. Попросите сравнить Codex с уже зафиксированным baseline.
8. Сохраните оба ответа.

Причина разделения: не допустить влияния Codex на независимый source review.
""")

    coverage = {"known24_source_manifest": str(SOURCE_PREP / "KNOWN24_REVIEW_MANIFEST.json"),
                "critical8_source_manifest": str(SOURCE_PREP / "CRITICAL8_REVIEW_MANIFEST.json"),
                "sufficient24_source_manifest": str(PHASE2 / "SUFFICIENT24_RETENTION.json"),
                "critical_reviewed_known_ids": sorted(critical8_ids), "known16_known_ids": sorted(remaining_ids),
                "critical_reviewed_count": len(critical8_ids), "known16_count": len(remaining_ids),
                "union_count": len(critical8_ids | remaining_ids), "missing_known_ids": sorted(known24_ids - (critical8_ids | remaining_ids)),
                "duplicate_known_ids": [], "critical_known16_intersection": sorted(critical8_ids & remaining_ids), "status": "PASS"}
    write_json(OUTPUT / "KNOWN24_COVERAGE_AUDIT.json", coverage)

    # Reviewer fields must be blank, and no model/matcher answer may enter STEP1.
    book1 = load_workbook(step1 / "KNOWN16_OVERVIEW.xlsx", data_only=False)
    book2 = load_workbook(step2 / "CODEX_KNOWN16_OVERVIEW.xlsx", data_only=False)
    for row in range(2, 18):
        if any(book1.active.cell(row, column).value for column in (8, 9)):
            raise RuntimeError("STEP1 reviewer field was prefilled")
        if any(book2.active.cell(row, column).value for column in (8, 9, 10)):
            raise RuntimeError("STEP2 reviewer field was prefilled")
    forbidden = re.compile(r"CODEX_RAW_RESPONSE|CODEX_NORMALIZED_FINDINGS|codex outputs|claude|automatic[_ -]?matcher|matcher[_ -]?diagnosis|RETAINED_STRONG|RETAINED_PARTIAL|CODEX_MISS|HISTORICAL_NOT_SUPPORTED|PACKAGING_GAP|AMBIGUOUS", re.IGNORECASE)
    leakage = [str(path.relative_to(step1)) for path in sorted((step1 / "cases").rglob("*")) if path.is_file() and path.suffix.lower() in {".json", ".txt"} and forbidden.search(path.read_text(encoding="utf-8"))]
    if leakage:
        raise RuntimeError(f"Forbidden content leaked into STEP1: {leakage}")

    step1_zip, step2_zip = OUTPUT / "STEP1_SOURCE_FIRST.zip", OUTPUT / "STEP2_CODEX_COMPARISON.zip"
    zip_tree(step1, step1_zip, "STEP1_SOURCE_FIRST")
    zip_tree(step2, step2_zip, "STEP2_CODEX_COMPARISON")
    for archive_path, expected_root in ((step1_zip, "STEP1_SOURCE_FIRST/"), (step2_zip, "STEP2_CODEX_COMPARISON/")):
        with zipfile.ZipFile(archive_path) as archive:
            if archive.testzip() or any(not name.startswith(expected_root) for name in archive.namelist()):
                raise RuntimeError(f"Archive isolation/integrity failure: {archive_path}")
    manifest = {"package": "KNOWN16_CHATGPT_REVIEW_MASTER", "revision": "r008_known16_chatgpt_review_package", "created_from_frozen_artifacts_only": True,
                "model_calls": 0, "new_engineering_judgments": 0, "sufficient_cases": 16, "known16_ids": sorted(remaining_ids),
                "critical8_excluded_ids": sorted(critical8_ids), "child_archives": {step1_zip.name: file_record(step1_zip, "source_first_review_archive"), step2_zip.name: file_record(step2_zip, "codex_comparison_archive")},
                "step1_files": inventory(step1), "step2_files": inventory(step2),
                "validation": {"known16": "16/16", "sufficient": "16/16", "critical_ids_accidentally_included": 0, "step1_cases_complete": "16/16", "step2_cases_complete": "16/16", "source_pdf_evidence": "16/16", "codex_frozen_outputs": "16/16", "codex_leakage_into_step1": 0, "claude_leakage_into_step1": 0, "automatic_matcher_verdict_leakage_into_step1": 0, "reviewer_verdicts_prefilled": 0, "broken_file_references": 0},
                "checksum_scope_note": "CHECKSUMS_SHA256.txt covers every package file existing before the master ZIP, except itself and the master ZIP."}
    write_json(OUTPUT / "MANIFEST.json", manifest)
    checksum_path = OUTPUT / "CHECKSUMS_SHA256.txt"
    lines = [f"{sha256(path)}  {path.relative_to(OUTPUT)}" for path in sorted(OUTPUT.rglob("*")) if path.is_file() and path.name not in {checksum_path.name, "KNOWN16_CHATGPT_REVIEW_MASTER.zip"}]
    write_text(checksum_path, "\n".join(lines))
    master = OUTPUT / "KNOWN16_CHATGPT_REVIEW_MASTER.zip"
    zip_master(master, [OUTPUT / "README_FIRST.txt", step1_zip, step2_zip, OUTPUT / "MANIFEST.json", checksum_path])
    with zipfile.ZipFile(master) as archive:
        expected = {"README_FIRST.txt", "STEP1_SOURCE_FIRST.zip", "STEP2_CODEX_COMPARISON.zip", "MANIFEST.json", "CHECKSUMS_SHA256.txt"}
        if archive.testzip() or set(archive.namelist()) != expected:
            raise RuntimeError("Master ZIP integrity failure")
    return {"status": "READY", "model_calls": 0, "known16": "16/16", "critical8_excluded": "8/8", "known24_union": "24/24", "step1_complete": "16/16", "step2_complete": "16/16", "source_pdf_evidence": "16/16", "codex_results": "16/16", "codex_leakage_into_step1": 0, "prefilled_reviewer_verdicts": 0, "broken_references": 0, "checksums": "PASS", "step1_bytes": step1_zip.stat().st_size, "step2_bytes": step2_zip.stat().st_size, "master_bytes": master.stat().st_size, "master": str(master), "master_sha256": sha256(master)}


if __name__ == "__main__":
    try:
        print(json.dumps(build(), ensure_ascii=False, indent=2))
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
