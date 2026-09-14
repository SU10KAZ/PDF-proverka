"""Evidence-linked engineer log and explicit DATA_BLOCKED research decision."""
from collections import Counter
from datetime import datetime,timezone
import html
import json
from pathlib import Path
import shutil
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from experiments.project_change_text_v1.contract import validate
from .run import ROOT,REPO,read,write,sha,now


def collect():
    events=[];origins={};sidecars={}
    baseline=ROOT.parent/'20260913_text_old_scope_recovery_v1/audited/project.json'
    for c in read(baseline)['project_changes']:
        events.append(c);origins[c['project_change_id']]='HISTORICAL_TEXT_DEV'
    for r in read(ROOT/'01_CYCLES/01_state/RESULTS.json'):
        if r['approach']!='typed_event':continue
        for c in r['result']['project_changes']:
            events.append(c);origins[c['project_change_id']]='CURRENT_TABLE_DEV';sidecars[c['project_change_id']]=r['result']
    for r in read(ROOT/'01_CYCLES/02_groups/RESULTS.json'):
        if r['approach']!='closed_interval':continue
        for c in r['result']['project_changes']:
            events.append(c);origins[c['project_change_id']]='CURRENT_GROUP_DEV';sidecars[c['project_change_id']]=r['result']
    for r in read(ROOT/'01_CYCLES/05_equipment_containers/RESULTS.json'):
        for c in r['result']['project_changes']:
            events.append(c);origins[c['project_change_id']]='CURRENT_FORM_REVIEW';sidecars[c['project_change_id']]=r
    quarantine=read(ROOT/'03_VALIDATIONS/recovered_v2/RESULTS.json')['project_changes']
    # Audit sidecar corrects the real source classification without rewriting
    # the frozen TEXT producer's four raw REVIEW predictions.
    write(ROOT/'05_FINAL/SOURCE_QUARANTINE.json',[
        dict(raw_prediction=c,actual_source_type='TABLE',audit='RASTER_CONFIRMED',reason='Parameter table routed as narrative',
             excluded_from_engineering_event_log=True) for c in quarantine])
    return events,origins,sidecars,quarantine


def locators(c,side):
    result=[]
    for e in c['evidence_'+side]:
        page=(e.get('locator') or {}).get('page') or next((r['page'] for r in e['source_refs']),None)
        result.append(f"{e['document_code']} · {e['document_version'][:10]} · PDF {page}")
    return '\n'.join(dict.fromkeys(result))


def build_log(events,origins,sidecars):
    final=ROOT/'05_FINAL';wb=Workbook();ws=wb.active;ws.title='Журнал изменений'
    columns=['Что изменилось','Где / OLD','Где / NEW','Было','Стало','Инженерный объект','Тип изменения','Одно событие — основание','Источники','Статус','Аудит','Подробности','ID']
    ws.append(columns);cards=[]
    types={'EQUIPMENT_REPLACED':'Замена оборудования','EQUIPMENT_ADDED':'Возможное добавление','EQUIPMENT_REMOVED':'Возможное удаление',
           'EQUIPMENT_COUNT_CHANGED':'Изменение количества','SYSTEM_TYPE_CHANGED':'Изменение типа системы',
           'SYSTEM_CONFIGURATION_CHANGED':'Изменение состава системы','CAPACITY_CHANGED':'Изменение расчётной характеристики',
           'REQUIREMENT_CHANGED':'Изменение требования','OTHER_ENGINEERING_CHANGE':'Изменение инженерного свойства'}
    properties={'model':'модель','flow':'расход','pressure':'давление','power':'мощность','capacity':'холодопроизводительность','composition':'состав','count':'количество'}
    for c in sorted(events,key=lambda c:(c['status']!='PROVEN',origins[c['project_change_id']],c['project_change_id'])):
        origin=origins[c['project_change_id']]
        audit=('Исторический DEV-аудит; не перепроверен заново' if origin=='HISTORICAL_TEXT_DEV' else
               'DEV: проверено по OLD/NEW растру' if c['status']=='PROVEN' else 'Требует проверки')
        details='\n'.join(f"{properties.get(f['property'],f['property'])}: {f['old']['value']} {f['old']['unit'] or ''} → {f['new']['value']} {f['new']['unit'] or ''}" for f in c['supporting_fact_changes'])
        why=('Модель и изменившиеся характеристики относятся к замене одного оборудования в том же инженерном месте.' if c['change_type']=='EQUIPMENT_REPLACED' else
             'Сравнивается состав одной группы в установленных границах; её строки сохранены в доказательствах.' if origin=='CURRENT_GROUP_DEV' else
             'Различия относятся к одному инженерному объекту и одному изменению его состояния.')
        where=[]
        for context in sidecars.get(c['project_change_id'],{}).get('scope_evidence',{}).get('new',[]):
            t=context.get('text','')
            for key,label in [('floor:','Этаж:'),('apartment_or_unit:','Помещение/квартира:'),('room:','Помещение:'),('system:','Система:')]:
                if t.startswith(key):where.append(t.replace(key,label,1))
        location='; '.join(dict.fromkeys(where))
        ws.append([c['short_summary_ru'],locators(c,'old'),locators(c,'new'),c['old_state'],c['new_state'],
                   c['engineering_subject']['semantic_subject'],types.get(c['change_type'],'Инженерное изменение'),why,
                   '+'.join(c['routes']),c['status'],audit,details or '\n'.join(c['review_reasons']),c['project_change_id']])
        ev=[]
        for side in ['old','new']:
            for e in c['evidence_'+side]:
                page=(e.get('locator') or {}).get('page') or next((r['page'] for r in e['source_refs']),None)
                href=e['source_receipts']['pdf']['path']+f'#page={page}'
                ev.append(f'<p><b>{side.upper()}, PDF {page}</b> <a href="{html.escape(href,quote=True)}">{html.escape(e["document_code"])}</a><br>{html.escape(e.get("quote") or "См. область источника")}</p>')
        cards.append(f'''<article data-status="{c['status']}"><small>{c['status']} · {' + '.join(c['routes'])} · {html.escape(audit)}</small><h2>{html.escape(c['short_summary_ru'])}</h2><p><b>Где:</b> {html.escape(location)} {html.escape(locators(c,'new'))}</p><div class="states"><p><b>Было</b><br>{html.escape(c['old_state'] or 'Не установлено')}</p><p><b>Стало</b><br>{html.escape(c['new_state'] or 'Не установлено')}</p></div><p><b>Почему одно событие:</b> {html.escape(why)}</p><details><summary>Подробности и источники ({len(c['supporting_fact_changes'])} различий)</summary><pre>{html.escape(details)}</pre>{''.join(ev)}<p>{html.escape('; '.join(c['review_reasons']))}</p><code>{c['project_change_id']}</code></details></article>''')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for row in ws:
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23445B')
    for col,width in {'A':65,'B':38,'C':38,'D':45,'E':45,'F':34,'G':30,'H':55,'I':14,'J':14,'K':38,'L':70,'M':32}.items():ws.column_dimensions[col].width=width
    notes=wb.create_sheet('Ограничения');notes.append(['Статус','DATA_BLOCKED']);notes.append(['Независимая точность','Не установлена'])
    notes.append(['Принято','10 исторических TEXT + 6 TABLE текущего DEV']);notes.append(['Утечка источников','4 новых TEXT REVIEW являются TABLE; вынесены в SOURCE_QUARANTINE.json'])
    notes.append(['Обобщение','Нет новых независимых парных проектов; результаты внутри ALIA не подтверждают переносимость'])
    wb.save(final/'ENGINEER_CHANGE_LOG.xlsx')
    (final/'ENGINEER_CHANGE_LOG.html').write_text('''<!doctype html><html lang="ru"><meta charset="utf-8"><title>Изменения проекта — исследовательский журнал</title><style>body{font:16px/1.55 system-ui;margin:40px auto;max-width:1100px;padding:0 20px;background:#f4f6f8;color:#203343}article{background:white;padding:24px;margin:18px 0;border-radius:8px;border-left:5px solid #387d63}article[data-status=REVIEW]{border-left-color:#c2993a}h1{font-size:30px}h2{font-size:20px}small{color:#536473}.states{display:grid;grid-template-columns:1fr 1fr;gap:24px}pre{white-space:pre-wrap}summary{cursor:pointer;font-weight:600}button{padding:9px 16px;margin-right:10px}</style><h1>Что изменилось в проекте</h1><p><b>DATA_BLOCKED.</b> 16 принятых DEV-событий: 10 из исторического TEXT-аудита и 6 TABLE-событий, проверенных в этой программе. Независимая точность не установлена. Все атомарные различия раскрываются внутри события.</p><p>72 REVIEW-события требуют проверки; ещё 4 ошибочно маршрутизированных TEXT REVIEW вынесены в карантин источников.</p><button onclick="filter('PROVEN')">Принятые</button><button onclick="filter('REVIEW')">Проверить</button><button onclick="filter('ALL')">Все</button>'''+''.join(cards)+'''<script>function filter(s){document.querySelectorAll('article').forEach(a=>a.hidden=s!=='ALL'&&a.dataset.status!==s)}filter('PROVEN')</script></html>''')


def diagnostic_packet():
    out=ROOT/'05_FINAL/SMALL_SOURCE_AUDIT';out.mkdir(exist_ok=True)
    specs=[('01_state','case_023_NEW','Наружный блок №3.6'),('02_groups','case_011_NEW','Состав внутренних блоков при №29.5'),
           ('02_groups','case_035_NEW','Состав внутренних блоков при №9.4'),('01_state','case_022_NEW','Наружный блок №30.5, 31.5'),
           ('05_equipment_containers',None,'Вентилятор системы ДУ4.1.1')]
    cases=[]
    for i,(cycle,prefix,focus) in enumerate(specs,1):
        case=dict(case_id=f'case_{i:02d}',focus=focus,sources={})
        for side in ['old','new']:
            source=ROOT/'01_CYCLES'/cycle/'audit'/((prefix+'_' if prefix else '')+side+'.png')
            dest=out/f'case_{i:02d}_{side}.png';shutil.copyfile(source,dest);case['sources'][side]=dest.name
        cases.append(case)
    write(out/'CASES.json',dict(cases=cases,label_kind='TARGETED_DEV_SOURCE_AUDIT_NOT_INDEPENDENT_BENCHMARK',candidate_answers_included=False))
    (out/'README.md').write_text('''# Пять локальных вопросов по исходникам

Для указанного объекта проверьте OLD/NEW PNG: это тот же объект, есть ли реальное
изменение и какое? ДА / НЕТ / НЕЯСНО / ПОВРЕЖДЁН. Сверьте модель, количество,
единицы и состав группы. Запишите ответ своими словами, с OLD/NEW свидетелями.
Предсказания и уверенность системы не включены. Это целевой DEV-аудит, выбранный
после исследования; он не заменяет независимый benchmark и не измеряет recall.
''')


def main():
    final=ROOT/'05_FINAL';final.mkdir(exist_ok=True)
    events,origins,sidecars,quarantine=collect()
    ids=[c['project_change_id'] for c in events]
    if len(ids)!=len(set(ids)):raise ValueError('Duplicate event IDs')
    receipts={}
    for c in events:
        validate(c,text_only=c['routes']==['TEXT'])
        for side in ['old','new']:
            for e in c['evidence_'+side]:
                for r in e['source_receipts'].values():receipts[r['path']]=r['sha256']
                for r in (e.get('locator') or {}).get('artifact_receipts',{}).values():receipts[r['path']]=r['sha256']
    bad=[p for p,h in receipts.items() if sha(p)!=h]
    if bad:raise ValueError('Source drift: '+str(bad))
    write(final/'SOURCE_TRACEABILITY_AUDIT.json',dict(events=len(events),schema_validated=len(events),source_receipts_verified=len(receipts),source_drift=bad,
                implementing_agent_current_accepted_raster_audit=6,historical_text_audit_not_repeated=10,
                semantic_recall_not_established=True))
    write(final/'PROJECT_CHANGES.json',events);write(final/'ORIGINS.json',origins);write(final/'SCOPE_IDENTITY_DETAILS.json',sidecars)
    build_log(events,origins,sidecars);diagnostic_packet()
    state=read(ROOT/'00_STATE/MASTER_STATE.json')
    current=[c for c in events if origins[c['project_change_id']]!='HISTORICAL_TEXT_DEV']
    summary=dict(final_status='DATA_BLOCKED',major_research_cycles=5,architectures_rejected=5,
        architectures_retained=['Dimension-aware state and event ownership (DEV)','Closed group composition (DEV)',
            'Lossless source-envelope adaptation','Orthogonal source/owner outcomes (diagnostic)',
            'Explicit equipment forms (narrow REVIEW producer)'],
        projects_used_for_DEV=2,projects_registered_DEV_known=3,independent_projects=0,new_source_test_document_pairs=14,
        TEXT=dict(raw_candidate=80,accepted=10,correct_historical_audit=10,correct_current_audit=None,false_historical_audit=0,review=70,quarantined=4),
        TABLE=dict(candidate=12,accepted=6,correct_current_audit=6,false_current_audit=0,review=6),
        GRAPHIC='NOT_YET_VALIDATED',
        ProjectChanges=dict(raw_candidates=92,engineer_log_events=88,accepted=16,
            correct_current_audit=6,correct_historical_audit=10,false_observed_accepted=0,
            review_engineer_events=72,quarantined_source_reviews=4,raw_review_total=76),
        ProjectChange_independent_precision=None,DEV_current_observed_precision='6/6; implementing-agent source audit, not independent',
        high_value_coverage=dict(full_corpus_recall=None,known_local_inventory_events=6,known_local_found_accepted=3,
            known_local_found_review=1,explanation='Historical ALIA inventory: focal replacement + four neighboring group observations; plus new DU fan replacement. Narrow DEV diagnostics, not exhaustive truth.'),
        false_additions_accepted=0,false_removals_accepted=0,current_audited_duplicates=0,
        current_audited_over_grouping=0,current_audited_under_grouping_emitted=0,current_audited_parameter_explosion=0,
        full_log_cross_document_duplicates=None,cross_source_fused_events=0,cross_source_conflicts=None,
        AI=dict(calls=0,input_tokens=0,output_tokens=0,cost_usd=0,median_context=None,p95_context=None,
                policy='No new provider calls. Unanswered LOCAL AI clarification: on-server-only assumed. Historical model receipts replayed as DEV data, not independent adjudication.'),
        runtime_seconds=(datetime.now(timezone.utc)-datetime.fromisoformat(state['started_at'])).total_seconds(),
        production_changed=False,push=0,deploy=0,tests_passed=81,
        next_action='Obtain complete OLD/NEW PDFs and same-version OCR/block exports from at least three previously uninspected real projects, including a project outside Sobytie; lock source manifest before prediction. Broaden equipment-container parsing only in a new DEV cycle.')
    if sum(c['status']=='PROVEN' for c in events)!=summary['ProjectChanges']['accepted'] or len(events)!=88:
        raise ValueError('Summary count mismatch')
    perf=[]
    for p in ROOT.rglob('*PERFORMANCE.json'):
        if p.is_file():perf.append(read(p))
    summary['peak_rss_kib']=max((p.get('peak_rss_kib',0) for p in perf),default=0)
    summary['artifact_bytes_at_closeout']=sum(p.stat().st_size for p in ROOT.rglob('*') if p.is_file())
    write(final/'SUMMARY.json',summary)
    state.update(status='DATA_BLOCKED',current_blocker='NO_PREVIOUSLY_UNINSPECTED_PAIRED_PROJECT_CORPUS',
                 current_metrics=summary,next_action=summary['next_action'],finished_at=now())
    state['accepted_components']=summary['architectures_retained']
    state['rejected_components']=['Literal-header state semantics','Unconditional row event grouping','Fixed-neighbor completeness',
                                  'Geometry-only input recovery without exporter envelope support','Coupled scope/purity gate']
    state['available_unseen_data']=dict(Mosfilmovskaya='100 version directories; zero actual PDF files',
        Sobytie_6_2='77 v001 sources; no second version',Sobytie_6_1='12 v001 sources; no second version',
        paired_unseen_projects=0,registered_scope_only_holdout='Preserved, no semantic inspection or predictions')
    write(ROOT/'00_STATE/MASTER_STATE.json',state)
    print(summary)


if __name__=='__main__':main()
