"""Post-freeze Sol/high comparison against frozen source, Astra, and Terra."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys


ROOT = Path("/home/coder/auditmanager/corpus-audits/20260914_project_change_272")
OUT = ROOT / "cheap_model_screening_v2_sol_high"
V1 = ROOT / "cheap_model_screening_v1"
PAIR_A = ROOT / "pair_a_ai_mapping_change_miner_v1"
PAIR_B = ROOT / "pair_b_ai_mapping_change_miner_v1"
DEPS = ROOT / "controlled_inference_f1_f4_f2_v3/runtime_deps"
sys.path.insert(0, str(DEPS))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def now():
    return datetime.now(timezone.utc).isoformat()


def read(path):
    return json.loads(Path(path).read_text())


def write_new(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify_sol_freeze():
    integrity = read(OUT / "SAMPLE_INTEGRITY.json")
    freeze = read(OUT / "SOL_HIGH_RESULT_FREEZE.json")
    if integrity["status"] != "PASS" or not integrity["same_group_ids"] or not integrity["same_input_hashes"]:
        raise RuntimeError("Sample integrity is not PASS")
    if freeze["model_calls"] != 8 or freeze["successful_calls"] != 8 or freeze["no_truth_leakage"] != "PASS":
        raise RuntimeError("Sol result freeze is incomplete")
    for field in ["source_truth_opened_before_freeze", "astra_outputs_opened_before_freeze", "terra_outputs_opened_before_freeze"]:
        if freeze[field]:
            raise RuntimeError(f"Truth leakage invariant failed: {field}")
    for relative, digest in freeze["hashes"].items():
        if sha(OUT / relative) != digest:
            raise RuntimeError(f"Sol freeze drift: {relative}")
    return integrity, freeze


SOL_ASSESSMENTS = {
    "A-REAL-08": ("CORRECT", ["G008_CHG_003"],
                   "Sol прямо восстановила существенное перераспределение площадей холла и кухни-столовой квартиры 1.3.47 внутри укрупнённой карточки."),
    "A-REAL-10": ("CORRECT", ["G008_CHG_005"],
                   "Sol прямо восстановила существенное перераспределение холлов и кухни-столовой квартиры 2.2.18 внутри укрупнённой карточки."),
    "F03": ("MISSED", [],
            "Sol не зафиксировала переход от двух раздельных вытяжек к общей шахте/показанному вентилятору; соответствующая GRAPHIC-суть отсутствует даже в hints."),
    "F04": ("PARTIAL", ["G017-H01"],
            "Четыре старые установки по 12000 и новый подбор 23600 распознаны, но оставлены unresolved hint и не выданы как доказанный ProjectChange."),
    "F09": ("CORRECT", ["G017-C02"],
            "Обе системы рампы, направление и точные рабочие точки 50990→21600 и 24000→10900 сохранены."),
    "F10": ("CORRECT", ["G017-C01"],
            "Переход от рекомендаций к СП 7.13130.2013 к АВОК 5.5.1-2023 и связанные расчётные параметры сохранены."),
    "G025-C01": ("CORRECT", ["G025-C01"],
                  "Переподбор ППП, числа и изменение размещения сохранены с TABLE/TEXT/GRAPHIC evidence."),
    "G025-C02": ("CORRECT", ["G025-C02"],
                  "Переподбор ПО, числа и изменение размещения сохранены с TABLE/TEXT/GRAPHIC evidence."),
    "G025-C03": ("CORRECT", ["G025-C01", "G025-C02"],
                  "Перенос обеих установок с −2 этажа на кровлю корректно включён в две предметные карточки без отдельной атомизации."),
    "G025-C04": ("CORRECT", ["G025-C03"],
                  "Sol сохранила температуры, геометрию дверей и переход от сопротивлений к новой модели щелей/газопроницания; итоговые расходы находятся в C01/C02."),
    "G025-C05": ("PARTIAL", ["G025-C03"],
                  "Изменение высотных исходных параметров сохранено, но расширение поэтажного расчёта до 16-го этажа не названо."),
}


FALSE_ROWS = [
    {"pair": "A", "group_id": "G002", "sol_id": "G002-C01",
     "subject": "Проектная организация и ответственные лица комплекта АР1",
     "rationale": "Смена организации, фамилий и ролей на титульном листе — административные метаданные, а не конкретное инженерное состояние проекта."},
    {"pair": "A", "group_id": "G002", "sol_id": "G002-C02",
     "subject": "Суффикс КОРР. в обозначении комплекта",
     "rationale": "Изменение идентификатора/оформления документа прямо исключено Change Miner contract; Astra корректно вернула пустой результат."},
]


GROUP_ROWS = [
    {"pair": "A", "group_id": "G002", "quality": "MAJOR_LOSS", "astra_changes": 0, "terra_changes": 3, "sol_changes": 2,
     "note": "Sol сократила число ложных карточек Terra с 3 до 2, но сохранила ту же титульную false-positive ошибку."},
    {"pair": "A", "group_id": "G008", "quality": "SAME_OR_BETTER", "astra_changes": 20, "terra_changes": 5, "sol_changes": 6,
     "note": "Обе REAL15-сущности сохранены полностью при полезном укрупнении 20→6 без новых false."},
    {"pair": "A", "group_id": "G024", "quality": "SAME_OR_BETTER", "astra_changes": 0, "terra_changes": 0, "sol_changes": 0,
     "note": "Пустой результат совпадает с Astra и Terra."},
    {"pair": "A", "group_id": "G031", "quality": "SAME_OR_BETTER", "astra_changes": 0, "terra_changes": 0, "sol_changes": 0,
     "note": "Пустой результат совпадает с Astra и Terra."},
    {"pair": "B", "group_id": "G004", "quality": "SAME_OR_BETTER", "astra_changes": 0, "terra_changes": 0, "sol_changes": 0,
     "note": "Пустой результат совпадает с Astra и Terra."},
    {"pair": "B", "group_id": "G017", "quality": "MAJOR_LOSS", "astra_changes": 4, "terra_changes": 3, "sol_changes": 2,
     "note": "F09/F10 сохранены, F04 понижен до hint, а GRAPHIC-суть F03 об общей шахте полностью пропущена."},
    {"pair": "B", "group_id": "G025", "quality": "ACCEPTABLE_LOSS", "astra_changes": 5, "terra_changes": 3, "sol_changes": 3,
     "note": "Оборудование, размещение и cross-modal evidence сильны; высотная модель сохранена без полного 16-этажного охвата."},
    {"pair": "B", "group_id": "G035", "quality": "SAME_OR_BETTER", "astra_changes": 2, "terra_changes": 1, "sol_changes": 1,
     "note": "Основное событие укрупнено; frozen source-first claims остаются insufficient to judge, подтверждённой reference-потери нет."},
]


def build_source_comparison():
    v1_source = read(V1 / "SOURCE_FIRST_COMPARISON.json")
    rows = []
    for reference in v1_source["rows"]:
        outcome, ids, rationale = SOL_ASSESSMENTS[reference["reference_id"]]
        rows.append({**reference, "sol_outcome": outcome, "sol_ids": ids, "sol_rationale": rationale})
    counts = Counter(row["sol_outcome"] for row in rows)
    value = {
        "created_at": now(), "evaluation_opened_after_verified_sol_freeze": True,
        "reference_definition": v1_source["reference_definition"], "reference_changes": len(rows),
        "rows": rows, "sol_reference_counts": dict(counts), "sol_false_changes": FALSE_ROWS,
        "sol_false_count": len(FALSE_ROWS),
        "sol_output_counts": {"CORRECT": 7, "PARTIAL": 4, "FALSE": 2, "INSUFFICIENT_TO_JUDGE": 1, "TOTAL": 14},
        "output_count_note": "Reference-event scoring is primary. Output-card counts differ because one Sol card can retain multiple reference events and G035 remains insufficient to judge.",
    }
    write_new(OUT / "SOURCE_FIRST_COMPARISON.json", value)
    return value


def build_three_model_comparison(source):
    quality = Counter(row["quality"] for row in GROUP_ROWS)
    model_rows = [
        {"model": "gpt-6-astra/xhigh", "correct": 10, "partial": 1, "false": 0, "missed": 0,
         "reference_retained": "11/11", "cross_modal_quality": "STRONG", "graphic_quality": "STRONG",
         "numeric_table_quality": "STRONG", "input_plus_output_tokens": 1146959},
        {"model": "gpt-5.6-terra/high", "correct": 6, "partial": 3, "false": 3, "missed": 2,
         "reference_retained": "9/11", "cross_modal_quality": "MIXED", "graphic_quality": "MAJOR_GAP_G017",
         "numeric_table_quality": "GOOD_WITH_MISSES", "input_plus_output_tokens": 694471},
        {"model": "gpt-5.6-sol/high", "correct": 8, "partial": 2, "false": 2, "missed": 1,
         "reference_retained": "10/11", "cross_modal_quality": "MIXED", "graphic_quality": "MAJOR_GAP_G017",
         "numeric_table_quality": "STRONG_WITH_GAPS", "input_plus_output_tokens": 699592},
    ]
    value = {
        "created_at": now(), "same_frozen_sample": True, "reference_changes": 11,
        "models": model_rows, "group_rows": GROUP_ROWS, "sol_vs_astra_quality_counts": dict(quality),
        "conclusion": "Sol is materially better than Terra on reference retention and Pair A grouping, but both cheaper candidates repeat the G002 false-positive class and miss the G017 graphic event.",
    }
    write_new(OUT / "ASTRA_TERRA_SOL_COMPARISON.json", value)
    return value


def usage_from_receipts(paths):
    rows = []
    for path in paths:
        receipt = read(path)
        usage = receipt["usage"][0]
        rows.append({"call": receipt.get("sample_id", path.parent.name),
                     **{key: usage.get(key, 0) for key in ["input_tokens", "cached_input_tokens", "output_tokens", "reasoning_output_tokens"]}})
    total = {key: sum(row[key] for row in rows)
             for key in ["input_tokens", "cached_input_tokens", "output_tokens", "reasoning_output_tokens"]}
    total["input_plus_output_tokens"] = total["input_tokens"] + total["output_tokens"]
    total["average_input_plus_output_per_group"] = total["input_plus_output_tokens"] / len(rows)
    return rows, total


def build_usage():
    rows, sol = usage_from_receipts(sorted((OUT / "sol_raw").glob("*/RECEIPT.json")))
    v1_usage = read(V1 / "TOKEN_USAGE.json")
    astra = v1_usage["astra_historical_same_sample"]
    terra = v1_usage["candidate"]
    saving = (1 - sol["input_plus_output_tokens"] / astra["input_plus_output_tokens"]) * 100
    value = {
        "created_at": now(), "model": "gpt-5.6-sol", "reasoning": "high", "calls": 8,
        "sol": sol, "sol_by_call": rows, "astra_historical_same_sample": astra,
        "terra_frozen_same_sample": terra, "saving_vs_astra_percent": round(saving, 1),
        "sol_difference_vs_terra_tokens": sol["input_plus_output_tokens"] - terra["input_plus_output_tokens"],
        "sol_difference_vs_terra_percent": round((sol["input_plus_output_tokens"] / terra["input_plus_output_tokens"] - 1) * 100, 1),
        "cached_input_note": "cached_input_tokens is included within input_tokens and is not added again.",
        "comparison_caveat": "Tokenization and cache behavior differ by model; subscription usage is primary and API price is not used as the main metric.",
    }
    write_new(OUT / "TOKEN_USAGE.json", value)
    return value


def build_projection(usage):
    full_a = read(PAIR_A / "TOKEN_USAGE.json")["total"]
    full_b = read(PAIR_B / "MODEL_USAGE.json")["total_tokens"]
    sol_a = sum(row["input_tokens"] + row["output_tokens"] for row in usage["sol_by_call"] if row["call"].startswith("PAIR_A_"))
    sol_b = sum(row["input_tokens"] + row["output_tokens"] for row in usage["sol_by_call"] if row["call"].startswith("PAIR_B_"))
    astra_rows = read(V1 / "TOKEN_USAGE.json")["astra_by_call"]
    astra_a = sum(row["input_tokens"] + row["output_tokens"] for row in astra_rows[:4])
    astra_b = sum(row["input_tokens"] + row["output_tokens"] for row in astra_rows[4:])
    historical_a = full_a["input_tokens"] + full_a["output_tokens"]
    historical_b = full_b["input_tokens"] + full_b["output_tokens"]
    ratios = {"A": sol_a / astra_a, "B": sol_b / astra_b}
    stratified = round(historical_a * ratios["A"] + historical_b * ratios["B"])
    pooled = round((historical_a + historical_b) * usage["sol"]["input_plus_output_tokens"] /
                   usage["astra_historical_same_sample"]["input_plus_output_tokens"])
    group_scaled = round(usage["sol"]["average_input_plus_output_per_group"] * 67)
    value = {
        "created_at": now(), "projection_only_not_fact": True,
        "scope": "Semantic Mapping V2 + ProjectChange Miner V3 for Pair A + Pair B",
        "historical_full_input_plus_output_tokens": historical_a + historical_b,
        "pair_specific_sample_ratios": ratios, "projected_full_v3_input_plus_output_tokens": stratified,
        "rounded_report_value": "~5.0 million tokens",
        "sensitivity": {"pooled_ratio_projection": pooled, "67_group_average_projection_mining_only": group_scaled,
                        "indicative_range": [min(pooled, stratified), group_scaled]},
        "assumptions": ["Full V3 has approximately the historical Pair A/Pair B mapping+mining volume.",
                        "Sample-specific output compression transfers to the full run.",
                        "No repair, retry, verifier, dedupe, or evaluation calls are included."],
        "warning": "Projection, not fact. The eight groups are structurally controlled rather than statistically random, and model caching/tokenization differ.",
    }
    write_new(OUT / "FULL_V3_COST_PROJECTION.json", value)
    return value


def add_sheet(book, title, rows):
    sheet = book.create_sheet(title)
    if not rows:
        return
    columns = list(rows[0])
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([json.dumps(row[column], ensure_ascii=False) if isinstance(row[column], (list, dict)) else row[column]
                      for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(80, max(12, max(len(str(cell.value or "")) for cell in cells) + 2))


def build_excel(source, comparison, usage, projection):
    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "Summary", [{
        "status": "COMPLETED_SOL_REJECTED", "model": "gpt-5.6-sol/high", "sample": "8 identical frozen groups",
        "calls": 8, "reference_changes": 11, "correct": 8, "partial": 2, "missed": 1, "false": 2,
        "same_or_better": 5, "acceptable_loss": 1, "major_loss": 2,
        "input": usage["sol"]["input_tokens"], "cached": usage["sol"]["cached_input_tokens"],
        "output": usage["sol"]["output_tokens"], "reasoning": usage["sol"]["reasoning_output_tokens"],
        "input_plus_output": usage["sol"]["input_plus_output_tokens"], "saving_vs_astra_percent": usage["saving_vs_astra_percent"],
        "projected_full_v3": projection["projected_full_v3_input_plus_output_tokens"],
        "no_truth_leakage": "PASS", "recommendation": "KEEP_ASTRA_XHIGH",
    }])
    add_sheet(book, "Reference comparison", source["rows"])
    add_sheet(book, "False changes", source["sol_false_changes"])
    add_sheet(book, "Three models", comparison["models"])
    add_sheet(book, "Group quality", comparison["group_rows"])
    add_sheet(book, "Sol usage", usage["sol_by_call"])
    path = OUT / "RESULTS.xlsx"
    if path.exists():
        raise FileExistsError(path)
    book.save(path)


def build_report(usage, projection):
    sol = usage["sol"]
    text = f"""# ProjectChange — cheap model screening V2 · Sol/high

STATUS: COMPLETED_SOL_REJECTED

MODEL: gpt-5.6-sol / high

SAMPLE: 8 identical frozen groups

MODEL CALLS: 8

REFERENCE CHANGES: 11

SOL:
- correct: 8
- partial: 2
- missed: 1
- false: 2
- reference retained: 10/11 (8 full + 2 partial)

VS ASTRA:
- same/better: 5
- acceptable loss: 1
- major loss: 2

TERRA:
- correct: 6
- partial: 3
- missed: 2
- false: 3
- reference retained: 9/11

ASTRA XHIGH:
- correct: 10
- partial: 1
- missed: 0
- false: 0
- reference retained: 11/11

TOKEN USAGE SOL:
- input: {sol['input_tokens']}
- cached: {sol['cached_input_tokens']} (included in input)
- output: {sol['output_tokens']}
- reasoning: {sol['reasoning_output_tokens']}
- input+output: {sol['input_plus_output_tokens']}

ASTRA SAME SAMPLE: 1,146,959 input+output tokens

TERRA SAME SAMPLE: 694,471 input+output tokens

SOL SAVING VS ASTRA: {usage['saving_vs_astra_percent']}%

PROJECTED FULL V3: ~5.0M input+output tokens ({projection['projected_full_v3_input_plus_output_tokens']} stratified estimate; indicative range {projection['sensitivity']['indicative_range'][0]}–{projection['sensitivity']['indicative_range'][1]}). Projection, not fact.

NO TRUTH LEAKAGE: PASS

QUALITY CONCLUSION: Sol/high заметно лучше Terra/high: полностью сохранила обе REAL15-сущности G008, сильнее прочитала числовые таблицы и восстановила почти всю расчётную модель G025. Но она повторила титульный false-positive класс G002 (2 ложных изменения), пропустила доказанную GRAPHIC-суть общей шахты F03 и оставила F04 лишь unresolved hint. Условие допуска к большому V3 не выполнено из-за false changes, real miss и двух major-loss групп.

RECOMMENDATION: KEEP_ASTRA_XHIGH

Sol/high не использовать для полного V3. Полный прогон автоматически не запускался.

PRODUCTION: UNCHANGED

VALIDATION: NOT OPENED

FINAL HOLDOUT: NOT OPENED

## Trace notes

- Sol received the exact 112 frozen input files from V1; all hashes matched.
- Sol result was hash-frozen before source truth, Astra outputs, Terra outputs, or prior verdicts were opened.
- Astra calls: 0; Terra calls: 0; OpenRouter calls: 0; Claude calls: 0; retries: 0.
- Official model capability source: https://developers.openai.com/api/docs/models/gpt-5.6-sol
"""
    with (OUT / "FINAL_REPORT.md").open("x") as handle:
        handle.write(text)


def main():
    verify_sol_freeze()
    source = build_source_comparison()
    comparison = build_three_model_comparison(source)
    usage = build_usage()
    projection = build_projection(usage)
    build_excel(source, comparison, usage, projection)
    build_report(usage, projection)
    required = ["SAMPLE_INTEGRITY.json", "SOL_HIGH_RESULT_FREEZE.json", "SOURCE_FIRST_COMPARISON.json",
                "ASTRA_TERRA_SOL_COMPARISON.json", "TOKEN_USAGE.json", "FULL_V3_COST_PROJECTION.json",
                "RESULTS.xlsx", "FINAL_REPORT.md"]
    write_new(OUT / "DELIVERY_MANIFEST.json", {
        "created_at": now(), "status": "COMPLETED_SOL_REJECTED",
        "files": {name: sha(OUT / name) for name in required},
        "sol_calls": 8, "astra_calls": 0, "terra_calls": 0, "openrouter_calls": 0, "claude_calls": 0,
        "no_truth_leakage": "PASS", "production": "UNCHANGED",
        "validation": "NOT OPENED", "final_holdout": "NOT OPENED",
    })
    print(json.dumps({"status": "COMPLETED_SOL_REJECTED", "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
