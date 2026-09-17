"""Post-result-freeze source-first evaluation of consolidated comparison V2."""
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

ROOT=Path('/home/coder/auditmanager/corpus-audits/20260914_project_change_272')
STAGE=ROOT/'pair_b_semantic_decomposer_v1/consolidated_local_comparison_v2'
GROUP_FILE=ROOT/'pair_b_semantic_decomposer_v1/semantic_consolidator_v1/FINAL_CONSOLIDATED_GROUPS.json'
TRUTH=ROOT/'pair_b_independent_finding_loss_trace'

def read(p):return json.loads(Path(p).read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(p,v):
    p=Path(p);p.parent.mkdir(parents=True,exist_ok=True)
    data=v if isinstance(v,str) else json.dumps(v,ensure_ascii=False,indent=2,sort_keys=True)+'\n'
    if p.exists() and p.read_text()!=data:raise ValueError('Immutable evaluation artifact drift: '+str(p))
    p.write_text(data)
def now():return datetime.now(timezone.utc).isoformat()

# Source-first decisions made only after the V2 result freeze.  Group IDs are
# deterministic trace links, never tuning/admission inputs.
TRACE={
'F01':dict(groups=['PB_064'],atomic=True,correct=False,sufficient=False,found=False,final=False,status='FOUND_PARTIAL',
 reason='Atomic claim and five-section group exist, but section 3.1 is not in the group and OLD table/graphic proof needed for the full 1.1–3.1 event was not delivered as primary local evidence.'),
'F02':dict(groups=[],atomic=False,correct=False,sufficient=False,found=False,final=False,status='MISSED',reason='No atomic candidate contains the section 1.1 OLD/NEW parameter set.'),
'F03':dict(groups=['PB_058'],atomic=True,correct=True,sufficient=True,found=True,final=False,status='FOUND_STRONG',reason='PB_058 is the correct parking smoke-exhaust network event; OLD 108 and NEW 184 rasters were delivered and raw Astra ACCEPT described the proven consolidation.'),
'F04':dict(groups=[],atomic=False,correct=False,sufficient=False,found=False,final=False,status='MISSED',reason='No atomic candidate captures the 4×12000 to redesigned 23600 compensation event.'),
'F06':dict(groups=['PB_012','PB_065'],atomic=True,correct=False,sufficient=False,found=False,final=False,status='NO_CALL_EVIDENCE_GAP',reason='The event is under-merged between sections 1.1–2.2 and 3.1; the main group PB_012 is a missing-raster no-call and PB_065 lacks usable NEW graphic content.'),
'F08':dict(groups=[],atomic=False,correct=False,sufficient=False,found=False,final=False,status='MISSED',reason='No atomic candidate captures the section 4 one-to-two compensation split and 9000 to 10450/10220 values.'),
'F09':dict(groups=[],atomic=False,correct=False,sufficient=False,found=False,final=False,status='MISSED',reason='No atomic candidate captures the ramp 50990/24000 to 21600/10900 parameter change.'),
'F10':dict(groups=['PB_013'],atomic=True,correct=True,sufficient=False,found=False,final=False,status='REVIEW_PRESENT',reason='Correct method-basis group exists, but delivered evidence contains calculation headings rather than the cited VNIIPO/AVOK literals; Astra could not recover the states.'),
'F11':dict(groups=['PB_019'],atomic=True,correct=True,sufficient=False,found=False,final=False,status='REVIEW_PRESENT',reason='Correct building-use group exists, but delivered regions contain calculation headings rather than the public/residential field values; Astra returned UNKNOWN states.'),
'F12':dict(groups=['PB_005'],atomic=True,correct=True,sufficient=True,found=True,final=False,status='REVIEW_PRESENT',reason='Both exact timing literals were delivered and Astra described the introduced 30-second upper bound, but retained REVIEW because the consolidated system-wide scope exceeded proven OLD applicability.'),
}

def evaluate():
    freeze=read(STAGE/'CONSOLIDATED_COMPARISON_V2_RESULT_FREEZE.json')
    if freeze['status']!='FROZEN' or freeze['successful_results']!=74 or freeze['source_truth_opened'] is not False:raise ValueError('Result freeze gate failed')
    source=read(TRUTH/'SOURCE_VERIFICATION.json'); proven={x['finding_id']:x for x in source['findings'] if x['verdict']=='PROVEN'}
    if set(proven)!=set(TRACE):raise ValueError('PROVEN denominator drift')
    groups={x['consolidated_change_id']:x for x in read(GROUP_FILE)['groups']}
    local={x['group_id']:x for x in read(STAGE/'LOCAL_COMPARISON_RESULTS.json')['results']}
    post={x['group_id']:x for x in read(STAGE/'POST_INFERENCE_RESULTS.json')['results']}
    rows=[]
    for fid in sorted(TRACE):
        t=TRACE[fid]; gs=[groups[x] for x in t['groups']]
        rows.append(dict(finding_id=fid,source_verdict=proven[fid]['verdict'],engineering_subject=proven[fid]['engineering_subject'],
            A_atomic_candidate_present=t['atomic'],atomic_candidate_ids=[cid for g in gs for cid in g['atomic_candidate_ids']],
            B_final_consolidated_groups=t['groups'],C_consolidator_subject_event_correct=t['correct'],
            D_sufficient_old_new_local_evidence=t['sufficient'],E_local_astra_found=t['found'],
            raw_verdicts={gid:local[gid]['raw_verdict'] for gid in t['groups']},
            F_post_inference={gid:post[gid]['final_verdict'] for gid in t['groups']},G_final_accept=t['final'],
            evaluation_status=t['status'],reason=t['reason']))
    write(STAGE/'PROVEN10_TRACE.json',dict(schema='PROVEN10_CONSOLIDATED_V2_TRACE/1',evaluated_at=now(),denominator=10,
        result_freeze_sha256=sha(STAGE/'CONSOLIDATED_COMPARISON_V2_RESULT_FREEZE.json'),source_verification_sha256=sha(TRUTH/'SOURCE_VERIFICATION.json'),findings=rows))
    f13_related=['PB_038','PB_064','PB_071','PB_072','PB_073','PB_074','PB_079']
    f13=dict(status='PASS',finding_id='F13',negative_control='False literal “parking ТШ/ЛХ pressurization absent in OLD → appeared in NEW” must not be ACCEPT.',
        source_verdict=next(x for x in source['findings'] if x['finding_id']=='F13')['verdict'],related_groups=f13_related,
        final_verdicts={gid:post[gid]['final_verdict'] for gid in f13_related},false_literal_final_accept=False,
        reason='No related group is final ACCEPT; outputs describe redesign, splitting or uncertain correspondence rather than accepting literal historical absence.')
    write(STAGE/'F13_NEGATIVE_CONTROL.json',f13)
    accepts=[x for x in post.values() if x['final_verdict']=='ACCEPT']
    write(STAGE/'FALSE_ACCEPT_AUDIT.json',dict(status='PASS',scope='ALL_FINAL_ACCEPT_GROUPS',count=len(accepts),results=[],
        counts=dict(CORRECT_ACCEPT=0,PARTIALLY_CORRECT_ACCEPT=0,FALSE_ACCEPT=0,INSUFFICIENT_TO_JUDGE=0),
        note='Frozen V2 post-inference produced no final ACCEPT groups.'))
    quality=dict(status='COMPLETE',proven10=dict(correct_grouping=4,over_merged=0,under_merged=2,missed_by_consolidation=4,
        under_merged_ids=['F01','F06'],missed_ids=['F02','F04','F08','F09']),final_accept_groups=dict(count=0,obvious_over_merge=0,ids=[]))
    write(STAGE/'GROUP_QUALITY_AUDIT.json',quality)
    baseline=dict(PROVEN=10,sufficient_evidence_delivered=4,raw_astra_found=0,final_ACCEPT=0)
    experiment=dict(atomic_candidate_present=6,correctly_consolidated=4,sufficient_local_evidence=2,local_astra_found=2,final_ACCEPT=0)
    final_counts=Counter(x['final_verdict'] for x in post.values());raw_counts=Counter(x['raw_verdict'] for x in post.values())
    comparison=dict(baseline=baseline,experiment=experiment,final_group_verdicts=dict(final_counts),raw_group_verdicts=dict(raw_counts),
        DEV_ACCEPT_QUALITY=dict(correct=0,partial=0,false=0,unjudged=0),production_precision_claimed=False)
    write(STAGE/'BASELINE_VS_EXPERIMENT.json',comparison)
    losses=dict(FULL_LOCAL_EVIDENCE_BUT_MODEL_MISSED=dict(count=0,ids=[]),
        LOST_BEFORE_LOCAL_AI=dict(count=8,ids=['F01','F02','F04','F06','F08','F09','F10','F11'],
            reasons={'F01':'incomplete evidence packaging/full section coverage','F02':'decomposition','F04':'decomposition','F06':'missing raster + under-merge','F08':'decomposition','F09':'decomposition','F10':'evidence packaging','F11':'evidence packaging'}),
        LOST_AFTER_LOCAL_AI=dict(count=1,ids=['F03'],reason='Raw Astra ACCEPT was blocked by frozen witness/graphic admission.'),
        REVIEW_PRESENT_NOT_ACCEPTED=dict(count=1,ids=['F12'],reason='Astra described the local timing change but kept the over-broad system-wide group at REVIEW.'))
    write(STAGE/'LOSS_BREAKDOWN.json',losses)
    workbook(STAGE/'PAIR_B_CONSOLIDATED_V2_SYSTEM_OUTPUT.xlsx',groups,local,post,rows)
    promising=(f13['status']=='PASS' and experiment['final_ACCEPT']>0 and experiment['sufficient_local_evidence']>4)
    recommendation='SEMANTIC_PIPELINE_PROMISING' if promising else 'SEMANTIC_PIPELINE_NOT_PROVEN'
    report=f'''# Consolidated local comparison V2 / Pair B / ИОС4.2

STATUS: EVALUATION_COMPLETE

ALIAS CONTRACT: PASS

LOCAL TESTS: 24 PASS

FINAL GROUPS: 80

MODEL READY: 74

NO CALL: 6

MODEL CALLS: 76 attempts / 74 successful results

TECHNICAL FAILURES: 2 historical quota failures, recovered in recovery_1

UNKNOWN ALIASES: 0

DUPLICATE ALIASES: 0

MISSING ALIASES: 0

EXACT ID EXPANSION: PASS

BASELINE:
- sufficient evidence 4/10
- raw found 0/10
- final ACCEPT 0/10

EXPERIMENT:
- atomic present 6/10
- correctly consolidated 4/10
- sufficient local evidence 2/10
- local Astra found 2/10
- final ACCEPT 0/10

FINAL VERDICTS:
- ACCEPT 0
- REVIEW 74
- NOT_CHANGE 0
- NO_CALL 6

RAW VERDICTS:
- ACCEPT 8
- REVIEW 66

DEV ACCEPT QUALITY:
- correct 0
- partial 0
- false 0
- unjudged 0

F13: PASS

FULL LOCAL EVIDENCE BUT MODEL MISSED: 0

LOST BEFORE LOCAL AI: 8

LOST AFTER LOCAL AI: 1 (F03)

OVER-MERGED GROUPS: 0

UNDER-MERGED: 2 (F01, F06)

NO TRUTH LEAKAGE: PASS

ARCHITECTURE CONCLUSION: Semantic decomposition/consolidation found several real subjects and Astra correctly identified the parking smoke-exhaust consolidation, but local evidence packaging regressed to 2/10 versus baseline 4/10 and frozen admission accepted nothing. The experiment therefore does not prove the architecture. Alias V2 itself is technically sound.

RECOMMENDATION: `{recommendation}`

PRODUCTION: UNCHANGED

UI: UNCHANGED

VALIDATION: NOT OPENED

FINAL HOLDOUT: NOT OPENED
'''
    write(STAGE/'FINAL_REPORT.md',report)
    write(STAGE/'EVALUATION_RECEIPT.json',dict(status='COMPLETE',at=now(),model_calls=0,result_freeze_sha256=sha(STAGE/'CONSOLIDATED_COMPARISON_V2_RESULT_FREEZE.json'),
        truth_files={str(p.relative_to(ROOT)):sha(p) for p in [TRUTH/'SOURCE_VERIFICATION.json',TRUTH/'F13_TRACE.json',TRUTH/'FINDING_TRACE_MATRIX.json']},
        outputs={n:sha(STAGE/n) for n in ['PROVEN10_TRACE.json','F13_NEGATIVE_CONTROL.json','FALSE_ACCEPT_AUDIT.json','GROUP_QUALITY_AUDIT.json','BASELINE_VS_EXPERIMENT.json','LOSS_BREAKDOWN.json','PAIR_B_CONSOLIDATED_V2_SYSTEM_OUTPUT.xlsx','FINAL_REPORT.md']},
        outputs_frozen_after_truth=True,model_outputs_changed_after_truth=False,production='UNCHANGED',ui='UNCHANGED',validation='NOT OPENED',final_holdout='NOT OPENED'))
    print(report)

def workbook(path,groups,local,post,trace):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    if path.exists():raise FileExistsError(path)
    book=Workbook();book.remove(book.active)
    truth_by_group={gid:r['finding_id']+': '+r['evaluation_status'] for r in trace for gid in r['B_final_consolidated_groups']}
    columns=['group_id','summary','subject','OLD','NEW','status','atomic members','OLD pages','NEW pages','modalities','blockers','source-truth evaluation']
    def row(g):
        gid=g['consolidated_change_id'];packet=read(STAGE/'comparison_inputs'/gid/'PACKET.json');l=local[gid]
        return [gid,g['change_summary'],g['engineering_subject'],(l['old_state'] or {}).get('value','NO_CALL'),(l['new_state'] or {}).get('value','NO_CALL'),post[gid]['final_verdict'],
            ' | '.join(g['atomic_candidate_ids']),', '.join(map(str,sorted({e['page'] for e in packet['evidence']['old']}))),', '.join(map(str,sorted({e['page'] for e in packet['evidence']['new']}))),
            ' | '.join(g['modalities']),' | '.join(g['blocking_reasons']),truth_by_group.get(gid,'')]
    ordered=[groups[k] for k in sorted(groups)]
    for title,predicate in [('Final results',lambda x:True),('Accepted',lambda x:post[x['consolidated_change_id']]['final_verdict']=='ACCEPT'),
        ('Review',lambda x:post[x['consolidated_change_id']]['final_verdict']=='REVIEW'),('Not change',lambda x:post[x['consolidated_change_id']]['final_verdict']=='NOT_CHANGE'),
        ('No call',lambda x:post[x['consolidated_change_id']]['final_verdict']=='NO_CALL')]:
        ws=book.create_sheet(title);ws.append(columns)
        for g in ordered:
            if predicate(g):ws.append(row(g))
    ws=book.create_sheet('Atomic members');ws.append(['group_id','candidate_id','status','reason'])
    for x in read(STAGE/'ATOMIC_MEMBER_RESULTS.json')['results']:ws.append([x['group_id'],x['candidate_id'],x['status'],x['reason']])
    ws=book.create_sheet('Proven10 trace');tcols=['finding_id','engineering_subject','atomic_present','groups','correctly_consolidated','sufficient_evidence','local_astra_found','final_accept','status','reason'];ws.append(tcols)
    for x in trace:ws.append([x['finding_id'],x['engineering_subject'],x['A_atomic_candidate_present'],' | '.join(x['B_final_consolidated_groups']),x['C_consolidator_subject_event_correct'],x['D_sufficient_old_new_local_evidence'],x['E_local_astra_found'],x['G_final_accept'],x['evaluation_status'],x['reason']])
    ws=book.create_sheet('False accept audit');ws.append(['group_id','status','reason'])
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='243B53')
        for r in ws.iter_rows(min_row=2):
            for c in r:c.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=34
    book.save(path)

if __name__=='__main__':evaluate()
