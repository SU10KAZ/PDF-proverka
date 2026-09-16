"""One guarded local replay of Pair A V2. No provider imports or calls.

Run with python -m experiments.project_change_post_inference_repair_272.replay.
The source audit is deliberately inaccessible until REPLAY_RESULT_FREEZE exists.
"""
from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path('/home/coder/auditmanager/corpus-audits/20260914_project_change_272')
REPO = Path(__file__).resolve().parents[2]
LIVE = ROOT / 'fresh_dev_pair_a_live_v2'
BASE = ROOT / 'fresh_dev_sample_f5_pipeline_v6'
OUT = ROOT / 'fresh_dev_pair_a_post_inference_repair_v1'
sys.path.insert(0, str(ROOT / 'controlled_inference_f1_f4_f2_v3/runtime_deps'))

from experiments.project_change_272.policy import admitted_pairs
from experiments.project_change_contracts_v2_272.sufficiency import PROFILES
from experiments.project_change_f2_binding_v4_272.normalization import normalize
from experiments.project_change_semantic_codex_v2_272.contracts import resolve_ownership
from .identity import contract
from .applicability import RULES


def read(path): return json.loads(Path(path).read_text())
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def now(): return datetime.now(timezone.utc).isoformat()
def ref(path): return dict(path=str(path), sha256=sha(path))


def save(name, value):
    path = OUT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    data = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + '\n'
    if path.exists() and path.read_text() != data:
        raise ValueError('Immutable replay output already exists: ' + name)
    path.write_text(data)


def offline_guard():
    def guard(event, args):
        if event.startswith('socket.') or event in {'subprocess.Popen', 'os.system', 'os.posix_spawn', 'os.exec', 'os.fork'}:
            raise PermissionError('Offline replay prohibits network/process execution: ' + event)
        if event == 'open' and isinstance(args[0], (str, bytes)):
            path = Path(args[0]).resolve()
            if path.is_relative_to(ROOT / 'validation') or 'FINAL_HOLDOUT' in path.parts:
                raise PermissionError('Reserve evidence is sealed')
            if path.is_relative_to(LIVE):
                flags = args[2] or 0
                if flags & 3 or flags & 512:
                    raise PermissionError('V2 is read-only')
                if any(x in path.name for x in ('SOURCE_AUDIT', 'SOURCE_FIRST_TRUTH', 'METRICS', 'MISSED_REAL')):
                    if not (OUT / 'REPLAY_RESULT_FREEZE.json').exists():
                        raise PermissionError('Frozen audit cannot be read before replay')
    sys.addaudithook(guard)


def verify_v2():
    expected = read(OUT / 'V2_BASELINE_HASHES.json')
    # Baseline includes audit hashes. Hash-only integrity is performed outside
    # the semantic read guard before replay, and again after results freeze.
    actual = {str(p.relative_to(LIVE)): sha(p) for p in sorted(LIVE.rglob('*')) if p.is_file()}
    if expected != actual: raise ValueError('V2 immutable artifacts changed')
    return dict(files=len(actual), unchanged=True)


def profile(raw):
    role = raw['new_state']['state_role']
    if raw.get('old_absence', {}).get('mode') in {'EXPLICIT_NEGATIVE', 'COMPLETE_BOUNDED_REPRESENTATION'} and role in {'TOPOLOGY', 'INSTALLED_CONFIGURATION'}:
        return 'NOVEL_SYSTEM'
    if role == 'TOPOLOGY':
        return 'TOPOLOGY_DECLARATION' if all(raw[s + '_state']['evidence_form'] == 'DECLARATION' for s in ('old', 'new')) else 'TOPOLOGY'
    if role in {'CALCULATED_RESULT', 'CAPACITY', 'COUNT'}: return 'AGGREGATE_CALCULATED_RESULT'
    return role if role in PROFILES else 'UNKNOWN'


def f2_verdict(n):
    f2 = n.get('f2') or {}
    exists = f2.get('exists_change', {}).get('status')
    if exists == 'NO': return 'NOT_CHANGE'
    if exists == 'YES' and f2['comparability']['status'] == 'COMPARABLE' and n.get('sufficiency', {}).get('sufficient'):
        return 'ACCEPT'
    return 'REVIEW'


def local_tests():
    modules = ['experiments.project_change_post_inference_repair_272.test_repair',
               'experiments.project_change_f2_binding_v4_272.test_binding',
               'experiments.project_change_contracts_272.test_comparability',
               'experiments.project_change_contracts_272.test_states']
    suite = unittest.TestSuite(unittest.defaultTestLoader.loadTestsFromName(m) for m in modules)
    stream = io.StringIO(); result = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
    receipt = dict(status='PASS' if result.wasSuccessful() else 'FAIL', tests=result.testsRun,
                   failures=len(result.failures), errors=len(result.errors), modules=modules, model_calls=0)
    save('LOCAL_TESTS.txt', stream.getvalue()); save('LOCAL_TEST_RECEIPT.json', receipt)
    if not result.wasSuccessful(): raise ValueError('Local tests failed')
    return receipt


def workbook(rows, groups):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook(); sheet = book.active; sheet.title = 'System output'
    columns = ['ID', 'summary', 'OLD', 'NEW', 'raw verdict', 'F2 verdict', 'final verdict',
               'OLD evidence pages', 'NEW evidence pages', 'binding status', 'ProjectChange group',
               'remaining reasons', 'claim type']
    sheet.append(columns)
    for r in rows:
        raw, n = r['raw'], r['normalized']
        reasons = n.get('issues', []) + n.get('sufficiency', {}).get('reasons', [])
        if n.get('f2'): reasons += n['f2']['comparability']['different'] + n['f2']['comparability']['unknown']
        sheet.append([r['package_id'], raw.get('project_change_summary', r['source_label']),
            raw.get('old_state', {}).get('value', 'NOT_EVALUATED'), raw.get('new_state', {}).get('value', 'NOT_EVALUATED'),
            raw.get('verdict', 'NO_CALL_MISSING'), r['f2_verdict'], r['status'],
            ', '.join(map(str, r['evidence_pages']['old'])), ', '.join(map(str, r['evidence_pages']['new'])),
            r['binding_status'], groups.get(r['package_id'], ''), '; '.join(dict.fromkeys(reasons)),
            (n.get('claim_applicability') or {}).get('claim_type', 'NOT_EVALUATED')])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    widths = [34, 80, 65, 65, 18, 18, 18, 25, 25, 25, 26, 65, 28]
    for i, width in enumerate(widths, 1): sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='244062')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')): cell.data_type = 's'
    path = OUT / 'PAIR_A_REPAIRED_SYSTEM_OUTPUT.xlsx'
    if path.exists(): raise ValueError('Workbook already exists')
    book.save(path)
    check = load_workbook(path, read_only=True)
    assert check.active.max_row == 60 and check.active.max_column == len(columns)
    check.close()


def main():
    if (OUT / 'REPLAY_RESULT_FREEZE.json').exists(): raise ValueError('Replay already frozen; do not rerun')
    integrity = verify_v2()
    offline_guard()
    pair = admitted_pairs('DEV', indices=[2])[0]
    assert pair['pair_key'] == 'ad0a31a342a666082f2ef66a'
    tests = local_tests()
    files = list(Path(__file__).parent.glob('*.py')) + [REPO / 'experiments/project_change_f2_binding_v4_272' / n
             for n in ('binding.py', 'normalization.py', 'admission.py')]
    save('REPAIR_CODE_FREEZE.json', dict(at=now(), code={str(p.relative_to(REPO)): sha(p) for p in files},
        source_audit_opened=False, model_calls=0, v2_integrity=integrity, tests=tests))
    save('BINDING_CONTRACT_V1.json', dict(schema='BINDING_CONTRACT/1',
        technical_candidate_id='Transport provenance only; never compared to model engineering_subject',
        canonical_engineering_subject='Stable source IDs + pair namespace + structured physical scope; cross-version IDs alone prove no equivalence',
        model_wording='Display only', source_pages='Side/version-specific provenance, not cross-version equality',
        evidence_admission=['Explicit raw role/path', 'side', 'version', 'source receipt', 'requirement/region link',
                            'candidate provenance', 'discovery ID', 'validated witness', 'scope/room witness'],
        uncertainty='REVIEW', legacy_callers='Unchanged unless subject_contract explicitly supplied'))
    save('CLAIM_APPLICABILITY_RULES.json', dict(rules=RULES,
        stage_phase='Existing phase_is_applicable: explicit construction/calculation phase remains applicable',
        materiality='Unchanged; OTHER remains unsupported by existing sufficiency profile',
        identity='Equal witnessed physical identity and members; room role changes additionally need room number and position mapping'))
    plan = read(LIVE / 'CALL_PLAN.json')['packages']
    frozen = read(LIVE / 'PAIR_A_INFERENCE_FREEZE.json')
    results, events, evidence, binding_rows, f2_rows, inputs = [], [], {}, [], [], {}
    for row in plan:
        key = row['key']; path = BASE / row['package']; p = read(path)
        assert p['provenance']['pair_index'] == 2 and p['provenance']['pair_key'] == pair['pair_key']
        assert sha(path) == frozen['input_files'][str(path)]
        inputs[str(path)] = sha(path)
        packet = deepcopy(p['evidence_packet'])
        for side in ('old', 'new'):
            for e in packet['evidence'][side]:
                assert e['source_receipt'] == pair[side]['artifacts']['pdf']
                if e.get('raster'):
                    e['raster']['path'] = str(BASE / e['raster']['path'])
                    assert sha(e['raster']['path']) == e['raster']['sha256']
                evidence[e['evidence_id']] = e
        before = read(LIVE / f'normalized_responses/{key}.json')
        if row['action'] == 'MODEL_CALL':
            raw_path = LIVE / f'raw_responses/{key}.json'; raw = read(raw_path)
            inputs[str(raw_path)] = sha(raw_path)
            assert raw['case_token'] == key
            n = normalize(raw, packet, profile(raw), subject_contract=contract(p['candidate_subject'], packet))
        else:
            raw, n = {}, deepcopy(before)
        save(f'normalized_responses/{key}.json', n)
        status = n['effective_verdict']
        bound_pages = {s: sorted({b['provenance']['page'] for b in n.get(s + '_state', {}).get('evidence_bindings', [])}) for s in ('old', 'new')}
        entry = dict(package_id=key, status=status, f2_verdict=f2_verdict(n), raw=raw, normalized=n,
            source_label=p['candidate_subject']['subject'], completeness=row['completeness'], evidence_pages=bound_pages,
            binding_status='NO_CALL_MISSING' if not raw else 'REVIEW' if n.get('rejected_evidence_bindings') else 'BOUND')
        results.append(entry)
        binding_rows.append(dict(package_id=key, before=before.get('rejected_evidence_bindings', []),
            after=n.get('rejected_evidence_bindings', []), before_issues=before.get('issues', []), after_issues=n.get('issues', []),
            bound_role_links=sum(len(n.get(s + '_state', {}).get('evidence_bindings', [])) for s in ('old', 'new'))))
        f2_rows.append(dict(package_id=key, before=f2_verdict(before), after=f2_verdict(n),
            before_details=before.get('f2'), after_details=n.get('f2'), applicability=n.get('claim_applicability')))
        if status == 'ACCEPT':
            # Existing event shape/grouping, with explicitly bound evidence only.
            ids = sorted({i for s in ('old', 'new') for i in n[s + '_state']['evidence_ids']})
            events.append(dict(event_id=key, object_id=272, status='ACCEPT', engineering_subject=raw['new_state']['engineering_subject'],
                summary_ru=raw['project_change_summary'], old_state=n['old_state'], new_state=n['new_state'], evidence_ids=ids))
    assert len(results) == 59 and sum(bool(r['raw']) for r in results) == 52
    grouped = resolve_ownership(events, [], evidence)
    group_ids = {i: 'PC_' + str(index).zfill(3) for index, g in enumerate(grouped['groups'], 1) for i in g['member_ids']}
    save('PROJECT_CHANGES_REPLAY.json', dict(**grouped, group_ids=group_ids, implementation='resolve_ownership unchanged; no pairwise proofs supplied'))
    save('PAIR_A_REPAIRED_RESULTS.json', results)
    save('BINDING_BEFORE_AFTER.json', dict(rows=binding_rows,
        subject_mismatch_before=sum('BINDING_SUBJECT_MISMATCH' in r['before_issues'] for r in binding_rows),
        subject_mismatch_after=sum('BINDING_SUBJECT_MISMATCH' in r['after_issues'] for r in binding_rows),
        any_binding_failure_after=sum(bool(r['after']) for r in binding_rows)))
    save('F2_BEFORE_AFTER.json', dict(rows=f2_rows,
        before=dict(Counter(r['before'] for r in f2_rows)), after=dict(Counter(r['after'] for r in f2_rows))))
    counts = dict(Counter(r['status'] for r in results))
    save('MECHANICAL_REPLAY_52.json', dict(status='REPLAY_COMPLETED', at=now(), model_calls=0,
        codex_astra_calls=0, openrouter_calls=0, claude_calls=0, replayed=52, no_call_missing=7,
        raw_counts=dict(Counter(r['raw']['verdict'] for r in results if r['raw'])), final_counts=counts,
        f2_counts=dict(Counter(r['f2_verdict'] for r in results)), project_changes=len(grouped['groups']),
        inputs=inputs, source_audit_opened=False, validation='NOT OPENED', final_holdout='NOT OPENED', pair_b='NOT RUN'))
    workbook(results, group_ids)
    save('REPLAY_RESULT_FREEZE.json', dict(at=now(), source_audit_opened=False, model_calls=0,
        files={str(p.relative_to(OUT)): sha(p) for p in sorted(OUT.rglob('*')) if p.is_file()}))
    print(json.dumps(dict(counts=counts, project_changes=len(grouped['groups']), tests=tests, immutable_v2=verify_v2()), ensure_ascii=False))


if __name__ == '__main__': main()
