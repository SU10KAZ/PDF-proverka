"""Traceable Excel export; model acceptance is never relabeled as source truth."""
import argparse
from pathlib import Path

from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

from experiments.project_change_272.inventory import ROOT,read,immutable,sha,now
from .access import authorize
from .packets import BASE
from .run import check_packet_scope


def append(sheet,values):
    sheet.append(values)
    # Source text is data, including strings which resemble spreadsheet formulas.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value,str):cell.data_type='s'


def export(name,ownership,partition='DEV',candidate=None,judgements=None,*,artifact_base=None):
    base=BASE if artifact_base is None else Path(artifact_base)
    allowed={p['index']:p for p in authorize(partition,candidate)}
    root=base/'ownership'/ownership
    manifest=read(root/'MANIFEST.json');pred=read(root/'PROJECT_CHANGES.json')
    if manifest['partition']!=partition:raise PermissionError('Wrong report partition')
    assessed={r['project_change_id']:r for r in read(judgements)} if judgements else {}
    out=base/'reports'/name
    immutable(out/'MANIFEST.json',dict(created_at=now(),partition=partition,source_ownership=ownership,
        source_sha256=sha(root/'PROJECT_CHANGES.json'),split_sha256=sha(ROOT/'SPLIT.json'),code_sha256=sha(__file__),
        candidate_manifest=str(candidate) if candidate else None,
        judgements_sha256=sha(judgements) if judgements else None,source_adjudication='EXPLICIT_PER_ROW_ONLY'))
    book=Workbook();summary=book.active;summary.title='Контекст'
    for row in [
        ['Объект','Садовническая 76 / Балчуг Эстейт, object 272'],
        ['Сравнение','OLD stage_1 -> NEW stage_2; logical v002 baseline'],
        ['Набор',partition],['Исторически слепой',False],
        ['Статус','Исследовательский экспорт. Принятие моделью не означает подтверждение по исходникам.'],
        ['Группы',len(pred['project_changes'])],['Принято моделью',pred['accepted']],
        ['Подтверждение по исходникам','Только явные оценки в колонке исходной проверки; незаполненное = НЕ ПРОВЕРЕНО'],
        ['Исходный результат',str(root/'PROJECT_CHANGES.json')],['SHA256 результата',sha(root/'PROJECT_CHANGES.json')],
        ['Замороженное разбиение',str(ROOT/'SPLIT.json')],['SHA256 разбиения',sha(ROOT/'SPLIT.json')],
        ['Ограничение','Полнота и precision определяются отдельной исходной оценкой и её знаменателями, а не количеством строк.']]:append(summary,row)
    changes=book.create_sheet('ProjectChange')
    append(changes,['ID','Пара','Шифр OLD','Шифр NEW','Решение модели','Исходная проверка','Инженерный предмет','Изменение','Объединение','Свидетельств-состояний','Проверка объединения'])
    facts=book.create_sheet('Состояния')
    append(facts,['ProjectChange ID','Member ID','№ факта','Параметр / состояние','OLD','NEW','OLD свидетельства','NEW свидетельства','OLD контекст','NEW контекст'])
    evidence=book.create_sheet('Свидетельства')
    append(evidence,['ProjectChange ID','Member ID','№ факта','Сторона','Evidence ID','Маршрут','Страница PDF','Точная цитата / транскрипция','Версия документа','PDF','SHA256 PDF','Координаты PDF','Растр','SHA256 растра','Пакет','SHA256 пакета'])
    reviews=book.create_sheet('Очередь проверки')
    append(reviews,['Тип','ID','Причина / состояние','Происхождение'])
    for r in read(root/'INPUT_REVIEWS.json'):
        append(reviews,['Входное состояние',r['member_id'],str(r['product_admission'])+' '+str(r['semantic_audit']),r['source_packet']['path']])
    for r in pred['review_members']:append(reviews,['Объединение',r['member_id'],str(r['reason']),str(root/'PROJECT_CHANGES.json')])
    evidence_count=0;fact_count=0
    for c in pred['project_changes']:
        pair=allowed.get(c['pair_index'])
        if pair is None:raise PermissionError('Foreign cipher in report')
        verdict=assessed.get(c['project_change_id'],{}).get('verdict','НЕ ПРОВЕРЕНО')
        append(changes,[c['project_change_id'],c['pair_index'],pair['old']['document_code'],pair['new']['document_code'],c['status'],
            verdict,c['engineering_subject'],c['summary_ru'],c['relation'],len(c['state_bundles']),str(c['ownership_audit'])])
        if c['status']=='REVIEW':append(reviews,['ProjectChange',c['project_change_id'],str(c['ownership_audit']),str(root/'PROJECT_CHANGES.json')])
        for m in c['state_bundles']:
            path=Path(m['source_packet']['path'])
            if sha(path)!=m['source_packet']['sha256']:raise ValueError('Report source packet drift')
            p=read(path);check_packet_scope(p,allowed,partition)
            by_id={e['evidence_id']:e for side in ['old','new'] for e in p['evidence'][side]}
            for n,f in enumerate(m['event']['facts'],1):
                append(facts,[c['project_change_id'],m['member_id'],n,f['property'],f['old_value'],f['new_value'],
                    ', '.join(w['evidence_id'] for w in f['old_witnesses']),', '.join(w['evidence_id'] for w in f['new_witnesses']),
                    m['event']['old_state'],m['event']['new_state']]);fact_count+=1
                for side in ['old','new']:
                    for w in f[side+'_witnesses']:
                        e=by_id[w['evidence_id']];pdf=e['source_receipt'];raster=e.get('raster',{})
                        append(evidence,[c['project_change_id'],m['member_id'],n,side,w['evidence_id'],w['route'],e['page'],w['quote'],
                            e['document_version'],pdf['path'],pdf['sha256'],str(e['bbox']),raster.get('path'),raster.get('sha256'),str(path),sha(path)])
                        evidence_count+=1
    for sheet in book:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29485F')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width=min(65,max(16,max(len(str(c.value or '')) for c in column[:50])*.7))
    target=out/'ProjectChange.xlsx';book.save(target)
    with target.open('rb') as stream:
        check=load_workbook(stream,read_only=True,data_only=False)
        if check['Состояния'].max_row!=fact_count+1 or check['Свидетельства'].max_row!=evidence_count+1:raise ValueError('Export row count mismatch')
        check.close()
    immutable(out/'EXPORT_RECEIPT.json',dict(path=str(target),sha256=sha(target),changes=len(pred['project_changes']),facts=fact_count,
        witness_rows=evidence_count,source_adjudication='Only explicit row judgements; unchanged model status preserved'))
    print(target,flush=True)
    return target


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--name',required=True);p.add_argument('--ownership',required=True)
    p.add_argument('--partition',default='DEV',choices=['DEV','VALIDATION','FINAL_HOLDOUT']);p.add_argument('--candidate',type=Path);p.add_argument('--judgements',type=Path)
    a=p.parse_args();export(a.name,a.ownership,a.partition,a.candidate,a.judgements)
