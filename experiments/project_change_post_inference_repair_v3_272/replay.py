"""Single immutable replay after local tests. Frozen source audit remains sealed."""
from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

from experiments.project_change_post_inference_repair_272.identity import contract
from experiments.project_change_post_inference_repair_272.replay import f2_verdict
from experiments.project_change_semantic_codex_v2_272.contracts import resolve_ownership
from experiments.project_change_272.policy import admitted_pairs
from .applicability import RULES, NEGATIVE_CONTRACT
from .repair import repair

ROOT = Path('/home/coder/auditmanager/corpus-audits/20260914_project_change_272')
REPO = Path(__file__).resolve().parents[2]
LIVE = ROOT / 'fresh_dev_pair_a_live_v2'
V1 = ROOT / 'fresh_dev_pair_a_post_inference_repair_v1'
BASE = ROOT / 'fresh_dev_sample_f5_pipeline_v6'
V2 = ROOT / 'fresh_dev_pair_a_post_inference_repair_v2'
OUT = ROOT / 'fresh_dev_pair_a_post_inference_repair_v3'


def read(path): return json.loads(Path(path).read_text())
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def now(): return datetime.now(timezone.utc).isoformat()


def save(name, value):
    p = OUT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.exists(): raise ValueError('Immutable output exists: ' + name)
    p.write_text(value if isinstance(value, str) else json.dumps(value,ensure_ascii=False,indent=2,sort_keys=True)+'\n')


def verify_inputs():
    files = read(OUT / 'INPUT_FREEZE.json')['files']
    changed = [p for p,h in files.items() if sha(p) != h]
    if changed: raise ValueError('Frozen input drift: ' + str(changed))
    metadata = read(OUT / 'INPUT_FREEZE.json')['immutable_tree_metadata']
    current = {str(p): [p.stat().st_size, p.stat().st_mtime_ns] for d in (LIVE, V1, V2) for p in d.rglob('*') if p.is_file()}
    if metadata != current: raise ValueError('Immutable prior artifact metadata changed')
    return dict(immutable_prior_artifacts=len(current), files=len(files), unchanged=True,
                raw_responses=sum('/raw_responses/' in p for p in files if p.startswith(str(LIVE))))


def offline_guard():
    sealed = {'SOURCE_AUDIT', 'SOURCE_FIRST_TRUTH', 'MISSED_REAL', 'AUDIT_COMPARISON', 'METRICS', 'FINAL_REPORT', 'RESULT_RECEIPT'}
    def guard(event, args):
        if event.startswith('socket.') or event in {'subprocess.Popen','os.system','os.posix_spawn','os.exec','os.fork'}:
            raise PermissionError('Mechanical replay prohibits network/process calls: ' + event)
        if event == 'open' and isinstance(args[0],(str,bytes)):
            path = Path(args[0]).resolve()
            if any(x.lower() in {'validation', 'final_holdout'} for x in path.parts):
                raise PermissionError('Reserve sealed')
            if path.is_relative_to(LIVE) or path.is_relative_to(V1) or path.is_relative_to(V2) or path.is_relative_to(BASE):
                if (args[2] or 0) & (3 | 512):
                    raise PermissionError('Frozen input write prohibited')
                if any(x in path.name for x in sealed) and not (OUT/'REPLAY_RESULT_FREEZE.json').exists():
                    raise PermissionError('Source audit not accessible before replay freeze')
    sys.addaudithook(guard)


TEST_MODULES = [
    'experiments.project_change_post_inference_repair_v3_272.test_repair',
    'experiments.project_change_post_inference_repair_v2_272.test_repair',
    'experiments.project_change_post_inference_repair_272.test_repair',
    'experiments.project_change_f2_binding_v4_272.test_binding',
    'experiments.project_change_contracts_272.test_comparability',
    'experiments.project_change_contracts_272.test_states',
]


def tests():
    suite=unittest.TestSuite(unittest.defaultTestLoader.loadTestsFromName(m) for m in TEST_MODULES)
    output=io.StringIO(); result=unittest.TextTestRunner(stream=output,verbosity=2).run(suite)
    receipt=dict(status='PASS' if result.wasSuccessful() else 'FAIL', tests=result.testsRun,
                 failures=len(result.failures),errors=len(result.errors),modules=TEST_MODULES,model_calls=0)
    if not result.wasSuccessful(): raise ValueError(output.getvalue())
    save('LOCAL_TESTS.txt',output.getvalue()); save('TEST_RECEIPT.json',receipt)
    return receipt


def reasons(n):
    c=(n.get('f2') or {}).get('comparability',{})
    return list(dict.fromkeys(n.get('issues',[])+n.get('sufficiency',{}).get('reasons',[])+c.get('unknown',[])+c.get('different',[])))


def workbook(rows, groups, before_after, negatives):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book=Workbook(); sheet=book.active; sheet.title='System output'
    sheet.append(['candidate_id','Claim','OLD','NEW','Raw verdict','F2 verdict','Final verdict',
        'OLD pages','NEW pages','Group','Reasons','Witness applicability','Graphic requirement',
        'Calculation basis','Original role','Normalized role','Numeric conflict'])
    for r in rows:
        raw,n=r['raw'],r['normalized']; a=n.get('claim_applicability') or {}
        sheet.append([r['package_id'],raw.get('project_change_summary',r['source_label']),
            raw.get('old_state',{}).get('value'),raw.get('new_state',{}).get('value'),
            raw.get('verdict','NO_CALL_MISSING'),r['f2_verdict'],r['status'],
            ', '.join(map(str,r['evidence_pages']['old'])),', '.join(map(str,r['evidence_pages']['new'])),
            groups.get(r['package_id']),'; '.join(reasons(n)),a.get('supporting_witness',{}).get('applicability'),
            a.get('graphic',{}).get('requirement'),a.get('calculation_basis',{}).get('applicability'),
            ', '.join(a.get('other_role',{}).get('original_role',[])),', '.join(a.get('other_role',{}).get('normalized_role',[])),
            json.dumps(n.get('numeric_conflict',{}).get('conflicts',[]),ensure_ascii=False)])
    detail=book.create_sheet('Raw ACCEPT before-after')
    detail.append(['candidate_id','Claim','Before','After','Old blockers','New reasons'])
    for r in before_after:
        detail.append([r['candidate_id'],r['raw_claim'],r['before_final'],r['final_verdict'],
                       '; '.join(r['old_blockers']),'; '.join(r['reason'])])
    neg=book.create_sheet('Raw NOT_CHANGE before-after')
    neg.append(['candidate_id','Raw negative claim','Scope','Before','After','Reasons','Graphic','Witness','Basis'])
    for r in negatives:
        neg.append([r['candidate_id'],r['raw_negative_claim'],json.dumps(r['scope'],ensure_ascii=False),r['V2_final_verdict'],r['final_verdict'],
            '; '.join(r['why']),r['graphic_applicability']['applicability'],r['witness_applicability']['applicability'],r['calculation_basis_applicability']['applicability']])
    for ws in book:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for i in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(i)].width=32 if i!=2 else 80
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='244062')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')): cell.data_type='s'
    path=OUT/'PAIR_A_REPAIRED_V3_SYSTEM_OUTPUT.xlsx'
    if path.exists(): raise ValueError('Workbook already exists')
    book.save(path)
    check=load_workbook(path,read_only=True)
    assert check['System output'].max_row==60 and check['Raw ACCEPT before-after'].max_row==14
    assert check['Raw NOT_CHANGE before-after'].max_row==6
    check.close()


def main():
    if (OUT/'REPAIR_CODE_FREEZE.json').exists(): raise ValueError('Replay already started/frozen; no automatic rerun')
    integrity=verify_inputs()
    offline_guard()
    pair=admitted_pairs('DEV',indices=[2])[0]
    assert pair['pair_key']=='ad0a31a342a666082f2ef66a'
    receipt=tests()
    numeric=read(V2/'DELIVERED_NUMERIC_FACTS.json')
    files=list(Path(__file__).parent.glob('*.py'))
    files += list((REPO/'experiments/project_change_post_inference_repair_272').glob('*.py'))
    for module in ['project_change_post_inference_repair_v2_272','project_change_f2_binding_v4_272','project_change_contracts_272','project_change_contracts_v2_272','project_change_semantic_codex_v2_272']:
        files += list((REPO/'experiments'/module).glob('*.py'))
    save('REPAIR_CODE_FREEZE.json',dict(at=now(),code={str(p.relative_to(REPO)):sha(p) for p in files},
        source_audit_opened=False,model_calls=0,tests=receipt,input_integrity=integrity,
        numeric_facts_sha256=sha(V2/'DELIVERED_NUMERIC_FACTS.json'),
        ledger_sha256=sha(OUT/'RAW_ACCEPT_REMAINING_4_LEDGER.json')))
    save('WITNESS_RULES_V3.json',dict(rules=RULES,source_truth_used=False))
    save('NEGATIVE_VERDICT_CONTRACT.json',NEGATIVE_CONTRACT)
    previous={r['package_id']:r for r in read(V2/'PAIR_A_REPAIRED_V2_RESULTS.json')}
    rows,events,evidence,ba=[],[],{},[]
    audits={name:[] for name in ['WITNESS_APPLICABILITY','GRAPHIC_APPLICABILITY','CALCULATION_BASIS','OTHER_ROLE','NUMERIC_CONFLICT']}
    for item in read(LIVE/'CALL_PLAN.json')['packages']:
        key=item['key']; p=read(BASE/item['package']); packet=deepcopy(p['evidence_packet'])
        assert p['provenance']['pair_key']==pair['pair_key']
        for side in ('old','new'):
            for e in packet['evidence'][side]:
                assert e['source_receipt']==pair[side]['artifacts']['pdf']
                if e.get('raster'):
                    e['raster']['path']=str(BASE/e['raster']['path'])
                    assert sha(e['raster']['path'])==e['raster']['sha256']
                evidence[e['evidence_id']]=e
        if item['action']=='MODEL_CALL':
            raw=read(LIVE/f'raw_responses/{key}.json')
            facts=[f for f in numeric['facts'] if f['candidate_id']==key]
            for f in facts:
                e=evidence[f['evidence_id']]
                assert f['document_version']==e['document_version'] and f['source_receipt']==e['source_receipt']
                assert f['raster']['sha256']==e['raster']['sha256']
                for o in f['outputs']: assert sha(o['path'])==o['sha256']
            n=repair(raw,packet,contract(p['candidate_subject'],packet),facts)
        else:
            raw={}; n=deepcopy(previous[key]['normalized'])
        save(f'normalized_responses/{key}.json',n)
        bound={s:sorted({b['provenance']['page'] for b in n.get(s+'_state',{}).get('evidence_bindings',[])}) for s in ('old','new')}
        row=dict(package_id=key,raw=raw,normalized=n,status=n['effective_verdict'],f2_verdict=f2_verdict(n),
            evidence_pages=bound,source_label=p['candidate_subject']['subject'],completeness=item['completeness'],
            binding_status='NO_CALL_MISSING' if not raw else 'BOUND_WITH_REJECTED_OPTIONAL_SUPPORT' if n.get('optional_rejected_supporting_references') else
                           'REVIEW' if n.get('rejected_evidence_bindings') else 'BOUND')
        rows.append(row)
        if raw:
            a=n.get('claim_applicability') or {}
            for label,field in [('WITNESS_APPLICABILITY','supporting_witness'),('GRAPHIC_APPLICABILITY','graphic'),
                                ('CALCULATION_BASIS','calculation_basis'),('OTHER_ROLE','other_role')]:
                audits[label].append(dict(candidate_id=key,raw_verdict=raw['verdict'],final_verdict=row['status'],
                    decision=a.get(field,{}),primary_evidence=a.get('primary_evidence'),
                    optional_rejected_supporting_references=n.get('optional_rejected_supporting_references',[])))
            audits['NUMERIC_CONFLICT'].append(dict(candidate_id=key,final_verdict=row['status'],decision=n.get('numeric_conflict',{})))
        if raw.get('verdict')=='ACCEPT':
            ba.append(dict(candidate_id=key,raw_claim=raw['project_change_summary'],old_blockers=reasons(previous[key]['normalized']),
                before_final=previous[key]['status'],new_applicability=n.get('claim_applicability'),
                evidence_requirement={k:(n.get('claim_applicability') or {}).get(k) for k in ['primary_evidence','supporting_witness','graphic','calculation_basis']},
                numeric_conflict=n.get('numeric_conflict'),f2_verdict=row['f2_verdict'],final_verdict=row['status'],reason=reasons(n),why='; '.join(reasons(n)) if reasons(n) else
                    'All claim dependencies satisfied by bound primary evidence and existing component proof; optional rejected references remain rejected.'))
        if row['status']=='ACCEPT':
            ids=sorted({eid for s in ('old','new') for eid in n[s+'_state']['evidence_ids']})
            events.append(dict(event_id=key,object_id=272,status='ACCEPT',engineering_subject=raw['new_state']['engineering_subject'],
                summary_ru=raw['project_change_summary'],old_state=n['old_state'],new_state=n['new_state'],evidence_ids=ids))
    assert len(rows)==59 and sum(bool(r['raw']) for r in rows)==52 and len(ba)==13
    grouped=resolve_ownership(events,[],evidence)
    gids={i:'PC_'+str(index).zfill(3) for index,g in enumerate(grouped['groups'],1) for i in g['member_ids']}
    save('PROJECT_CHANGES_REPLAY.json',dict(**grouped,group_ids=gids,implementation='resolve_ownership unchanged; no new pairwise proofs'))
    for name,data in audits.items(): save(name+'_AUDIT.json',dict(rows=data,source_truth_used=False))
    save('RAW_ACCEPT_13_BEFORE_AFTER.json',ba)
    save('PAIR_A_REPAIRED_V3_RESULTS.json',rows)
    previous_accept={k for k,v in previous.items() if v['status']=='ACCEPT'}
    structural=dict(
        nine_v2_accepts_preserved=len(previous_accept)==9 and previous_accept <= {r['package_id'] for r in rows if r['status']=='ACCEPT'},
        witness_required_has_reason=all(bool(r['normalized']['claim_applicability']['supporting_witness']['witness_required_reason']) for r in rows if r['raw'] and r['normalized']['claim_applicability']['supporting_witness']['required']),
        negative_scope_local=all(not r['normalized']['claim_applicability']['NEGATIVE_SCOPE']['whole_page_unchanged'] for r in rows if r['raw'].get('verdict')=='NOT_CHANGE'),
        no_new_binding=all(r['normalized'][side+'_state']['evidence_bindings']==previous[r['package_id']]['normalized'][side+'_state']['evidence_bindings'] for r in rows if r['raw'] for side in ('old','new')), 
        packages_59=len(rows)==59, saved_responses_52=sum(bool(r['raw']) for r in rows)==52,
        raw_accept_13=len(ba)==13, no_call_missing_7=sum(not r['raw'] for r in rows)==7,
        binding_subject_mismatch_zero=not any('BINDING_SUBJECT_MISMATCH' in r['normalized'].get('issues',[]) for r in rows),
        exact_conflicts_not_accepted=all(r['status']!='ACCEPT' for r in rows if r['normalized'].get('numeric_conflict',{}).get('blocking')),
        required_missing_not_accepted=all(r['status']!='ACCEPT' for r in rows if any(x in r['normalized'].get('issues',[]) for x in
            ['GRAPHIC_REQUIRED_NOT_DELIVERED','CALCULATION_BASIS_REQUIRED_UNKNOWN','AMBIGUOUS_OTHER_ROLE'])),
        values_unchanged=all(r['raw'][s+'_state']['value']==r['normalized'][s+'_state']['value'] for r in rows if r['raw'] and r['normalized'].get('old_state') for s in ('old','new')),
        group_members_exact=set(gids)=={e['event_id'] for e in events},
        accepted_primary_bound=all(r['normalized'].get('claim_applicability',{}).get('primary_evidence',{}).get('complete') for r in rows if r['status']=='ACCEPT'),
    )
    save('STRUCTURAL_AUDIT.json',dict(status='PASS' if all(structural.values()) else 'FAIL',checks=structural))
    counts=dict(Counter(r['status'] for r in rows))
    save('MECHANICAL_REPLAY_52.json',dict(at=now(),model_calls=0,codex_astra_calls=0,openrouter_calls=0,claude_calls=0,
        replayed=52,no_call_missing=7,raw_counts=dict(Counter(r['raw']['verdict'] for r in rows if r['raw'])),
        final_counts=counts,f2_counts=dict(Counter(r['f2_verdict'] for r in rows)),project_changes=len(grouped['groups']),
        binding_subject_mismatch=0 if structural['binding_subject_mismatch_zero'] else 'FAIL',
        source_audit_opened=False,validation='NOT OPENED',final_holdout='NOT OPENED',pair_b='NOT RUN',production='UNCHANGED',
        local_ocr='No new OCR. Reused V2 hash-checked numeric facts from identical delivered raster crops',
        rows=[dict(candidate_id=r['package_id'],raw_verdict=r['raw']['verdict'],f2_verdict=r['f2_verdict'],final_verdict=r['status']) for r in rows if r['raw']]))
    four=[x for x in ba if x['before_final']=='REVIEW']
    assert len(four)==4
    save('RAW_ACCEPT_4_BEFORE_AFTER.json',four)
    negatives=[]
    for r in rows:
        if r['raw'].get('verdict')!='NOT_CHANGE': continue
        a=r['normalized']['claim_applicability']
        negatives.append(dict(candidate_id=r['package_id'],raw_negative_claim=r['raw']['project_change_summary'],
            scope=a['NEGATIVE_SCOPE'], evidence={side:dict(state=r['raw'][side+'_state'],bindings=r['normalized'][side+'_state']['evidence_bindings']) for side in ('old','new')},
            V2_blocker=reasons(previous[r['package_id']]['normalized']), V2_final_verdict=previous[r['package_id']]['status'],
            graphic_applicability=a['graphic'],witness_applicability=a['supporting_witness'],calculation_basis_applicability=a['calculation_basis'],
            final_verdict=r['status'],why=reasons(r['normalized'])))
    assert len(negatives)==5
    save('RAW_NOT_CHANGE_5_BEFORE_AFTER.json',negatives)
    save('POST_INFERENCE_REGRESSION_AUDIT.json',dict(status='PASS' if all(structural.values()) else 'FAIL',checks=structural,
        previous_accept_ids=sorted(previous_accept), remaining_raw_accept_review=[r['candidate_id'] for r in four if r['final_verdict']=='REVIEW'],
        truth_evaluation='PENDING_POST_REPLAY',raw_hashes=verify_inputs()))
    workbook(rows,gids,ba,negatives)
    save('REPLAY_RESULT_FREEZE.json',dict(at=now(),source_audit_opened=False,model_calls=0,
        files={str(p.relative_to(OUT)):sha(p) for p in sorted(OUT.rglob('*')) if p.is_file()}))
    save('INPUT_INTEGRITY_AFTER_REPLAY.json',verify_inputs())
    print(json.dumps(dict(counts=counts,structural=structural,tests=receipt),ensure_ascii=False))


if __name__=='__main__': main()
