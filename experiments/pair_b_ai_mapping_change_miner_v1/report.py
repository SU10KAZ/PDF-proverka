"""Render frozen experiment/evaluation artifacts; never invoke a model."""
from collections import Counter
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .run import OUT, read, write, sha, now, verify_freeze


def workbook(tabs):
    book = Workbook()
    book.remove(book.active)
    for name, rows in tabs.items():
        sheet = book.create_sheet(name[:31])
        if not rows:
            rows = [{'status':'NO_ROWS'}]
        keys = list(dict.fromkeys(k for r in rows for k in r))
        sheet.append(keys)
        for row in rows:
            vals = []
            for key in keys:
                v = row.get(key)
                if isinstance(v, (dict,list)):
                    v = json.dumps(v,ensure_ascii=False)
                if isinstance(v,str):
                    v = v[:32760]
                    if v.startswith(('=','+','-','@')):
                        v = "'" + v
                vals.append(v)
            sheet.append(vals)
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24476A')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(70,max(18,len(str(col[0].value))+2))
            for cell in col:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
    target=OUT/'RESULTS.xlsx'
    if target.exists():
        raise FileExistsError(target)
    book.save(target)


def stopped():
    stops=[read(p) for p in sorted(OUT.glob('STOP_*.json'))]
    assert stops
    calls=list((OUT/'raw').glob('*/INVOCATION.json'))
    receipts=[read(p) for p in sorted((OUT/'raw').glob('*/RECEIPT.json'))]
    reason=stops[-1]['error']
    mapping_stats={}
    map_freeze=OUT/'DOCUMENT_MAP_FREEZE.json'
    if map_freeze.exists() and read(map_freeze).get('hashes'):
        verify_freeze('DOCUMENT_MAP_FREEZE.json')
        mapping_stats=read(map_freeze)['stats']
    partial=[read(p) for p in sorted((OUT/'raw').glob('PASS_B_*/parsed.json'))]
    status=dict(status='NOT_COMPLETED', reason=reason, model='gpt-6-astra', reasoning='xhigh',
        model_calls=len(calls), completed_calls=sum(bool(r['usage']) for r in receipts),
        openrouter=0, claude=0, evaluation='NOT_RUN', proven10='NOT_OPENED', f13='NOT_OPENED',
        production='UNCHANGED', validation='NOT OPENED', final_holdout='NOT OPENED',
        retries=0, repairs=0, mapping_stats=mapping_stats,
        completed_mining_groups=len(partial),
        unfrozen_partial_concrete_changes=sum(len(r['concrete_changes']) for r in partial),
        unfrozen_partial_hints=sum(len(r['unresolved_hints']) for r in partial))
    write(OUT/'RUN_STATUS.json', status)
    for name in ['DOCUMENT_MAP.json','DOCUMENT_MAP_FREEZE.json','CHANGE_MINER_RESULTS.json',
                 'CHANGE_MINER_FREEZE.json','PROVEN10_EVALUATION.json','F13_CHECK.json',
                 'FALSE_POSITIVE_AUDIT.json']:
        if not (OUT/name).exists():
            write(OUT/name,dict(status='NOT_CREATED' if 'FREEZE' in name else 'NOT_RUN',
                               reason=reason, artifact_is_placeholder=True))
    write(OUT/'BASELINE_VS_AI_MAPPING.json',dict(baseline_final_accept=0, baseline_proven_total=10,
        baseline_source='user-supplied; previous truth not opened', ai_found=None,
        hypothesis='NOT_EVALUATED', reason=reason))
    text=f'''STATUS: NOT_COMPLETED

MODEL: gpt-6-astra / xhigh
MODEL CALLS: {len(calls)} (completed: {status['completed_calls']})
OpenRouter: 0; Claude: 0

Причина остановки: {reason}
Повторных запусков и исправления результата не было.

MAPPING GROUPS: {mapping_stats.get('mapping_groups','NOT_AVAILABLE')}
OLD PAGES MAPPED: {mapping_stats.get('old_pages_mapped','NOT_AVAILABLE')} / 108
NEW PAGES MAPPED: {mapping_stats.get('new_pages_mapped','NOT_AVAILABLE')} / 188
CONCRETE CHANGES: NOT_AVAILABLE
UNRESOLVED HINTS: NOT_AVAILABLE
COMPLETED MINING GROUPS: {len(partial)}
UNFROZEN PARTIAL OUTPUT: {status['unfrozen_partial_concrete_changes']} changes / {status['unfrozen_partial_hints']} hints
PROVEN10: NOT_EVALUATED (NOT_OPENED)
F13: NOT_EVALUATED (NOT_OPENED)
FALSE CHANGES: NOT_EVALUATED
CURRENT PIPELINE PROVEN FOUND: 0/10 (из задания пользователя)
AI-MAPPING PIPELINE FOUND: NOT_EVALUATED
HYPOTHESIS: NOT_EVALUATED

ARCHITECTURE CONCLUSION:
Эксперимент не завершён, поэтому вывод о гипотезе сделать нельзя.
Источник остановки сохранён в raw и STOP; неподтверждённые результаты не посчитаны нулями.

PRODUCTION: UNCHANGED
VALIDATION: NOT OPENED
FINAL HOLDOUT: NOT OPENED

Подготовлено 294 разрешённых страницы. OLD p4 и NEW p8 исключены по frozen split.
Это исторически DEV-known пара; опыт не является исторически слепым тестом.
Файлы с artifact_is_placeholder=true обозначают невыполненные этапы, а не реальные freeze.
'''
    write(OUT/'FINAL_REPORT.md',text)
    workbook({'Status':[status], 'Baseline':[read(OUT/'BASELINE_VS_AI_MAPPING.json')],
              'Calls':[dict(path=str(p),**read(p)) for p in calls], 'Stop':stops})
    write(OUT/'DELIVERY_MANIFEST.json',dict(at=now(),hashes={p.name:sha(p) for p in OUT.iterdir() if p.is_file()}))


def complete():
    verify_freeze('CHANGE_MINER_FREEZE.json')
    mapping=read(OUT/'DOCUMENT_MAP.json')
    stats=read(OUT/'DOCUMENT_MAP_FREEZE.json')['stats']
    results=read(OUT/'CHANGE_MINER_RESULTS.json')
    evaluation=read(OUT/'PROVEN10_EVALUATION.json')
    audit=read(OUT/'FALSE_POSITIVE_AUDIT.json')
    f13=read(OUT/'F13_CHECK.json')
    conclusion=read(OUT/'EVALUATION_CONCLUSION.json')
    mapping_quality=read(OUT/'PAGE_MAPPING_QUALITY.json')
    coverage=read(OUT/'MAPPING_COVERAGE.json')
    counts=dict(Counter(r['outcome'] for r in evaluation['rows']))
    assert len(evaluation['rows'])==10
    counts={k:counts.get(k,0) for k in ['STRONG','PARTIAL','MISSED']}
    fp_counts=Counter(r['audit_status'] for r in audit['rows'])
    fp={k:fp_counts[k] for k in ['correct','partial','false','insufficient to judge']}
    audited={r['change_id'] for r in audit['rows']}
    assert all(c['change_id'] in audited for c in results['concrete_changes'] if c['confidence']>=0.85)
    usage=[]
    for path in sorted((OUT/'raw').glob('*/RECEIPT.json')):
        record=read(path)
        assert record['exit_code']==0 and record['tool_items']==0
        usage.append(dict(call=path.parent.name,**record))
    assert len(usage)==1+len(mapping['groups'])
    total_tokens=Counter()
    for record in usage:
        for u in record['usage']:
            for k,v in u.items():
                if isinstance(v,(int,float)):
                    total_tokens[k]+=v
    summary=dict(status='COMPLETED',model='gpt-6-astra',reasoning='xhigh',model_calls=len(usage),
        openrouter=0,claude=0,**stats,concrete_changes=len(results['concrete_changes']),
        unresolved_hints=len(results['unresolved_hints']),proven10=counts,f13=f13['status'],
        false_positive_audit=fp,false_changes=fp.get('false',0),
        unaudited_changes=len(results['concrete_changes'])-len(audited),
        baseline_proven_found=0,ai_mapping_strong_found=counts['STRONG'],
        hypothesis=conclusion['hypothesis'],production='UNCHANGED',validation='NOT OPENED',
        final_holdout='NOT OPENED',tokens=dict(total_tokens),retries=0,repairs=0)
    write(OUT/'RUN_STATUS.json',summary)
    write(OUT/'MODEL_USAGE.json',dict(model='gpt-6-astra',reasoning='xhigh',model_calls=len(usage),
        openrouter=0,claude=0,total_tokens=dict(total_tokens),calls=usage))
    baseline=dict(current_pipeline=dict(proven10_found=0,total=10,metric='final ACCEPT',source='user request'),
        ai_mapper_change_miner=dict(proven10=counts,found=counts['STRONG'],found_definition='STRONG only',
            total=10,f13=f13['status'],concrete_changes=summary['concrete_changes'],
            unresolved_hints=summary['unresolved_hints'],false_positive_audit=fp),
        hypothesis=conclusion['hypothesis'],caveat='One historically DEV-known pair; differing output gates; no reserve/generalization claim')
    write(OUT/'BASELINE_VS_AI_MAPPING.json',baseline)
    report=f'''STATUS: COMPLETED

MODEL: gpt-6-astra / xhigh
MODEL CALLS: {len(usage)}
OpenRouter: 0; Claude: 0

MAPPING GROUPS: {stats['mapping_groups']}
OLD PAGES MAPPED: {stats['old_pages_mapped']} / 108
NEW PAGES MAPPED: {stats['new_pages_mapped']} / 188
UNMATCHED OLD: {mapping['unmatched_old']}
UNMATCHED NEW: {mapping['unmatched_new']}
1→1: {coverage['cardinalities'].get('1→1',0)}; 1→N: {len(stats['one_to_many'])}; N→1: {len(stats['many_to_one'])}; N→N: {coverage['cardinalities'].get('N→N',0)}
EMBARGO: OLD p4, NEW p8

CONCRETE CHANGES: {summary['concrete_changes']}
UNRESOLVED HINTS: {summary['unresolved_hints']}

PROVEN10:

- STRONG {counts['STRONG']}
- PARTIAL {counts['PARTIAL']}
- MISSED {counts['MISSED']}

F13: {f13['status']}
FALSE CHANGES: {summary['false_changes']}
SOURCE-FIRST AUDIT: {json.dumps(fp,ensure_ascii=False)}
UNAUDITED CHANGES: {summary['unaudited_changes']}

CURRENT PIPELINE PROVEN FOUND: 0/10
AI-MAPPING PIPELINE FOUND: {counts['STRONG']}/10 (STRONG; PARTIAL отдельно)
HYPOTHESIS: {conclusion['hypothesis']}

ARCHITECTURE CONCLUSION:
{conclusion['architecture_conclusion']}

PRODUCTION: UNCHANGED
VALIDATION: NOT OPENED
FINAL HOLDOUT: NOT OPENED

Ограничения: одна исторически DEV-known пара, два embargo-листа; покрытие
страниц само по себе не доказывает правильность соответствий. Сопоставление
0 final ACCEPT с source-audited STRONG показывает результат этого опыта,
но не эквивалентность критериев приёмки двух pipeline. После freeze результат,
карта и промпты не исправлялись; повторных запросов не было.

PROVEN10 по находкам:

| Находка | Оценка | Изменения | Обоснование |
|---|---|---|---|
'''
    report+='\n'.join('| '+ ' | '.join([r['finding_id'],r['outcome'],', '.join(r['found_change_ids']),r['rationale']])+' |' for r in evaluation['rows'])
    report+='''

F13 — PASS с ограничением экспозиции: NEW p4 с ложной фразой «отсутствовали
системы» был в snippets mapper, но остался unmatched и не попал в miner.
В final concrete changes нет ложного «не было → появилось». Это проверка
конечного выхода, а не доказанная способность miner отвергать эту фразу
при прямом предъявлении.

Качество mapping: достаточное соответствие для всех 10 benchmark-кейсов;
для F10 использованы эквивалентные страницы NEW66/68 вместо предложенной NEW22.
Общая семантическая точность всех 35 групп не оценена. NEW187 передавалась
без шапки NEW186 в группах G026–G031 и G033; данные мощности нагревателей
остались hints. В G034/G035 не доказана преемственность старых и новых
помещений ТШ. Повторное использование страниц привело к дубликатам:
G007-C01/G008_C04/G010-C01 и перекрытию G017-C04 с G018-C02/C03.
62 карточки не равны 62 уникальным инженерным событиям.

Source-first audit охватил все 62 изменения, включая все 61 с confidence ≥0,85.
G008_C01 — PARTIAL: общий перенос лестничных вентиляторов с подземного уровня
на кровлю чрезмерно обобщён; OLD105/106 уже показывают часть вентиляторов
на кровле. G035-C01/G035-C02 — INSUFFICIENT TO JUDGE: различия показаны,
но тождество старых и новых тамбуров не доказано. Полностью ложных событий
не установлено; это не означает, что все 62 карточки пригодны к приёмке.

Аудит выполнен оркестратором по исходным PDF, растру, native/OCR и проверяемым
фактическим якорям; отдельного запроса модели на самооценку не было.
Это не независимая экспертиза человеком и не повторный нормативный расчёт.
Проверялись документированные события и ключевые состояния/параметры,
а не соответствие построенного объекта документации.

В таблицах OLD10–13 есть расхождение шифра в штампе
22-0121-ОК-1/Н-1.2-ИОС4.1ТЧ. Они сохранены как страницы разрешённого
hash-verified v002 PDF; выводы относятся к этому baseline. Другой документ
для подмены источника не открывался.

Mapping freeze: '''+read(OUT/'DOCUMENT_MAP_FREEZE.json')['frozen_at']+'''
Result freeze: '''+read(OUT/'CHANGE_MINER_FREEZE.json')['frozen_at']+'''
Первый доступ к evaluation разрешён после проверки freeze: '''+read(OUT/'EVALUATION_ACCESS_RECEIPT.json')['at']+'''

Файлы: RESULTS.xlsx; DOCUMENT_MAP.json; DOCUMENT_MAP_FREEZE.json;
CHANGE_MINER_RESULTS.json; CHANGE_MINER_FREEZE.json; PROVEN10_EVALUATION.json;
F13_CHECK.json; FALSE_POSITIVE_AUDIT.json; BASELINE_VS_AI_MAPPING.json.
Дополнительная трассировка: FACT_ASSERTIONS.json, audit_material/,
audit_visual/, PAGE_MAPPING_QUALITY.json, raw/, DELIVERY_MANIFEST.json.
'''
    write(OUT/'FINAL_REPORT.md',report)
    workbook({'Summary':[summary],'Mapping':mapping['groups'],
        'Concrete changes':results['concrete_changes'],'Unresolved hints':results['unresolved_hints'],
        'PROVEN10':evaluation['rows'],'F13':[f13],'Source audit':audit['rows'],
        'Mapping quality':[mapping_quality],'Baseline':[baseline],'Model calls':usage})
    verify_freeze('CHANGE_MINER_FREEZE.json')
    write(OUT/'DELIVERY_MANIFEST.json',dict(at=now(),hashes={p.name:sha(p) for p in OUT.iterdir() if p.is_file()}))


if __name__=='__main__':
    import sys
    complete() if sys.argv[1:] == ['complete'] else stopped()
