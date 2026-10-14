"""Состав Excel GPT (EM2_реестр_сравнения.xlsx): листы, пояснения, заполненность колонок «Листы».

Запуск: python -I xlsx_probe.py <путь к xlsx>. См. docs/version_comparison_plan.md §15, §16.
"""
import sys
from collections import Counter

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
for ws in wb.worksheets:
    print(f"=== {ws.title}: rows={ws.max_row} cols={ws.max_column}; note: {str(ws['A2'].value)[:160]}")
ws = wb["Листы"]
rows = list(ws.iter_rows(min_row=6, values_only=True))
hdr = [c.value for c in ws[5]]
for i, h in enumerate(hdr):
    vals = [r[i] for r in rows]
    filled = sum(v not in (None, "") for v in vals)
    c = Counter(str(v)[:40] for v in vals if v not in (None, ""))
    print(f"  {h}: filled {filled}/{len(rows)}; top: {c.most_common(4) if len(c) < 60 else str(len(c)) + ' distinct'}")
