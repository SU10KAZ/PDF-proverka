"""Post-freeze evaluation for the optimized Sol/high screening."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys


ROOT = Path("/home/coder/auditmanager/corpus-audits/20260914_project_change_272")
OUT = ROOT / "cheap_model_screening_v3_sol_optimized"
V1 = ROOT / "cheap_model_screening_v1"
PAIR_A = ROOT / "pair_a_ai_mapping_change_miner_v1"
PAIR_B = ROOT / "pair_b_ai_mapping_change_miner_v1"
DEPS = ROOT / "controlled_inference_f1_f4_f2_v3/runtime_deps"
sys.path.insert(0, str(DEPS))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def now():
    return datetime.now(timezone.utc).isoformat()


def read(path):
    return json.loads(Path(path).read_text())


def write_new(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify_freeze():
    integrity = read(OUT / "SAMPLE_INTEGRITY.json")
    freeze = read(OUT / "SOL_OPTIMIZED_RESULT_FREEZE.json")
    if integrity["status"] != "PASS" or not integrity["only_prompt_changed"]:
        raise RuntimeError("Sample integrity is not PASS")
    if freeze["model_calls"] != 8 or freeze["successful_calls"] != 8 or freeze["no_truth_leakage"] != "PASS":
        raise RuntimeError("Optimized result freeze is incomplete")
    for relative, digest in freeze["hashes"].items():
        if sha(OUT / relative) != digest:
            raise RuntimeError(f"Frozen artifact drift: {relative}")
    return integrity, freeze


ASSESSMENTS = {
    "A-REAL-08": ("PARTIAL", ["G008-C01"],
        "Общая площадь квартиры 1.3.47 сохранена, но ключевое перераспределение площади между холлом и кухней-столовой и повторяемость события не восстановлены."),
    "A-REAL-10": ("PARTIAL", ["G008-C02"],
        "Общая площадь квартиры 2.2.18 сохранена, но перераспределение кухни и холлов и повторяемость решения не восстановлены."),
    "F03": ("MISSED", [],
        "Переход от двух раздельных вытяжек к общей шахте/показанному вентилятору не выдан; общий topology hint не называет доказанную reference-суть."),
    "F04": ("PARTIAL", ["G017-H02"],
        "Старые установки 12000 и новый подбор 23600, давление и мощность найдены, но оставлены unresolved из-за групповой identity и не опубликованы как ProjectChange."),
    "F09": ("MISSED", [],
        "Обе системы рампы и рабочие точки 50990→21600 и 24000→10900 отсутствуют в concrete changes и hints."),
    "F10": ("PARTIAL", ["G017-C01"],
        "Изменение геометрических и температурных расчётных условий сохранено, но переход нормативной основы СП→АВОК не назван."),
    "G025-C01": ("CORRECT", ["G025-C01"],
        "Переподбор ППП сохранён: тип, расход, давление, мощность и обороты подтверждены обеими сторонами."),
    "G025-C02": ("CORRECT", ["G025-C03"],
        "Переподбор ПО сохранён: тип, расход, давление, мощность и обороты подтверждены обеими сторонами."),
    "G025-C03": ("CORRECT", ["G025-C02", "G025-C04"],
        "Перенос обоих вентиляторов из подземной венткамеры на кровлю и изменение вводов восстановлены по GRAPHIC+TABLE evidence."),
    "G025-C04": ("MISSED", [],
        "Размеры дверей, температурные допущения, модель утечек и переход к единому итоговому расходу не сформированы как отдельное событие."),
    "G025-C05": ("MISSED", [],
        "Изменение высотной модели и расширение расчёта до 16-го этажа отсутствуют."),
}


GROUP_ROWS = [
    {"pair": "A", "group_id": "G002", "quality": "SAME_OR_BETTER", "astra_changes": 0, "optimized_changes": 0,
     "note": "Metadata-only различия корректно исключены; прежние два Sol false устранены."},
    {"pair": "A", "group_id": "G008", "quality": "ACCEPTABLE_LOSS", "astra_changes": 20, "optimized_changes": 6,
     "note": "Инженерные изменения площадей сохранены без false, но обе reference-перепланировки retained лишь частично."},
    {"pair": "A", "group_id": "G024", "quality": "SAME_OR_BETTER", "astra_changes": 0, "optimized_changes": 0,
     "note": "Пустой результат совпадает с Astra; OCR-артефакты не приняты за изменение."},
    {"pair": "A", "group_id": "G031", "quality": "SAME_OR_BETTER", "astra_changes": 0, "optimized_changes": 0,
     "note": "Concrete changes отсутствуют; недоказанное различие корректно оставлено hint."},
    {"pair": "B", "group_id": "G004", "quality": "SAME_OR_BETTER", "astra_changes": 0, "optimized_changes": 0,
     "note": "Пустой результат совпадает с Astra."},
    {"pair": "B", "group_id": "G017", "quality": "MAJOR_LOSS", "astra_changes": 4, "optimized_changes": 1,
     "note": "F03 и F09 пропущены, F04 понижен до hint, F10 сохранён частично."},
    {"pair": "B", "group_id": "G025", "quality": "MAJOR_LOSS", "astra_changes": 5, "optimized_changes": 4,
     "note": "Два переподбора и перенос сохранены, но аэродинамическая и высотная модели пропущены."},
    {"pair": "B", "group_id": "G035", "quality": "SAME_OR_BETTER", "astra_changes": 2, "optimized_changes": 1,
     "note": "Основное событие укрупнено; reference-denominator исключает frozen insufficient-to-judge claims."},
]


GRAPHIC_REFS = ["F03", "F04", "G025-C03"]
CROSS_MODAL_REFS = ["F03", "F04", "F10", "G025-C03", "G025-C04", "G025-C05"]


def build_source_comparison():
    source = read(V1 / "SOURCE_FIRST_COMPARISON.json")
    rows = []
    for reference in source["rows"]:
        outcome, ids, rationale = ASSESSMENTS[reference["reference_id"]]
        rows.append({**reference, "optimized_outcome": outcome,
                     "optimized_ids": ids, "optimized_rationale": rationale})
    counts = Counter(row["optimized_outcome"] for row in rows)
    value = {
        "created_at": now(), "evaluation_opened_after_verified_optimized_freeze": True,
        "reference_definition": source["reference_definition"], "reference_changes": len(rows),
        "rows": rows, "optimized_reference_counts": dict(counts),
        "optimized_reference_retained": sum(counts[k] for k in ["CORRECT", "PARTIAL"]),
        "optimized_false_changes": [], "optimized_false_count": 0,
        "optimized_output_counts": {"CONCRETE": 12, "UNRESOLVED_HINT": 8,
                                    "FALSE": 0, "INSUFFICIENT_TO_JUDGE": 1},
        "output_count_note": "Reference-event scoring is primary. G035 output remains insufficient to judge and is not called false.",
    }
    write_new(OUT / "SOURCE_FIRST_COMPARISON.json", value)
    return value


def build_comparison(source):
    quality = Counter(row["quality"] for row in GROUP_ROWS)
    models = [
        {"configuration": "ASTRA XHIGH", "model": "gpt-6-astra/xhigh", "correct": 10, "partial": 1,
         "missed": 0, "false": 0, "input_plus_output_tokens": 1146959},
        {"configuration": "TERRA HIGH", "model": "gpt-5.6-terra/high", "correct": 6, "partial": 3,
         "missed": 2, "false": 3, "input_plus_output_tokens": 694471},
        {"configuration": "SOL HIGH ORIGINAL", "model": "gpt-5.6-sol/high", "correct": 8, "partial": 2,
         "missed": 1, "false": 2, "input_plus_output_tokens": 699592},
        {"configuration": "SOL HIGH OPTIMIZED", "model": "gpt-5.6-sol/high", "correct": 3, "partial": 4,
         "missed": 4, "false": 0, "input_plus_output_tokens": None},
    ]
    value = {"created_at": now(), "same_frozen_sample": True, "reference_changes": 11,
             "configurations": models, "group_rows": GROUP_ROWS,
             "optimized_vs_astra_quality_counts": dict(quality),
             "conclusion": "Prompt eliminated metadata false positives but over-constrained evidence formation and materially reduced reference recall."}
    write_new(OUT / "ASTRA_TERRA_SOL_COMPARISON.json", value)
    return value


def usage_from_receipts():
    rows = []
    keys = ["input_tokens", "cached_input_tokens", "output_tokens", "reasoning_output_tokens"]
    for path in sorted((OUT / "sol_raw").glob("*/RECEIPT.json")):
        receipt = read(path)
        usage = receipt["usage"][0]
        rows.append({"call": receipt["sample_id"], **{key: usage.get(key, 0) for key in keys}})
    total = {key: sum(row[key] for row in rows) for key in keys}
    total["input_plus_output_tokens"] = total["input_tokens"] + total["output_tokens"]
    total["average_input_plus_output_per_group"] = total["input_plus_output_tokens"] / len(rows)
    return rows, total


def build_usage(comparison):
    rows, sol = usage_from_receipts()
    saving = (1 - sol["input_plus_output_tokens"] / 1146959) * 100
    original_delta = sol["input_plus_output_tokens"] - 699592
    value = {
        "created_at": now(), "model": "gpt-5.6-sol", "reasoning": "high", "calls": 8,
        "optimized": sol, "optimized_by_call": rows,
        "astra_same_sample": {"input_plus_output_tokens": 1146959},
        "terra_same_sample": {"input_plus_output_tokens": 694471},
        "sol_original_same_sample": {"input_plus_output_tokens": 699592},
        "saving_vs_astra_percent": round(saving, 1),
        "increase_vs_sol_original_tokens": original_delta,
        "increase_vs_sol_original_percent": round(original_delta / 699592 * 100, 1),
        "cached_input_note": "cached_input_tokens is included within input_tokens and is not added again.",
    }
    comparison["configurations"][-1]["input_plus_output_tokens"] = sol["input_plus_output_tokens"]
    Path(OUT / "ASTRA_TERRA_SOL_COMPARISON.json").write_text(json.dumps(comparison, ensure_ascii=False, indent=2) + "\n")
    write_new(OUT / "TOKEN_USAGE.json", value)
    return value


def build_projection(usage):
    full_a = read(PAIR_A / "TOKEN_USAGE.json")["total"]
    full_b = read(PAIR_B / "MODEL_USAGE.json")["total_tokens"]
    astra_rows = read(V1 / "TOKEN_USAGE.json")["astra_by_call"]
    sol_a = sum(row["input_tokens"] + row["output_tokens"] for row in usage["optimized_by_call"] if row["call"].startswith("PAIR_A_"))
    sol_b = sum(row["input_tokens"] + row["output_tokens"] for row in usage["optimized_by_call"] if row["call"].startswith("PAIR_B_"))
    astra_a = sum(row["input_tokens"] + row["output_tokens"] for row in astra_rows[:4])
    astra_b = sum(row["input_tokens"] + row["output_tokens"] for row in astra_rows[4:])
    historical_a = full_a["input_tokens"] + full_a["output_tokens"]
    historical_b = full_b["input_tokens"] + full_b["output_tokens"]
    ratios = {"A": sol_a / astra_a, "B": sol_b / astra_b}
    stratified = round(historical_a * ratios["A"] + historical_b * ratios["B"])
    pooled = round((historical_a + historical_b) * usage["optimized"]["input_plus_output_tokens"] / 1146959)
    group_scaled = round(usage["optimized"]["average_input_plus_output_per_group"] * 67)
    value = {
        "created_at": now(), "projection_only_not_fact": True,
        "scope": "Semantic Mapping V2 + ProjectChange Miner V3 for Pair A + Pair B",
        "historical_full_input_plus_output_tokens": historical_a + historical_b,
        "pair_specific_sample_ratios": ratios,
        "projected_full_v3_input_plus_output_tokens": stratified,
        "rounded_report_value": "~5.8 million tokens",
        "sensitivity": {"pooled_ratio_projection": pooled,
                        "67_group_average_projection_mining_only": group_scaled,
                        "indicative_range": [min(pooled, stratified), group_scaled]},
        "assumptions": ["Full V3 has approximately the historical Pair A/Pair B mapping+mining volume.",
                        "Sample-specific token ratios transfer to the full run.",
                        "No repair, retry, verifier, dedupe, or evaluation calls are included."],
    }
    write_new(OUT / "FULL_V3_COST_PROJECTION.json", value)
    return value


def build_error_audit(source):
    outcome = {row["reference_id"]: row["optimized_outcome"] for row in source["rows"]}
    retained = lambda ids: sum(outcome[x] in {"CORRECT", "PARTIAL"} for x in ids)
    value = {
        "created_at": now(), "labels_opened_only_after_result_freeze": True,
        "metadata_only_false_changes": {"count": 0, "groups": ["G002"],
                                        "note": "G002 returned no changes or hints and explicitly marked NOT_ENGINEERING_CHANGE."},
        "graphic_reference_events": {"definition": "Reference events whose frozen Astra evidence included GRAPHIC.",
                                     "ids": GRAPHIC_REFS, "total": len(GRAPHIC_REFS),
                                     "retained": retained(GRAPHIC_REFS),
                                     "retained_definition": "CORRECT or PARTIAL",
                                     "outcomes": {x: outcome[x] for x in GRAPHIC_REFS}},
        "cross_modal_events": {"definition": "Reference events whose frozen Astra evidence used more than one modality.",
                               "ids": CROSS_MODAL_REFS, "total": len(CROSS_MODAL_REFS),
                               "retained": retained(CROSS_MODAL_REFS),
                               "retained_definition": "CORRECT or PARTIAL",
                               "outcomes": {x: outcome[x] for x in CROSS_MODAL_REFS}},
    }
    write_new(OUT / "ERROR_CLASS_AUDIT.json", value)
    return value


def add_sheet(book, title, rows):
    sheet = book.create_sheet(title)
    if not rows:
        sheet.append(["empty"])
        return
    columns = list(rows[0])
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([json.dumps(row.get(column), ensure_ascii=False) if isinstance(row.get(column), (list, dict)) else row.get(column)
                      for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(80, max(12, max(len(str(cell.value or "")) for cell in cells) + 2))


def build_excel(source, comparison, usage, projection, errors):
    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "Summary", [{
        "status": "COMPLETED_SOL_REJECTED", "model": "gpt-5.6-sol/high", "prompt": "SOL_OPTIMIZED_V1",
        "sample": "8 identical frozen groups", "calls": 8, "reference_changes": 11,
        "correct": 3, "partial": 4, "missed": 4, "false": 0,
        "same_or_better": 5, "acceptable_loss": 1, "major_loss": 2,
        "metadata_false": 0, "graphic_total": 3, "graphic_retained": 2,
        "cross_modal_total": 6, "cross_modal_retained": 3,
        "input": usage["optimized"]["input_tokens"], "cached": usage["optimized"]["cached_input_tokens"],
        "output": usage["optimized"]["output_tokens"], "reasoning": usage["optimized"]["reasoning_output_tokens"],
        "input_plus_output": usage["optimized"]["input_plus_output_tokens"],
        "saving_vs_astra_percent": usage["saving_vs_astra_percent"],
        "projected_full_v3": projection["projected_full_v3_input_plus_output_tokens"],
        "no_truth_leakage": "PASS", "recommendation": "KEEP_ASTRA_XHIGH",
    }])
    add_sheet(book, "Reference comparison", source["rows"])
    add_sheet(book, "Configurations", comparison["configurations"])
    add_sheet(book, "Group quality", comparison["group_rows"])
    add_sheet(book, "Usage", usage["optimized_by_call"])
    add_sheet(book, "Error classes", [
        {"class": "METADATA_ONLY_FALSE", "total": 0, "retained": "n/a"},
        {"class": "GRAPHIC_REFERENCE", "total": 3, "retained": 2},
        {"class": "CROSS_MODAL_REFERENCE", "total": 6, "retained": 3},
    ])
    book.save(OUT / "RESULTS.xlsx")


def build_report(usage, projection):
    sol = usage["optimized"]
    text = f"""# ProjectChange — Sol-optimized screening V3

STATUS: COMPLETED_SOL_REJECTED

MODEL: gpt-5.6-sol / high

PROMPT: SOL_OPTIMIZED_V1

SAMPLE: 8 identical frozen groups

MODEL CALLS: 8

REFERENCE CHANGES: 11

SOL OPTIMIZED:
- correct: 3
- partial: 4
- missed: 4
- false: 0
- reference retained: 7/11

METADATA FALSE: 0

GRAPHIC EVENTS:
- total: 3
- retained: 2

CROSS-MODAL EVENTS:
- total: 6
- retained: 3

VS ASTRA:
- same/better: 5
- acceptable loss: 1
- major loss: 2

TOKEN USAGE:
- input: {sol['input_tokens']}
- cached: {sol['cached_input_tokens']} (included in input)
- output: {sol['output_tokens']}
- reasoning: {sol['reasoning_output_tokens']}
- input+output: {sol['input_plus_output_tokens']}
- average per group: {sol['average_input_plus_output_per_group']:.0f}

ASTRA: 1,146,959

TERRA: 694,471

SOL ORIGINAL: 699,592

SOL OPTIMIZED: {sol['input_plus_output_tokens']}

SAVING VS ASTRA: {usage['saving_vs_astra_percent']}%

PROJECTED FULL V3: ~5.8M tokens ({projection['projected_full_v3_input_plus_output_tokens']} stratified estimate; sensitivity {projection['sensitivity']['indicative_range'][0]}–{projection['sensitivity']['indicative_range'][1]}). Projection only.

NO TRUTH LEAKAGE: PASS

QUALITY CONCLUSION: Новый contract устранил оба metadata-only false changes, но стал слишком консервативным. F03 не восстановлен как доказанное GRAPHIC-событие, F09 полностью пропущен, F04 оставлен hint, F10 сохранён лишь частично, а в G025 потеряны аэродинамическая и высотная модели. Correct снизился с 8 до 3, missed вырос с 1 до 4, token usage вырос относительно Sol original на {usage['increase_vs_sol_original_percent']}%.

RECOMMENDATION: KEEP_ASTRA_XHIGH

Условия USE_SOL_HIGH_FOR_V3 не выполнены: missed != 0, graphic reference event retained не полностью, есть две MAJOR_LOSS groups и заметная деградация против Astra. Дополнительный prompt-tuning cycle и полный V3 автоматически не запускались.

PRODUCTION: UNCHANGED

VALIDATION: NOT OPENED

FINAL HOLDOUT: NOT OPENED

## Trace notes

- Из 112 invocation input files изменены ровно восемь prompt.txt; остальные 104 файла идентичны frozen V1 по SHA-256.
- Все восемь Sol-контекстов были свежими и изолированными: они получили только frozen source input и общий prompt; truth и прежние model results в их filesystem/context отсутствовали. Результаты hash-frozen до запуска post-freeze evaluator.
- Astra calls: 0; Terra calls: 0; OpenRouter calls: 0; Claude calls: 0; retries: 0.
"""
    (OUT / "FINAL_REPORT.md").write_text(text)


def main():
    verify_freeze()
    source = build_source_comparison()
    comparison = build_comparison(source)
    usage = build_usage(comparison)
    projection = build_projection(usage)
    errors = build_error_audit(source)
    build_excel(source, comparison, usage, projection, errors)
    build_report(usage, projection)
    required = ["SAMPLE_INTEGRITY.json", "OPTIMIZED_PROMPT.txt", "SOL_OPTIMIZED_RESULT_FREEZE.json",
                "SOURCE_FIRST_COMPARISON.json", "ASTRA_TERRA_SOL_COMPARISON.json", "ERROR_CLASS_AUDIT.json",
                "TOKEN_USAGE.json", "FULL_V3_COST_PROJECTION.json", "RESULTS.xlsx", "FINAL_REPORT.md"]
    write_new(OUT / "DELIVERY_MANIFEST.json", {
        "created_at": now(), "status": "COMPLETED_SOL_REJECTED",
        "files": {name: sha(OUT / name) for name in required},
        "sol_calls": 8, "astra_calls": 0, "terra_calls": 0, "openrouter_calls": 0,
        "claude_calls": 0, "retries": 0, "no_truth_leakage": "PASS",
        "production": "UNCHANGED", "validation": "NOT OPENED", "final_holdout": "NOT OPENED",
    })


if __name__ == "__main__":
    main()
