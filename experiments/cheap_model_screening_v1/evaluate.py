"""Post-freeze comparison and reporting for cheap model screening V1."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys


ROOT = Path("/home/coder/auditmanager/corpus-audits/20260914_project_change_272")
OUT = ROOT / "cheap_model_screening_v1"
PAIR_A = ROOT / "pair_a_ai_mapping_change_miner_v1"
PAIR_B = ROOT / "pair_b_ai_mapping_change_miner_v1"
DEPS = ROOT / "controlled_inference_f1_f4_f2_v3/runtime_deps"
sys.path.insert(0, str(DEPS))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def now():
    return datetime.now(timezone.utc).isoformat()


def read(path):
    return json.loads(Path(path).read_text())


def write_new(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify_candidate_freeze():
    freeze = read(OUT / "CANDIDATE_RESULT_FREEZE.json")
    if freeze["model_calls"] != 8 or freeze["successful_calls"] != 8:
        raise RuntimeError("Candidate result is not a complete eight-call freeze")
    if freeze["source_truth_opened_before_freeze"] or freeze["astra_outputs_opened_before_freeze"]:
        raise RuntimeError("Pre-freeze access invariant failed")
    for relative, digest in freeze["hashes"].items():
        if sha(OUT / relative) != digest:
            raise RuntimeError(f"Candidate freeze drift: {relative}")
    return freeze


REFERENCE_ROWS = [
    dict(pair="A", group_id="G008", reference_id="A-REAL-08",
         reference_change="Перераспределение холлов и кухонь квартир 1.3.37/42/47/52",
         astra_outcome="CORRECT", candidate_outcome="PARTIAL", candidate_ids=["G008-CHG-03"],
         rationale="Terra сохранила изменение общей площади 1.3.47, но не восстановила ключевое перераспределение кухни и холла и полный повторяющийся source event."),
    dict(pair="A", group_id="G008", reference_id="A-REAL-10",
         reference_change="Перераспределение кухонь и холлов квартир 2.2.18/20/22",
         astra_outcome="CORRECT", candidate_outcome="PARTIAL", candidate_ids=["G008-CHG-05"],
         rationale="Terra агрегировала рост площадей корпуса 2, но не описала инженерную суть перераспределения кухни и холлов и полный повторяющийся event."),
    dict(pair="B", group_id="G017", reference_id="F03",
         reference_change="Две раздельные вытяжки заменены общей шахтой/показанным вентилятором",
         astra_outcome="PARTIAL", candidate_outcome="MISSED", candidate_ids=[],
         rationale="Terra описала оборудование ДУ-1.1п, но не заявила переход от двух раздельных вытяжек к общей шахте; ключевая GRAPHIC-суть пропущена."),
    dict(pair="B", group_id="G017", reference_id="F04",
         reference_change="Компенсация паркинга: старые установки 12000 → новый подбор 23600 м³/ч",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G017-C02"],
         rationale="Расходы, давление, мощность и привязка паркинга −1 сохранены."),
    dict(pair="B", group_id="G017", reference_id="F09",
         reference_change="Рампа: 50990→21600 и 24000→10900 м³/ч",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G017-C01"],
         rationale="Обе системы рампы, направление и числовые состояния сохранены."),
    dict(pair="B", group_id="G017", reference_id="F10",
         reference_change="Основание расчёта: рекомендации к СП 7.13130.2013 → АВОК 5.5.1-2023",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G017-C03"],
         rationale="Нормативное основание и связанные температурно-геометрические параметры сохранены."),
    dict(pair="B", group_id="G025", reference_id="G025-C01",
         reference_change="Переподбор вентилятора подпора лифта ППП",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G025-C01"],
         rationale="Тип, расход, давление, обороты и направление изменения сохранены."),
    dict(pair="B", group_id="G025", reference_id="G025-C02",
         reference_change="Переподбор вентилятора подпора лифта ПО",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G025-C02"],
         rationale="Расход, давление, мощность и направление изменения сохранены."),
    dict(pair="B", group_id="G025", reference_id="G025-C03",
         reference_change="Перенос двух вентиляторов и вводов воздуха из подземной венткамеры на кровлю",
         astra_outcome="CORRECT", candidate_outcome="CORRECT", candidate_ids=["G025-C01", "G025-C02"],
         rationale="Событие не атомизировано: размещение корректно включено в две карточки соответствующих режимов."),
    dict(pair="B", group_id="G025", reference_id="G025-C04",
         reference_change="Переработка аэродинамической модели, дверей, температур и утечек",
         astra_outcome="CORRECT", candidate_outcome="PARTIAL", candidate_ids=["G025-C03"],
         rationale="Сохранено 20→16 °C, но потеряны размеры дверей, способ расчёта утечек и переход к единому итоговому расходу."),
    dict(pair="B", group_id="G025", reference_id="G025-C05",
         reference_change="Изменение высотной модели и расширение расчёта до 16-го этажа",
         astra_outcome="CORRECT", candidate_outcome="MISSED", candidate_ids=[],
         rationale="Terra оставила это только как unresolved hint и ошибочно сочла доказательство недостаточным, хотя frozen source-first audit подтвердил событие."),
]


FALSE_ROWS = [
    dict(pair="A", group_id="G002", candidate_id="G002-C01", subject="Разработчик проектной документации",
         rationale="Смена организации на титульном листе — административные метаданные, а не инженерное состояние проекта."),
    dict(pair="A", group_id="G002", candidate_id="G002-C02", subject="Ответственные лица за проектирование",
         rationale="Смена фамилий и ролей на титульном листе не является конкретным инженерным изменением."),
    dict(pair="A", group_id="G002", candidate_id="G002-C03", subject="Суффикс -КОРР. в обозначении комплекта",
         rationale="Суффикс документа — оформление/обозначение, прямо исключённое из инженерных changes; Astra корректно вернула пустой результат."),
]


GROUP_ROWS = [
    dict(pair="A", group_id="G002", quality="MAJOR_LOSS", astra_changes=0, candidate_changes=3,
         note="Три новых false changes из титульных метаданных при пустом Astra output."),
    dict(pair="A", group_id="G008", quality="ACCEPTABLE_LOSS", astra_changes=20, candidate_changes=5,
         note="Сильное полезное укрупнение без новых false, но обе REAL15-сущности сохранены лишь частично."),
    dict(pair="A", group_id="G024", quality="SAME_OR_BETTER", astra_changes=0, candidate_changes=0,
         note="Обе модели корректно вернули пустой результат."),
    dict(pair="A", group_id="G031", quality="SAME_OR_BETTER", astra_changes=0, candidate_changes=0,
         note="Обе модели корректно вернули пустой результат."),
    dict(pair="B", group_id="G004", quality="SAME_OR_BETTER", astra_changes=0, candidate_changes=0,
         note="Обе модели вернули пустой результат; новых false нет."),
    dict(pair="B", group_id="G017", quality="MAJOR_LOSS", astra_changes=4, candidate_changes=3,
         note="Сохранены F04/F09/F10, но полностью потеряна доказанная GRAPHIC-суть F03 об общей шахте."),
    dict(pair="B", group_id="G025", quality="MAJOR_LOSS", astra_changes=5, candidate_changes=3,
         note="Два подбора и перенос сохранены, аэродинамическая модель частична, высотная модель пропущена."),
    dict(pair="B", group_id="G035", quality="SAME_OR_BETTER", astra_changes=2, candidate_changes=1,
         note="Terra укрупнила основное событие; frozen source-first статусы Astra для обеих карточек были insufficient to judge, поэтому подтверждённой reference-потери нет."),
]


def usage_from_receipts(paths):
    rows = []
    for path in paths:
        receipt = read(path)
        usage = receipt["usage"][0]
        rows.append({
            "call": receipt.get("sample_id", path.parent.name),
            **{key: usage.get(key, 0) for key in ["input_tokens", "cached_input_tokens", "output_tokens", "reasoning_output_tokens"]},
        })
    total = {key: sum(row[key] for row in rows)
             for key in ["input_tokens", "cached_input_tokens", "output_tokens", "reasoning_output_tokens"]}
    total["input_plus_output_tokens"] = total["input_tokens"] + total["output_tokens"]
    total["average_input_plus_output_per_group"] = total["input_plus_output_tokens"] / len(rows)
    return rows, total


def build_usage():
    candidate_paths = sorted((OUT / "candidate_raw").glob("*/RECEIPT.json"))
    astra_paths = [
        *[PAIR_A / "miner_raw" / f"PASS_B_{group}" / "RECEIPT.json" for group in ["G002", "G008", "G024", "G031"]],
        *[PAIR_B / "raw" / f"PASS_B_{group}" / "RECEIPT.json" for group in ["G004", "G017", "G025", "G035"]],
    ]
    candidate_rows, candidate = usage_from_receipts(candidate_paths)
    astra_rows, astra = usage_from_receipts(astra_paths)
    token_saving = 1 - candidate["input_plus_output_tokens"] / astra["input_plus_output_tokens"]
    # Base API-price equivalent only; the run used ChatGPT subscription access.
    candidate_cost = ((candidate["input_tokens"] - candidate["cached_input_tokens"]) * 2.0 +
                      candidate["cached_input_tokens"] * 0.2 + candidate["output_tokens"] * 12.0) / 1_000_000
    astra_cost = ((astra["input_tokens"] - astra["cached_input_tokens"]) * 10.0 +
                  astra["cached_input_tokens"] * 1.0 + astra["output_tokens"] * 50.0) / 1_000_000
    value = {
        "created_at": now(), "candidate_model": "gpt-5.6-terra", "candidate_reasoning": "high",
        "candidate_calls": 8, "candidate": candidate, "candidate_by_call": candidate_rows,
        "astra_historical_same_sample": astra, "astra_by_call": astra_rows,
        "estimated_input_plus_output_token_saving_fraction": token_saving,
        "estimated_input_plus_output_token_saving_percent": round(token_saving * 100, 1),
        "base_api_price_equivalent_usd": {"candidate": round(candidate_cost, 4), "astra": round(astra_cost, 4),
                                           "saving_percent": round((1 - candidate_cost / astra_cost) * 100, 1)},
        "cost_caveat": "Runs used Codex/ChatGPT subscription access, not metered API billing. Price-equivalent is indicative only; direct monetary cost is N/A. Token counts and caching differ by model.",
        "cached_input_note": "cached_input_tokens is included within input_tokens; it is reported separately, not added again.",
    }
    write_new(OUT / "TOKEN_USAGE.json", value)
    return value


def build_projection(usage):
    full_a = read(PAIR_A / "TOKEN_USAGE.json")["total"]
    full_b = read(PAIR_B / "MODEL_USAGE.json")["total_tokens"]
    candidate_by_pair = {}
    astra_by_pair = {}
    for pair in ["A", "B"]:
        prefix = f"PAIR_{pair}_"
        candidate_by_pair[pair] = sum(row["input_tokens"] + row["output_tokens"] for row in usage["candidate_by_call"] if row["call"].startswith(prefix))
    astra_by_pair["A"] = sum(row["input_tokens"] + row["output_tokens"] for row in usage["astra_by_call"][:4])
    astra_by_pair["B"] = sum(row["input_tokens"] + row["output_tokens"] for row in usage["astra_by_call"][4:])
    historical = {
        "A": full_a["input_tokens"] + full_a["output_tokens"],
        "B": full_b["input_tokens"] + full_b["output_tokens"],
    }
    ratios = {pair: candidate_by_pair[pair] / astra_by_pair[pair] for pair in ["A", "B"]}
    stratified = round(sum(historical[pair] * ratios[pair] for pair in ["A", "B"]))
    direct = round(sum(historical.values()) * usage["candidate"]["input_plus_output_tokens"] /
                   usage["astra_historical_same_sample"]["input_plus_output_tokens"])
    group_scaled = round(usage["candidate"]["average_input_plus_output_per_group"] * 67)
    value = {
        "created_at": now(), "projection_only_not_fact": True,
        "historical_full_pair_a_plus_b_input_plus_output_tokens": sum(historical.values()),
        "pair_specific_sample_ratios": ratios,
        "projected_full_v3_input_plus_output_tokens": stratified,
        "rounded_report_value": "~5.0 million tokens",
        "sensitivity": {"pooled_ratio_projection": direct, "67_group_average_projection_mining_only": group_scaled,
                        "indicative_range": [min(direct, stratified), group_scaled]},
        "assumptions": ["V3 has approximately the same Pair A/Pair B volume as the historical full mapping+mining runs.",
                        "Sample-specific output compression transfers to the full run.",
                        "No extra repair, retry, verifier, dedupe, or evaluation calls are included."],
        "warning": "The eight groups are structurally controlled, not a statistical random sample; model tokenization and caching differ. Treat ~5.0M as an order-of-magnitude projection.",
    }
    write_new(OUT / "FULL_RUN_COST_PROJECTION.json", value)
    return value


def build_comparisons():
    reference_counts = Counter(row["candidate_outcome"] for row in REFERENCE_ROWS)
    astra_counts = Counter(row["astra_outcome"] for row in REFERENCE_ROWS)
    source = {
        "created_at": now(), "evaluation_opened_after_verified_candidate_freeze": True,
        "reference_definition": "Frozen REAL15/PROVEN10 events intersecting the sample plus frozen source-first CORRECT Pair B G025 events; insufficient-to-judge G035 claims excluded from denominator.",
        "reference_changes": len(REFERENCE_ROWS), "rows": REFERENCE_ROWS,
        "candidate_reference_counts": dict(reference_counts), "candidate_false_changes": FALSE_ROWS,
        "candidate_false_count": len(FALSE_ROWS),
        "candidate_output_counts": {"CORRECT": 5, "PARTIAL": 6, "FALSE": 3, "INSUFFICIENT_TO_JUDGE": 1, "TOTAL": 15},
        "insufficient_note": "G035 candidate claim is not forced into correct/false because the frozen source-first audit rated the corresponding Astra claims insufficient to judge.",
    }
    write_new(OUT / "SOURCE_FIRST_COMPARISON.json", source)
    quality_counts = Counter(row["quality"] for row in GROUP_ROWS)
    comparison = {
        "created_at": now(), "rows": GROUP_ROWS, "quality_counts": dict(quality_counts),
        "astra_reference_metrics": {**dict(astra_counts), "FALSE": 0, "MISSED": 0},
        "candidate_reference_metrics": {**dict(reference_counts), "FALSE": len(FALSE_ROWS)},
        "astra_source_audited_output_counts": {"CORRECT": 11, "PARTIAL": 18, "FALSE": 0, "INSUFFICIENT_TO_JUDGE": 2, "TOTAL": 31},
        "candidate_source_audited_output_counts": source["candidate_output_counts"],
        "atomicity": {"astra_output_changes": 31, "candidate_output_changes": 15,
                      "conclusion": "Candidate is materially less atomic, especially G008, but compression also removed reference content."},
    }
    write_new(OUT / "ASTRA_VS_CANDIDATE.json", comparison)
    return source, comparison


def add_sheet(book, title, rows):
    sheet = book.create_sheet(title)
    if not rows:
        return
    columns = list(rows[0])
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([json.dumps(row[column], ensure_ascii=False) if isinstance(row[column], (list, dict)) else row[column]
                      for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(80, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def build_excel(source, comparison, usage, projection):
    book = Workbook()
    book.remove(book.active)
    summary = [{
        "status": "COMPLETED_CANDIDATE_REJECTED", "candidate": "gpt-5.6-terra", "reasoning": "high",
        "sample_groups": 8, "model_calls": 8, "reference_changes": 11,
        "retained_full": 6, "partial": 3, "missed": 2, "false": 3,
        "same_or_better_groups": 4, "acceptable_loss_groups": 1, "major_loss_groups": 3,
        "input_tokens": usage["candidate"]["input_tokens"], "cached_input_tokens": usage["candidate"]["cached_input_tokens"],
        "output_tokens": usage["candidate"]["output_tokens"], "reasoning_tokens": usage["candidate"]["reasoning_output_tokens"],
        "input_plus_output": usage["candidate"]["input_plus_output_tokens"],
        "token_saving_percent": usage["estimated_input_plus_output_token_saving_percent"],
        "projected_full_v3_tokens": projection["projected_full_v3_input_plus_output_tokens"],
        "recommendation": "TEST_SECOND_CHEAP_CANDIDATE",
    }]
    add_sheet(book, "Summary", summary)
    add_sheet(book, "Reference comparison", source["rows"])
    add_sheet(book, "False changes", source["candidate_false_changes"])
    add_sheet(book, "Group quality", comparison["rows"])
    add_sheet(book, "Candidate usage", usage["candidate_by_call"])
    add_sheet(book, "Astra historical usage", usage["astra_by_call"])
    path = OUT / "RESULTS.xlsx"
    if path.exists():
        raise FileExistsError(path)
    book.save(path)


def build_report(source, comparison, usage, projection):
    u = usage["candidate"]
    a = usage["astra_historical_same_sample"]
    text = f"""# ProjectChange — cheap model screening V1

STATUS: COMPLETED_CANDIDATE_REJECTED

CANDIDATE MODEL: gpt-5.6-terra

REASONING: high

SAMPLE: 8 groups

PAIR A: 4

PAIR B: 4

MODEL CALLS: 8

REFERENCE CHANGES: 11

RETAINED: 6

PARTIAL: 3

MISSED: 2

FALSE: 3

ASTRA XHIGH REFERENCE-EVENT METRICS:
- correct: 10
- partial: 1
- false: 0
- missed: 0

CHEAP CANDIDATE REFERENCE-EVENT METRICS:
- correct: 6
- partial: 3
- false: 3
- missed: 2

VS ASTRA XHIGH:
- same/better: 4
- acceptable loss: 1
- major loss: 3

TOKEN USAGE:
- input: {u['input_tokens']}
- cached: {u['cached_input_tokens']} (included in input)
- output: {u['output_tokens']}
- reasoning: {u['reasoning_output_tokens']}
- input+output: {u['input_plus_output_tokens']}
- average input+output per group: {u['average_input_plus_output_per_group']:.1f}

ASTRA HISTORICAL SAME-SAMPLE USAGE: {a['input_plus_output_tokens']} input+output tokens (input {a['input_tokens']}, cached {a['cached_input_tokens']}, output {a['output_tokens']}, reasoning {a['reasoning_output_tokens']})

ESTIMATED SAVING: {usage['estimated_input_plus_output_token_saving_percent']}% input+output tokens. Base-API-price equivalent suggests {usage['base_api_price_equivalent_usd']['saving_percent']}%, but direct monetary cost is N/A because these calls used ChatGPT subscription access and the models tokenize/cache differently.

PROJECTED FULL V3 COST: ~5.0M input+output tokens ({projection['projected_full_v3_input_plus_output_tokens']} stratified estimate; indicative range {projection['sensitivity']['indicative_range'][0]}–{projection['sensitivity']['indicative_range'][1]}). Projection only.

QUALITY CONCLUSION: Terra/high хорошо укрупнила атомарный Pair A output и уверенно прочитала многие таблицы, числа и несколько cross-modal связей. Но она создала три ложных инженерных изменения из титульных метаданных, полностью пропустила доказанный переход к общей шахте в G017 и потеряла часть расчётной модели в G025. Экономия заметная, но качество не проходит заданный порог: новые false changes есть, а два reference events пропущены.

RECOMMENDATION: TEST_SECOND_CHEAP_CANDIDATE

Следующий разумный кандидат — `gpt-5.6-sol / high` на тех же восьми frozen inputs отдельным контролируемым тестом. До такого теста Terra/high не использовать для полного V3; полный прогон автоматически не запускался.

PRODUCTION: UNCHANGED

VALIDATION: NOT OPENED

FINAL HOLDOUT: NOT OPENED

## Trace notes

- Astra xhigh was not rerun; all Astra values are historical frozen outputs/usages for the same groups.
- Candidate result was hash-frozen before REAL15, PROVEN10, source-first audits, or Astra outputs were opened.
- G035 remains `INSUFFICIENT_TO_JUDGE` in frozen source-first evidence and is excluded from the 11-change reference denominator.
- OpenRouter calls: 0. Claude calls: 0. Candidate retries: 0.
- Official model capability/pricing source checked before selection: https://developers.openai.com/api/docs/models/compare
"""
    path = OUT / "FINAL_REPORT.md"
    with path.open("x") as handle:
        handle.write(text)


def main():
    verify_candidate_freeze()
    source, comparison = build_comparisons()
    usage = build_usage()
    projection = build_projection(usage)
    build_excel(source, comparison, usage, projection)
    build_report(source, comparison, usage, projection)
    deliverables = [
        "AVAILABLE_MODEL_CANDIDATES.json", "SCREENING_SAMPLE_FREEZE.json", "CANDIDATE_RESULT_FREEZE.json",
        "SOURCE_FIRST_COMPARISON.json", "ASTRA_VS_CANDIDATE.json", "TOKEN_USAGE.json",
        "FULL_RUN_COST_PROJECTION.json", "RESULTS.xlsx", "FINAL_REPORT.md",
    ]
    write_new(OUT / "DELIVERY_MANIFEST.json", {
        "created_at": now(), "status": "COMPLETED_CANDIDATE_REJECTED",
        "files": {name: sha(OUT / name) for name in deliverables},
        "candidate_calls": 8, "astra_calls": 0, "openrouter_calls": 0, "claude_calls": 0,
        "production": "UNCHANGED", "validation": "NOT OPENED", "final_holdout": "NOT OPENED",
    })
    print(json.dumps({"status": "COMPLETED_CANDIDATE_REJECTED", "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
