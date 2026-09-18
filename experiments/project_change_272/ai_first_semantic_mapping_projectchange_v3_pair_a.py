#!/usr/bin/env python3
"""Frozen-architecture staged execution of V3 for Pair A only.

This runner is intentionally separate from the completed Pair B runner.  It
consumes the pre-Pair-B Pair A map and prepared inputs, makes no mapper call,
and opens Pair A truth/controls only after the final result freeze.
"""

from __future__ import annotations

import argparse
import asyncio
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import signal
import sys
import time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from experiments.project_change_272 import ai_first_semantic_mapping_projectchange_v3 as v3  # noqa: E402


OUT = v3.OUT
PAIR = "A"
PAIR_A_ROOT = v3.CORPUS / "pair_a_ai_mapping_change_miner_v1"
PAIR_A_TRUTH = v3.CORPUS / "fresh_dev_pair_a_live_v2/SOURCE_FIRST_TRUTH.json"
PAIR_A_SOURCE_AUDIT = PAIR_A_ROOT / "FALSE_POSITIVE_AUDIT.json"
PAIR_A_CONTROLS = PAIR_A_ROOT / "CONTROL_AUDIT.json"
PAIR_B_EVALUATION_FREEZE = OUT / "PAIR_B_EVALUATION_FREEZE.json"
SOFT_WARNING = 8_000_000
HARD_CAP = 9_000_000
FUTURE_DEDUPE_ALLOWANCE = 320_000


S = {"type": "string"}
STRINGS = {"type": "array", "items": S}
EVAL_ROW = v3.obj(
    projectchange_id=S,
    quality={"type": "string", "enum": ["CORRECT", "PARTIAL", "FALSE", "INSUFFICIENT_TO_JUDGE"]},
    rationale=S,
    grouping_quality={"type": "string", "enum": ["GOOD_HUMAN_LEVEL", "TOO_ATOMIC", "OVER_MERGED", "DUPLICATE"]},
    grouping_rationale=S,
    source_refs=STRINGS,
)
TRUTH_ROW = v3.obj(
    truth_id=S,
    outcome={"type": "string", "enum": ["STRONG", "PARTIAL", "MISSED"]},
    matched_projectchange_ids=STRINGS,
    rationale=S,
)
EVAL_SCHEMA = v3.obj(
    pair={"type": "string", "enum": ["A"]},
    projectchange_rows={"type": "array", "items": EVAL_ROW},
    truth_rows={"type": "array", "items": TRUTH_ROW},
    notes=STRINGS,
)


EVALUATION_PROMPT = """Ты post-freeze source-first evaluator AI-FIRST
PROJECTCHANGE MINER V3 для Pair A / АР1. Final output уже frozen: не меняй
его, не создавай ProjectChanges и не предлагай tuning.

Оцени каждый final ProjectChange ровно один раз: CORRECT, PARTIAL, FALSE или
INSUFFICIENT_TO_JUDGE. Затем ровно одна grouping label: GOOD_HUMAN_LEVEL,
TOO_ATOMIC, OVER_MERGED или DUPLICATE. Проверяй engineering subject,
scope/location, OLD и NEW states, направление, evidence с обеих сторон и
отсутствие metadata-only claims. FALSE — только при доказанной ошибке или
контрдоказательстве; PARTIAL — доказанное ядро с существенным overclaim либо
неполнотой; INSUFFICIENT_TO_JUDGE — когда source пакета недостаточно для
решения. TOO_ATOMIC — часть одного более широкого human-level события;
OVER_MERGED — независимые события соединены; DUPLICATE — тот же event остался
другой final card (назови ID); иначе GOOD_HUMAN_LEVEL.

Для каждого frozen REAL15 верни STRONG/PARTIAL/MISSED только по final V3, не
по hints. STRONG требует корректного смысла и достаточного охвата; PARTIAL —
смысл найден, но существенная часть пропущена или искажена. Existing Pair A
source-first controls и prior same-version source audit — дополнительная
проверенная evidence; старые miner IDs не являются V3 совпадениями, а сходство
формулировок само по себе не доказывает V3 claim. Верни exact partition всех
ProjectChange IDs и REAL15 IDs, только JSON по схеме, по-русски."""


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def stage_receipts(kind: str | None = None) -> list[dict[str, Any]]:
    roots = []
    if kind in (None, "MINING"):
        roots.append(OUT / "miner_raw")
    if kind in (None, "DEDUPE"):
        roots.append(OUT / "dedupe_raw")
    if kind in (None, "EVALUATION"):
        roots.append(OUT / "evaluation_raw")
    rows = []
    for root in roots:
        if not root.exists():
            continue
        for path in sorted(root.glob("PAIR_A*/RECEIPT.json")):
            receipt = v3.read_json(path)
            if receipt.get("pair") == "A" or receipt.get("call_id", "").startswith("PAIR_A"):
                rows.append(receipt)
    return rows


def usage_total(receipts: list[dict[str, Any]]) -> int:
    return sum(
        (usage.get("input_tokens") or 0) + (usage.get("output_tokens") or 0)
        for receipt in receipts
        for usage in receipt.get("usage", [])
    )


def optimized_preflight() -> dict[str, Any]:
    value = v3.read_json(OUT / "TOKEN_PREFLIGHT_OPTIMIZED.json")
    if value["pair_a_projected_mining_input_output"] != 7_884_077:
        raise RuntimeError("Pair A frozen projection drift")
    if value["mapping_freeze_sha256"] != v3.sha256(OUT / "SEMANTIC_MAPPING_V2_FREEZE.json"):
        raise RuntimeError("Mapping freeze drift")
    return value


def verify_pair_b_complete() -> None:
    freeze = v3.read_json(PAIR_B_EVALUATION_FREEZE)
    if freeze.get("status") != "PAIR_B_EVALUATION_COMPLETE_STOP":
        raise RuntimeError("Pair B evaluation gate is not complete")
    comparison = v3.read_json(OUT / "PAIR_B_V2_VS_V3.json")
    if comparison.get("recommendation") != "PAIR_B_GATE_PASS_RUN_PAIR_A":
        raise RuntimeError("Pair B gate does not authorize Pair A")
    for name, digest in freeze["hashes"].items():
        if v3.sha256(OUT / name) != digest:
            raise RuntimeError(f"Pair B evaluation freeze drift: {name}")


def freeze_architecture() -> None:
    target = OUT / "PAIR_A_STAGED_ARCHITECTURE_FREEZE.json"
    if target.exists():
        raise FileExistsError("Pair A staged architecture already frozen")
    verify_pair_b_complete()
    preflight = optimized_preflight()
    mapping = v3.read_json(OUT / "PAIR_A_SEMANTIC_MAP.json")
    if len(mapping["regions"]) != 25:
        raise RuntimeError("Pair A region count drift")
    if stage_receipts("MINING") or stage_receipts("DEDUPE"):
        raise RuntimeError("Pair A inference already started before staged freeze")
    calls = [row for row in preflight["calls"] if row["pair"] == "A"]
    rows = {row["region_id"]: row for row in calls}
    if set(rows) != {region["region_id"] for region in mapping["regions"]}:
        raise RuntimeError("Optimized Pair A preflight is not exact region partition")
    hashes = {
        "PAIR_A_SEMANTIC_MAP.json": v3.sha256(OUT / "PAIR_A_SEMANTIC_MAP.json"),
        "SEMANTIC_MAPPING_V2_FREEZE.json": v3.sha256(OUT / "SEMANTIC_MAPPING_V2_FREEZE.json"),
        "MINER_PROMPT.txt": v3.sha256(OUT / "MINER_PROMPT.txt"),
        "DEDUPE_PROMPT.txt": v3.sha256(OUT / "DEDUPE_PROMPT.txt"),
        "TOKEN_PREFLIGHT_OPTIMIZED.json": v3.sha256(OUT / "TOKEN_PREFLIGHT_OPTIMIZED.json"),
        "PAIR_B_EVALUATION_FREEZE.json": v3.sha256(PAIR_B_EVALUATION_FREEZE),
        "runner": v3.sha256(Path(__file__)),
    }
    for region in mapping["regions"]:
        call_id = f"PAIR_A_{region['region_id']}"
        for name in ("MODEL_INPUT.json", "EXACT_PROMPT.txt", "IMAGE_MANIFEST.json"):
            path = OUT / "miner_inputs_optimized" / call_id / name
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
        if v3.sha256(OUT / "miner_inputs_optimized" / call_id / "EXACT_PROMPT.txt") != rows[region["region_id"]]["input_sha256"]:
            raise RuntimeError(f"Optimized input drift: {call_id}")
    v3.write_new(target, {
        "frozen_at": now(), "status": "FROZEN_BEFORE_PAIR_A_MINING",
        "model": v3.MODEL, "reasoning": v3.REASONING, "provider": "codex_chatgpt",
        "openrouter": 0, "pair_a_regions": 25, "mapper_calls": 0,
        "pair_a_truth_opened": False,
        "pair_a_projected_mining_input_output": preflight["pair_a_projected_mining_input_output"],
        "pair_a_mining_dedupe_soft_warning": SOFT_WARNING,
        "pair_a_mining_dedupe_hard_cap": HARD_CAP,
        "soft_warning_active_at_start": preflight["pair_a_projected_mining_input_output"] + FUTURE_DEDUPE_ALLOWANCE > SOFT_WARNING,
        "future_dedupe_allowance": FUTURE_DEDUPE_ALLOWANCE,
        "prompt_semantics_changed_after_pair_b": False,
        "mapping_changed_after_pair_b": False,
        "evidence_changed_after_pair_b": False,
        "schema_changed_after_pair_b": False,
        "dedupe_contract_changed_after_pair_b": False,
        "gate_thresholds_frozen_before_pair_a_truth": {
            "real15_missed_max": 1,
            "real15_recovered_min": 14,
            "false_max": 3,
            "false_share_max": 0.05,
            "final_projectchanges_max": 110,
            "too_atomic_max": 50,
            "good_human_level_share_min": 0.50,
            "over_merged_max": 5,
            "token_hard_cap": HARD_CAP,
        },
        "schema": v3.MINER_SCHEMA,
        "hashes": hashes,
    })
    print("PAIR_A_STAGED_ARCHITECTURE_FROZEN", flush=True)


def verify_architecture() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_A_STAGED_ARCHITECTURE_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        path = Path(__file__) if relative == "runner" else OUT / relative
        if v3.sha256(path) != digest:
            raise RuntimeError(f"Pair A architecture drift: {relative}")
    verify_pair_b_complete()
    return freeze


async def call(stage: str, call_id: str, prompt: str, data: Any, schema: dict[str, Any], images: list[dict[str, Any]], prepared_input: Path | None = None) -> dict[str, Any]:
    raw_dir = {"MINING": "miner_raw", "DEDUPE": "dedupe_raw", "EVALUATION": "evaluation_raw"}[stage]
    input_dir_name = {"MINING": "miner_inputs_optimized", "DEDUPE": "dedupe_inputs", "EVALUATION": "evaluation_inputs"}[stage]
    target = OUT / raw_dir / call_id
    target.mkdir(parents=True, exist_ok=False)
    input_dir = OUT / input_dir_name / call_id
    if stage != "MINING":
        input_dir.mkdir(parents=True, exist_ok=False)
    names, labels, seen = [], [], set()
    for row in images:
        source = Path(row["path"])
        digest = v3.sha256(source)
        if digest in seen:
            continue
        seen.add(digest)
        name = f"image_{len(names):03d}.png"
        shutil.copyfile(source, target / name)
        names.append(name)
        labels.append({**row["label"], "image": len(names)})
    payload = prompt + "\nIMAGES:\n" + json.dumps(labels, ensure_ascii=False) + "\nSOURCE DATA:\n" + json.dumps(data, ensure_ascii=False)
    if prepared_input is not None:
        if prepared_input.read_text(encoding="utf-8") != payload or v3.read_json(prepared_input.parent / "MODEL_INPUT.json") != data:
            raise RuntimeError(f"Prepared optimized input drift: {call_id}")
    else:
        v3.write_new(input_dir / "MODEL_INPUT.json", data)
        v3.write_new(input_dir / "EXACT_PROMPT.txt", payload)
    v3.write_new(target / "prompt.txt", payload)
    v3.write_new(target / "schema.json", schema)
    command = v3.sandbox_command(target, v3.cli_command(names))
    input_hashes = {path.name: v3.sha256(path) for path in target.iterdir()}
    v3.write_new(target / "INVOCATION.json", {
        "call_id": call_id, "stage": stage, "pair": "A", "at": now(),
        "provider": "codex_chatgpt", "model": v3.MODEL, "reasoning": v3.REASONING,
        "openrouter": 0, "command": command, "environment_keys": list(v3.safe_env()),
        "input_hashes": input_hashes,
    })
    started = time.monotonic()
    process = None
    try:
        with (target / "raw.jsonl").open("wb") as raw, (target / "stderr.txt").open("wb") as err:
            process = await asyncio.create_subprocess_exec(
                *command, stdin=asyncio.subprocess.PIPE, stdout=raw, stderr=err,
                env=v3.safe_env(), start_new_session=True,
            )
            await asyncio.wait_for(process.communicate(payload.encode("utf-8")), 1800)
    except BaseException:
        if process is not None and process.returncode is None:
            os.killpg(process.pid, signal.SIGKILL)
            await process.wait()
        raise
    records = []
    for line in (target / "raw.jsonl").read_text(encoding="utf-8").splitlines():
        try:
            records.append(json.loads(line))
        except ValueError:
            pass
    usage = [row["usage"] for row in records if row.get("type") == "turn.completed" and row.get("usage")]
    tools = [row for row in records if row.get("type", "").startswith("item.") and row.get("item", {}).get("type") not in {None, "agent_message", "reasoning", "error"}]
    receipt = {
        "call_id": call_id, "stage": stage, "pair": "A", "at": now(),
        "exit_code": process.returncode, "wall_time_seconds": time.monotonic() - started,
        "usage": usage, "tool_items": len(tools), "raw_sha256": v3.sha256(target / "raw.jsonl"),
    }
    v3.write_new(target / "RECEIPT.json", receipt)
    if process.returncode or len(usage) != 1 or tools:
        raise RuntimeError(f"Model call failed; no automatic retry: {call_id}")
    for name, digest in input_hashes.items():
        if v3.sha256(target / name) != digest:
            raise RuntimeError(f"Inference input drift: {call_id}/{name}")
    value = v3.read_json(target / "final.txt")
    v3.jsonschema.validate(value, schema)
    v3.write_new(target / "parsed.json", value)
    print(json.dumps({"call": call_id, "status": "SUCCESS", "usage": usage[0]}, ensure_ascii=False), flush=True)
    return value


async def mine_pair_a() -> None:
    verify_architecture()
    preflight = optimized_preflight()
    calls = [row for row in preflight["calls"] if row["pair"] == "A"]
    mapping = v3.read_json(OUT / "PAIR_A_SEMANTIC_MAP.json")
    results = []
    for index, region in enumerate(mapping["regions"]):
        actual = usage_total(stage_receipts("MINING"))
        projected_remaining = sum(row["projected_tokens_after"] for row in calls[index:]) + FUTURE_DEDUPE_ALLOWANCE
        if actual + projected_remaining > HARD_CAP:
            raise RuntimeError("PAIR_A_TOKEN_BUDGET_REVIEW_REQUIRED")
        if actual + projected_remaining > SOFT_WARNING:
            print(json.dumps({"status": "PAIR_A_TOKEN_SOFT_WARNING", "actual": actual, "projected_remaining": projected_remaining, "projected_total": actual + projected_remaining}), flush=True)
        data, images, _ = v3.optimized_region_bundle("A", region)
        call_id = f"PAIR_A_{region['region_id']}"
        value = await call("MINING", call_id, v3.MINER_PROMPT, data, v3.MINER_SCHEMA, images, OUT / "miner_inputs_optimized" / call_id / "EXACT_PROMPT.txt")
        v3.validate_miner("A", region, value)
        results.append(value)
    result_path = OUT / "PAIR_A_PROJECTCHANGE_MINER_RESULTS.json"
    v3.write_new(result_path, {
        "created_at": now(), "model": v3.MODEL, "reasoning": v3.REASONING,
        "regions": results,
        "projectchanges": [change for region in results for change in region["projectchanges"]],
        "unresolved_hints": [hint for region in results for hint in region["unresolved_hints"]],
    })
    hashes = {
        "PAIR_A_PROJECTCHANGE_MINER_RESULTS.json": v3.sha256(result_path),
        "PAIR_A_STAGED_ARCHITECTURE_FREEZE.json": v3.sha256(OUT / "PAIR_A_STAGED_ARCHITECTURE_FREEZE.json"),
    }
    for path in sorted((OUT / "miner_raw").glob("PAIR_A*/*")):
        if path.is_file():
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
    actual = usage_total(stage_receipts("MINING"))
    v3.write_new(OUT / "PAIR_A_PROJECTCHANGE_MINER_V3_FREEZE.json", {
        "frozen_at": now(), "status": "PAIR_A_MINER_FROZEN", "model": v3.MODEL,
        "reasoning": v3.REASONING, "regions": 25, "calls": 25,
        "projectchanges": len([c for r in results for c in r["projectchanges"]]),
        "unresolved_hints": len([h for r in results for h in r["unresolved_hints"]]),
        "actual_input_output": actual, "hard_cap": HARD_CAP, "truth_opened": False,
        "hashes": hashes,
    })
    print("PAIR_A_PROJECTCHANGE_MINER_V3_FROZEN", flush=True)


def verify_miner_freeze() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_A_PROJECTCHANGE_MINER_V3_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        if v3.sha256(OUT / relative) != digest:
            raise RuntimeError(f"Pair A Miner freeze drift: {relative}")
    return freeze


async def dedupe_pair_a() -> None:
    verify_miner_freeze()
    mined = v3.read_json(OUT / "PAIR_A_PROJECTCHANGE_MINER_RESULTS.json")
    payload = {"pair": "A", "projectchanges": [v3.compact_change(change) for change in mined["projectchanges"]]}
    if usage_total(stage_receipts("MINING")) + FUTURE_DEDUPE_ALLOWANCE > HARD_CAP:
        raise RuntimeError("PAIR_A_TOKEN_BUDGET_REVIEW_REQUIRED")
    raw = await call("DEDUPE", "PAIR_A_DEDUPE", v3.DEDUPE_PROMPT, payload, v3.DEDUPE_SCHEMA, [])
    if raw["pair"] != "A":
        raise RuntimeError("Dedupe pair mismatch")
    final = v3.apply_dedupe("A", mined["projectchanges"], raw)
    v3.write_new(OUT / "PAIR_A_DEDUPE.json", raw)
    final_path = OUT / "PAIR_A_FINAL_PROJECTCHANGES.json"
    v3.write_new(final_path, {
        "created_at": now(), "pair": "A", "model": v3.MODEL, "reasoning": v3.REASONING,
        "projectchanges": final, "unresolved_hints": mined["unresolved_hints"],
    })
    actual = usage_total(stage_receipts("MINING") + stage_receipts("DEDUPE"))
    if actual > HARD_CAP:
        raise RuntimeError("PAIR_A_TOKEN_BUDGET_REVIEW_REQUIRED")
    hashes = {name: v3.sha256(OUT / name) for name in (
        "PAIR_A_PROJECTCHANGE_MINER_V3_FREEZE.json", "PAIR_A_DEDUPE.json", "PAIR_A_FINAL_PROJECTCHANGES.json"
    )}
    for path in sorted((OUT / "dedupe_raw" / "PAIR_A_DEDUPE").glob("*")):
        if path.is_file():
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
    v3.write_new(OUT / "PAIR_A_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json", {
        "frozen_at": now(), "status": "PAIR_A_FINAL_FROZEN", "model": v3.MODEL,
        "reasoning": v3.REASONING, "mined": len(mined["projectchanges"]), "final": len(final),
        "unresolved_hints": len(mined["unresolved_hints"]),
        "mining_dedupe_input_output": actual, "soft_warning": SOFT_WARNING,
        "hard_cap": HARD_CAP, "truth_opened": False,
        "semantic_mutation_after_pair_b": False, "hashes": hashes,
    })
    print("PAIR_A_AI_FIRST_PROJECTCHANGE_V3_RESULT_FROZEN", flush=True)


def verify_final_freeze() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_A_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        if v3.sha256(OUT / relative) != digest:
            raise RuntimeError(f"Pair A final freeze drift: {relative}")
    return freeze


def compact_change(change: dict[str, Any]) -> dict[str, Any]:
    keys = ("projectchange_id", "engineering_subject", "scope", "locations", "change_summary", "old_state", "new_state", "changed_parameters", "old_pages", "new_pages", "modalities", "confidence", "why_one_event", "dedupe_lineage", "dedupe_reason")
    return {key: change[key] for key in keys if key in change} | {
        "evidence_items": [{key: item[key] for key in ("side", "physical_page", "block_id", "block_type", "relevant_fragment", "evidence_role")} for item in change["evidence_items"]]
    }


def compact_source_audit() -> list[dict[str, Any]]:
    keys = ("change_id", "audit_status", "engineering_subject", "change_summary", "reason", "old_pages", "new_pages")
    return [{key: row.get(key) for key in keys} for row in v3.read_json(PAIR_A_SOURCE_AUDIT)["rows"]]


def evaluation_payload() -> tuple[dict[str, Any], list[str], list[str]]:
    final = v3.read_json(OUT / "PAIR_A_FINAL_PROJECTCHANGES.json")["projectchanges"]
    truth = v3.read_json(PAIR_A_TRUTH)
    real = [entry for entry in truth["entries"] if entry.get("classification") == "REAL_CHANGE"]
    if len(real) != 15:
        raise RuntimeError(f"Expected REAL15, got {len(real)}")
    truth_keys = ("id", "title", "old_pages", "new_pages", "old_state", "new_state", "basis", "route")
    control_keys = ("control_id", "classification", "title", "basis")
    controls = v3.read_json(PAIR_A_CONTROLS)["rows"]
    payload = {
        "pair": "A",
        "frozen_v3_projectchanges": [compact_change(item) for item in final],
        "source_first_truth": [{key: item.get(key) for key in truth_keys} for item in real],
        "source_first_controls": [{key: item.get(key) for key in control_keys} for item in controls],
        "prior_same_version_source_audit": compact_source_audit(),
    }
    return payload, [item["projectchange_id"] for item in final], [item["id"] for item in real]


def validate_evaluation(value: dict[str, Any], projectchange_ids: list[str], truth_ids: list[str]) -> None:
    if value["pair"] != "A":
        raise RuntimeError("Evaluation pair mismatch")
    got_changes = [row["projectchange_id"] for row in value["projectchange_rows"]]
    got_truth = [row["truth_id"] for row in value["truth_rows"]]
    if len(got_changes) != len(set(got_changes)) or set(got_changes) != set(projectchange_ids):
        raise RuntimeError("Evaluation is not exact ProjectChange partition")
    if len(got_truth) != len(set(got_truth)) or set(got_truth) != set(truth_ids):
        raise RuntimeError("Evaluation is not exact REAL15 partition")
    matched = {item for row in value["truth_rows"] for item in row["matched_projectchange_ids"]}
    if not matched <= set(projectchange_ids):
        raise RuntimeError("Evaluation refers to unknown ProjectChange")


def summarize(rows: list[dict[str, Any]], field: str, labels: list[str]) -> dict[str, int]:
    return {label: sum(row[field] == label for row in rows) for label in labels}


def usage_report() -> dict[str, Any]:
    def section(receipts: list[dict[str, Any]]) -> dict[str, Any]:
        usages = [usage for receipt in receipts for usage in receipt.get("usage", [])]
        return {
            "calls": len(usages),
            "input_tokens": sum(item.get("input_tokens") or 0 for item in usages),
            "cached_input_tokens": sum(item.get("cached_input_tokens") or 0 for item in usages),
            "output_tokens": sum(item.get("output_tokens") or 0 for item in usages),
            "reasoning_output_tokens": sum(item.get("reasoning_output_tokens") or 0 for item in usages),
            "input_output": sum((item.get("input_tokens") or 0) + (item.get("output_tokens") or 0) for item in usages),
        }
    mining = section(stage_receipts("MINING"))
    dedupe = section(stage_receipts("DEDUPE"))
    evaluation = section(stage_receipts("EVALUATION"))
    total = {key: mining[key] + dedupe[key] for key in mining}
    pair_b = v3.read_json(OUT / "PAIR_B_TOKEN_USAGE.json")["pair_b_mining_dedupe_total"]
    combined = {key: pair_b[key] + total[key] for key in total}
    return {
        "accounting_note": "cached_input_tokens are a subset of input_tokens and are not added again",
        "pair_a_mining": mining, "pair_a_dedupe": dedupe,
        "pair_a_mining_dedupe_total": total, "pair_a_evaluation_separate": evaluation,
        "pair_b_v3_actual": pair_b, "combined_v3_mining_dedupe": combined,
    }


def write_results(evaluation: dict[str, Any]) -> None:
    quality = summarize(evaluation["projectchange_rows"], "quality", ["CORRECT", "PARTIAL", "FALSE", "INSUFFICIENT_TO_JUDGE"])
    grouping = summarize(evaluation["projectchange_rows"], "grouping_quality", ["GOOD_HUMAN_LEVEL", "TOO_ATOMIC", "OVER_MERGED", "DUPLICATE"])
    real15 = summarize(evaluation["truth_rows"], "outcome", ["STRONG", "PARTIAL", "MISSED"])
    final = v3.read_json(OUT / "PAIR_A_FINAL_PROJECTCHANGES.json")
    usage = usage_report()
    thresholds = v3.read_json(OUT / "PAIR_A_STAGED_ARCHITECTURE_FREEZE.json")["gate_thresholds_frozen_before_pair_a_truth"]
    n = len(final["projectchanges"])
    false_share = quality["FALSE"] / n if n else 0
    gate_checks = {
        "real15_no_substantial_degradation": real15["MISSED"] <= thresholds["real15_missed_max"] and real15["STRONG"] + real15["PARTIAL"] >= thresholds["real15_recovered_min"],
        "false_low": quality["FALSE"] <= thresholds["false_max"] and false_share <= thresholds["false_share_max"],
        "final_substantially_below_v2": n <= thresholds["final_projectchanges_max"],
        "too_atomic_substantially_below_v2": grouping["TOO_ATOMIC"] <= thresholds["too_atomic_max"],
        "good_is_primary_useful_share": n > 0 and grouping["GOOD_HUMAN_LEVEL"] / n >= thresholds["good_human_level_share_min"],
        "no_mass_over_merge": grouping["OVER_MERGED"] <= thresholds["over_merged_max"],
        "within_token_cap": usage["pair_a_mining_dedupe_total"]["input_output"] <= thresholds["token_hard_cap"],
    }
    if all(gate_checks.values()):
        recommendation = "SEMANTIC_MAPPING_PROJECTCHANGE_PROMISING"
    elif gate_checks["within_token_cap"] and gate_checks["real15_no_substantial_degradation"] and sum(gate_checks.values()) >= 5:
        recommendation = "MORE_CONTROLLED_TESTING_REQUIRED"
    else:
        recommendation = "SEMANTIC_MAPPING_PROJECTCHANGE_NOT_PROVEN"
    comparison = {
        "pair_a_v2": {"final_projectchanges": 146, "GOOD_HUMAN_LEVEL": 65, "TOO_ATOMIC": 77, "OVER_MERGED": 3, "DUPLICATE": 1, "REAL15": {"STRONG": 14, "PARTIAL": 1, "MISSED": 0}},
        "pair_a_v3": {"final_projectchanges": n, **grouping, "REAL15": real15},
        "gate_thresholds_frozen_before_truth": thresholds,
        "gate_checks": gate_checks,
        "recommendation": recommendation,
    }
    v3.write_new(OUT / "PAIR_A_EVALUATION.json", evaluation)
    v3.write_new(OUT / "PAIR_A_GROUPING_QUALITY.json", {
        "counts": grouping,
        "rows": [{"projectchange_id": row["projectchange_id"], "grouping_quality": row["grouping_quality"], "rationale": row["grouping_rationale"]} for row in evaluation["projectchange_rows"]],
    })
    v3.write_new(OUT / "PAIR_A_TOKEN_USAGE.json", usage)
    v3.write_new(OUT / "PAIR_A_V2_VS_V3.json", comparison)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "V2", "V3"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in [
        ("Final ProjectChanges", 146, n),
        ("Unresolved hints", "", len(final["unresolved_hints"])),
        *[(label, comparison["pair_a_v2"].get(label, ""), grouping[label]) for label in grouping],
        *[(f"REAL15 {label}", comparison["pair_a_v2"]["REAL15"][label], real15[label]) for label in real15],
        ("Recommendation", "", recommendation),
    ]:
        summary.append(row)
    changes = workbook.create_sheet("ProjectChanges")
    changes.append(["projectchange_id", "quality", "grouping_quality", "rationale", "grouping_rationale", "source_refs"])
    for row in evaluation["projectchange_rows"]:
        changes.append([row["projectchange_id"], row["quality"], row["grouping_quality"], row["rationale"], row["grouping_rationale"], "\n".join(row["source_refs"])])
    truth = workbook.create_sheet("REAL15")
    truth.append(["truth_id", "outcome", "matched_projectchange_ids", "rationale"])
    for row in evaluation["truth_rows"]:
        truth.append([row["truth_id"], row["outcome"], ", ".join(row["matched_projectchange_ids"]), row["rationale"]])
    usage_sheet = workbook.create_sheet("Token usage")
    usage_sheet.append(["Stage", "calls", "input", "cached", "output", "reasoning", "input+output"])
    for key in ("pair_a_mining", "pair_a_dedupe", "pair_a_mining_dedupe_total", "pair_a_evaluation_separate", "pair_b_v3_actual", "combined_v3_mining_dedupe"):
        row = usage[key]
        usage_sheet.append([key, row["calls"], row["input_tokens"], row["cached_input_tokens"], row["output_tokens"], row["reasoning_output_tokens"], row["input_output"]])
    workbook.save(OUT / "PAIR_A_RESULTS.xlsx")
    pair_b = v3.read_json(OUT / "PAIR_B_V2_VS_V3.json")["pair_b_v3"]
    report = f"""# ProjectChange V3 — staged Pair A final report

STATUS: COMPLETE_PAIR_A_V3_FROZEN

PAIR A:
semantic regions: 25
Miner calls: {usage['pair_a_mining']['calls']}
Final ProjectChanges: {n}
Unresolved hints: {len(final['unresolved_hints'])}

REAL15: STRONG {real15['STRONG']}; PARTIAL {real15['PARTIAL']}; MISSED {real15['MISSED']}.
Quality: CORRECT {quality['CORRECT']}; PARTIAL {quality['PARTIAL']}; FALSE {quality['FALSE']}; INSUFFICIENT {quality['INSUFFICIENT_TO_JUDGE']}.
Grouping: GOOD_HUMAN_LEVEL {grouping['GOOD_HUMAN_LEVEL']}; TOO_ATOMIC {grouping['TOO_ATOMIC']}; OVER_MERGED {grouping['OVER_MERGED']}; DUPLICATE {grouping['DUPLICATE']}.

TOKEN USAGE Pair A mining+dedupe: {usage['pair_a_mining_dedupe_total']['input_output']} / {HARD_CAP}.
Combined Pair B + Pair A mining+dedupe: {usage['combined_v3_mining_dedupe']['input_output']}.

PAIR A V2 → V3: final 146 → {n}; GOOD 65 → {grouping['GOOD_HUMAN_LEVEL']}; TOO_ATOMIC 77 → {grouping['TOO_ATOMIC']}; OVER_MERGED 3 → {grouping['OVER_MERGED']}; DUPLICATE 1 → {grouping['DUPLICATE']}; REAL15 14/1/0 → {real15['STRONG']}/{real15['PARTIAL']}/{real15['MISSED']}.

PAIR B V3: {pair_b['final_projectchanges']} final; {pair_b['GOOD_HUMAN_LEVEL']} GOOD; {pair_b['DUPLICATE']} duplicates; {pair_b['PROVEN10']['STRONG']}/10 STRONG; {pair_b['PROVEN10']['PARTIAL']}/10 PARTIAL; {pair_b['PROVEN10']['MISSED']} MISSED.

COMBINED V3 CONCLUSION: Pair A проверена той же frozen V3 после Pair B; решение основано на заранее зафиксированных gate thresholds: {json.dumps(gate_checks, ensure_ascii=False)}.

RECOMMENDATION: {recommendation}

NO POST-TRUTH TUNING OR RERUN.
PRODUCTION: UNCHANGED
UI: UNCHANGED
VALIDATION: NOT OPENED
FINAL HOLDOUT: NOT OPENED
"""
    v3.write_new(OUT / "PAIR_A_FINAL_REPORT.md", report)


async def evaluate_pair_a() -> None:
    freeze = verify_final_freeze()
    if freeze["truth_opened"] is not False:
        raise RuntimeError("Invalid pre-evaluation truth state")
    payload, projectchange_ids, truth_ids = evaluation_payload()
    hashes = {
        "pair_a_truth": v3.sha256(PAIR_A_TRUTH),
        "pair_a_source_audit": v3.sha256(PAIR_A_SOURCE_AUDIT),
        "pair_a_controls": v3.sha256(PAIR_A_CONTROLS),
        "pair_a_result_freeze": v3.sha256(OUT / "PAIR_A_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json"),
    }
    v3.write_new(OUT / "PAIR_A_EVALUATION_INPUT_FREEZE.json", {
        "opened_at": now(), "pair_a_truth_opened_after_final_freeze": True, "hashes": hashes,
    })
    value = await call("EVALUATION", "PAIR_A_SOURCE_FIRST_EVALUATION", EVALUATION_PROMPT, payload, EVAL_SCHEMA, [])
    validate_evaluation(value, projectchange_ids, truth_ids)
    write_results(value)
    names = ("PAIR_A_EVALUATION.json", "PAIR_A_GROUPING_QUALITY.json", "PAIR_A_TOKEN_USAGE.json", "PAIR_A_V2_VS_V3.json", "PAIR_A_RESULTS.xlsx", "PAIR_A_FINAL_REPORT.md")
    v3.write_new(OUT / "PAIR_A_EVALUATION_FREEZE.json", {
        "frozen_at": now(), "status": "PAIR_A_EVALUATION_COMPLETE_STOP",
        "no_post_truth_tuning": True, "no_rerun": True,
        "production": "UNCHANGED", "ui": "UNCHANGED",
        "validation": "NOT_OPENED", "final_holdout": "NOT_OPENED",
        "hashes": {name: v3.sha256(OUT / name) for name in names},
    })
    print((OUT / "PAIR_A_FINAL_REPORT.md").read_text(encoding="utf-8"), flush=True)


def run(action: str) -> None:
    try:
        asyncio.run({"mine": mine_pair_a, "dedupe": dedupe_pair_a, "evaluate": evaluate_pair_a}[action]())
    except BaseException as exc:
        v3.write_new(OUT / f"PAIR_A_STOP_{action}_{int(time.time())}.json", {
            "at": now(), "status": "STOPPED_NO_AUTOMATIC_RETRY", "action": action,
            "error": str(exc), "error_type": type(exc).__name__,
            "truth_opened": (OUT / "PAIR_A_EVALUATION_INPUT_FREEZE.json").exists(),
        })
        raise


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["freeze", "mine", "dedupe", "evaluate"])
    action = parser.parse_args().action
    if action == "freeze":
        freeze_architecture()
    else:
        run(action)


if __name__ == "__main__":
    main()
