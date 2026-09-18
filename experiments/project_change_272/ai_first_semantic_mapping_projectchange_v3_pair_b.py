#!/usr/bin/env python3
"""Frozen-architecture staged execution of V3 for Pair B only."""

from __future__ import annotations

import argparse
import asyncio
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import signal
import sys
import time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from experiments.project_change_272 import ai_first_semantic_mapping_projectchange_v3 as v3  # noqa: E402


OUT = v3.OUT
PAIR = "B"
PAIR_B_ROOT = v3.CORPUS / "pair_b_ai_mapping_change_miner_v1"
PAIR_B_TRUTH = v3.CORPUS / "pair_b_independent_finding_loss_trace/SOURCE_VERIFICATION.json"
PAIR_B_SOURCE_AUDIT = PAIR_B_ROOT / "FALSE_POSITIVE_AUDIT.json"
PAIR_B_F13 = PAIR_B_ROOT / "F13_CHECK.json"
STAGE_BUDGET = 4_000_000
FUTURE_DEDUPE_ALLOWANCE = 320_000


S = {"type": "string"}
STRINGS = {"type": "array", "items": S}
EVAL_ROW = v3.obj(
    projectchange_id=S,
    quality={"type": "string", "enum": ["CORRECT", "PARTIAL", "FALSE", "INSUFFICIENT"]},
    rationale=S,
    grouping_quality={"type": "string", "enum": ["GOOD_HUMAN_LEVEL", "TOO_ATOMIC", "OVER_MERGED", "DUPLICATE"]},
    grouping_rationale=S,
    source_refs=STRINGS,
)
TRUTH_ROW = v3.obj(
    truth_id=S,
    outcome={"type": "string", "enum": ["STRONG", "PARTIAL", "MISSED"]},
    matched_projectchange_ids=STRINGS,
    rationale=S,
)
EVAL_SCHEMA = v3.obj(
    pair={"type": "string", "enum": ["B"]},
    projectchange_rows={"type": "array", "items": EVAL_ROW},
    truth_rows={"type": "array", "items": TRUTH_ROW},
    f13_status={"type": "string", "enum": ["PASS", "FAIL"]},
    f13_rationale=S,
    notes=STRINGS,
)


EVALUATION_PROMPT = """Ты post-freeze source-first evaluator AI-FIRST
PROJECTCHANGE MINER V3 для Pair B / ИОС4.2. Final output уже frozen: не меняй
его, не создавай ProjectChanges и не предлагай tuning.

Оцени каждый final ProjectChange ровно один раз: CORRECT, PARTIAL, FALSE или
INSUFFICIENT. Затем ровно одна grouping label: GOOD_HUMAN_LEVEL, TOO_ATOMIC,
OVER_MERGED или DUPLICATE. Проверяй engineering subject, scope/location, OLD и
NEW states, направление, evidence с обеих сторон и отсутствие metadata-only
claims. FALSE — только при доказанной ошибке/контрдоказательстве; PARTIAL —
доказанное ядро с существенным overclaim или неполнотой; INSufficient — когда
source пакета недостаточно для решения. TOO_ATOMIC — часть одного более широкого
human-level события; OVER_MERGED — независимые события соединены; DUPLICATE —
тот же event остался другой final card (назови ID); иначе GOOD_HUMAN_LEVEL.

Для каждого из frozen PROVEN10 верни STRONG/PARTIAL/MISSED только по final V3,
не по hints. STRONG требует корректного смысла и достаточного охвата; PARTIAL —
смысл найден, но существенная часть пропущена/искажена. F13: FAIL, если V3
утверждает отсутствие OLD систем подпора ТШ/ЛХ паркинга вопреки counterevidence;
иначе PASS. Prior same-version source audit — дополнительная проверенная source
evidence, но сходство формулировок не доказывает V3 claim. Верни exact partition
всех ProjectChange IDs и truth IDs, только JSON по схеме, по-русски."""


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def stage_receipts(kind: str | None = None) -> list[dict[str, Any]]:
    roots = []
    if kind in (None, "MINING"):
        roots.append(OUT / "miner_raw")
    if kind in (None, "DEDUPE"):
        roots.append(OUT / "dedupe_raw")
    if kind in (None, "EVALUATION"):
        roots.append(OUT / "evaluation_raw")
    rows = []
    for root in roots:
        if not root.exists():
            continue
        for path in sorted(root.glob("PAIR_B*/RECEIPT.json")):
            receipt = v3.read_json(path)
            if receipt.get("pair") == "B" or receipt.get("call_id", "").startswith("PAIR_B"):
                rows.append(receipt)
    return rows


def usage_total(receipts: list[dict[str, Any]]) -> int:
    return sum((u.get("input_tokens") or 0) + (u.get("output_tokens") or 0) for receipt in receipts for u in receipt.get("usage", []))


def optimized_preflight() -> dict[str, Any]:
    value = v3.read_json(OUT / "TOKEN_PREFLIGHT_OPTIMIZED.json")
    if value["recommendation"] not in {"PROCEED", "STILL_TOO_LARGE"}:
        raise RuntimeError("Unknown optimized preflight status")
    if value["mapping_freeze_sha256"] != v3.sha256(OUT / "SEMANTIC_MAPPING_V2_FREEZE.json"):
        raise RuntimeError("Mapping freeze drift")
    return value


def freeze_architecture() -> None:
    if (OUT / "PAIR_B_STAGED_ARCHITECTURE_FREEZE.json").exists():
        raise FileExistsError("Pair B staged architecture already frozen")
    preflight = optimized_preflight()
    mapping = v3.read_json(OUT / "PAIR_B_SEMANTIC_MAP.json")
    if len(mapping["regions"]) != 13:
        raise RuntimeError("Pair B region count drift")
    if stage_receipts("MINING") or stage_receipts("DEDUPE"):
        raise RuntimeError("Pair B inference already started before staged freeze")
    rows = {row["region_id"]: row for row in preflight["calls"] if row["pair"] == "B"}
    if set(rows) != {region["region_id"] for region in mapping["regions"]}:
        raise RuntimeError("Optimized Pair B preflight is not exact region partition")
    hashes = {
        "PAIR_B_SEMANTIC_MAP.json": v3.sha256(OUT / "PAIR_B_SEMANTIC_MAP.json"),
        "SEMANTIC_MAPPING_V2_FREEZE.json": v3.sha256(OUT / "SEMANTIC_MAPPING_V2_FREEZE.json"),
        "MINER_PROMPT.txt": v3.sha256(OUT / "MINER_PROMPT.txt"),
        "DEDUPE_PROMPT.txt": v3.sha256(OUT / "DEDUPE_PROMPT.txt"),
        "TOKEN_PREFLIGHT_OPTIMIZED.json": v3.sha256(OUT / "TOKEN_PREFLIGHT_OPTIMIZED.json"),
        "runner": v3.sha256(Path(__file__)),
    }
    for region in mapping["regions"]:
        call_id = f"PAIR_B_{region['region_id']}"
        for name in ("MODEL_INPUT.json", "EXACT_PROMPT.txt", "IMAGE_MANIFEST.json"):
            path = OUT / "miner_inputs_optimized" / call_id / name
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
        if v3.sha256(OUT / "miner_inputs_optimized" / call_id / "EXACT_PROMPT.txt") != rows[region["region_id"]]["input_sha256"]:
            raise RuntimeError(f"Optimized input drift: {call_id}")
    v3.write_new(OUT / "PAIR_B_STAGED_ARCHITECTURE_FREEZE.json", {
        "frozen_at": now(), "status": "FROZEN_BEFORE_PAIR_B_MINING",
        "model": v3.MODEL, "reasoning": v3.REASONING, "provider": "codex_chatgpt",
        "openrouter": 0, "pair_b_regions": 13, "pair_a_regions": 25,
        "pair_a_status": "FROZEN_NOT_RUN", "pair_a_miner_calls": 0,
        "pair_a_truth_opened": False, "pair_b_truth_opened": False,
        "pair_b_projected_mining_input_output": preflight["pair_b_projected_mining_input_output"],
        "pair_b_mining_dedupe_budget": STAGE_BUDGET,
        "prompt_semantics_changed": False, "mapping_changed": False,
        "evidence_changed_after_optimized_preflight": False,
        "schema": v3.MINER_SCHEMA, "hashes": hashes,
    })
    print("PAIR_B_STAGED_ARCHITECTURE_FROZEN", flush=True)


def verify_architecture() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_B_STAGED_ARCHITECTURE_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        path = Path(__file__) if relative == "runner" else OUT / relative
        if v3.sha256(path) != digest:
            raise RuntimeError(f"Pair B architecture drift: {relative}")
    return freeze


async def call(stage: str, call_id: str, prompt: str, data: Any, schema: dict[str, Any], images: list[dict[str, Any]], prepared_input: Path | None = None) -> dict[str, Any]:
    raw_dir = {"MINING": "miner_raw", "DEDUPE": "dedupe_raw", "EVALUATION": "evaluation_raw"}[stage]
    input_dir_name = {"MINING": "miner_inputs_optimized", "DEDUPE": "dedupe_inputs", "EVALUATION": "evaluation_inputs"}[stage]
    target = OUT / raw_dir / call_id
    target.mkdir(parents=True, exist_ok=False)
    input_dir = OUT / input_dir_name / call_id
    if stage != "MINING":
        input_dir.mkdir(parents=True, exist_ok=False)
    names, labels, seen = [], [], set()
    for row in images:
        source = Path(row["path"])
        digest = v3.sha256(source)
        if digest in seen:
            continue
        seen.add(digest)
        name = f"image_{len(names):03d}.png"
        shutil.copyfile(source, target / name)
        names.append(name)
        labels.append({**row["label"], "image": len(names)})
    payload = prompt + "\nIMAGES:\n" + json.dumps(labels, ensure_ascii=False) + "\nSOURCE DATA:\n" + json.dumps(data, ensure_ascii=False)
    if prepared_input is not None:
        if prepared_input.read_text(encoding="utf-8") != payload or v3.read_json(prepared_input.parent / "MODEL_INPUT.json") != data:
            raise RuntimeError(f"Prepared optimized input drift: {call_id}")
    else:
        v3.write_new(input_dir / "MODEL_INPUT.json", data)
        v3.write_new(input_dir / "EXACT_PROMPT.txt", payload)
    v3.write_new(target / "prompt.txt", payload)
    v3.write_new(target / "schema.json", schema)
    command = v3.sandbox_command(target, v3.cli_command(names))
    input_hashes = {path.name: v3.sha256(path) for path in target.iterdir()}
    v3.write_new(target / "INVOCATION.json", {
        "call_id": call_id, "at": now(), "stage": stage, "pair": "B",
        "model": v3.MODEL, "reasoning": v3.REASONING, "provider": "codex_chatgpt",
        "openrouter": 0, "command": command, "input_hashes": input_hashes,
        "retries": 0, "tools_disabled": True, "images": len(names),
    })
    started = time.monotonic()
    with (target / "raw.jsonl").open("wb") as stdout, (target / "stderr.txt").open("wb") as stderr:
        process = await asyncio.create_subprocess_exec(*command, stdin=asyncio.subprocess.PIPE, stdout=stdout, stderr=stderr, env=v3.safe_env(), start_new_session=True)
        try:
            await asyncio.wait_for(process.communicate(payload.encode()), timeout=3600)
        except BaseException:
            os.killpg(process.pid, signal.SIGKILL)
            await process.wait()
            raise
    records = []
    for line in (target / "raw.jsonl").read_text(encoding="utf-8").splitlines():
        try:
            records.append(json.loads(line))
        except ValueError:
            pass
    usage = [row["usage"] for row in records if row.get("type") == "turn.completed" and row.get("usage")]
    tools = [row for row in records if row.get("type", "").startswith("item.") and row.get("item", {}).get("type") not in {None, "agent_message", "reasoning", "error"}]
    receipt = {"call_id": call_id, "stage": stage, "pair": "B", "at": now(), "exit_code": process.returncode, "wall_time_seconds": time.monotonic() - started, "usage": usage, "tool_items": len(tools), "raw_sha256": v3.sha256(target / "raw.jsonl")}
    v3.write_new(target / "RECEIPT.json", receipt)
    if process.returncode or len(usage) != 1 or tools:
        raise RuntimeError(f"Model call failed; no automatic retry: {call_id}")
    for name, digest in input_hashes.items():
        if v3.sha256(target / name) != digest:
            raise RuntimeError(f"Inference input drift: {call_id}/{name}")
    value = v3.read_json(target / "final.txt")
    v3.jsonschema.validate(value, schema)
    v3.write_new(target / "parsed.json", value)
    print(json.dumps({"call": call_id, "status": "SUCCESS", "usage": usage[0]}, ensure_ascii=False), flush=True)
    return value


async def mine_pair_b() -> None:
    freeze = verify_architecture()
    if freeze["pair_a_miner_calls"] != 0:
        raise RuntimeError("Pair A must remain not run")
    preflight = optimized_preflight()
    calls = [row for row in preflight["calls"] if row["pair"] == "B"]
    mapping = v3.read_json(OUT / "PAIR_B_SEMANTIC_MAP.json")
    results = []
    for index, region in enumerate(mapping["regions"]):
        projected_remaining = sum(row["projected_tokens_after"] for row in calls[index:]) + FUTURE_DEDUPE_ALLOWANCE
        if usage_total(stage_receipts("MINING")) + projected_remaining > STAGE_BUDGET:
            raise RuntimeError("PAIR_B_4M_TOKEN_SAFETY_STOP")
        data, images, _ = v3.optimized_region_bundle("B", region)
        call_id = f"PAIR_B_{region['region_id']}"
        value = await call("MINING", call_id, v3.MINER_PROMPT, data, v3.MINER_SCHEMA, images, OUT / "miner_inputs_optimized" / call_id / "EXACT_PROMPT.txt")
        v3.validate_miner("B", region, value)
        results.append(value)
    result_path = OUT / "PAIR_B_PROJECTCHANGE_MINER_RESULTS.json"
    v3.write_new(result_path, {"created_at": now(), "model": v3.MODEL, "reasoning": v3.REASONING, "regions": results, "projectchanges": [change for region in results for change in region["projectchanges"]], "unresolved_hints": [hint for region in results for hint in region["unresolved_hints"]]})
    hashes = {"PAIR_B_PROJECTCHANGE_MINER_RESULTS.json": v3.sha256(result_path), "PAIR_B_STAGED_ARCHITECTURE_FREEZE.json": v3.sha256(OUT / "PAIR_B_STAGED_ARCHITECTURE_FREEZE.json")}
    for path in sorted((OUT / "miner_raw").glob("PAIR_B*/*")):
        if path.is_file():
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
    v3.write_new(OUT / "PAIR_B_PROJECTCHANGE_MINER_V3_FREEZE.json", {"frozen_at": now(), "status": "PAIR_B_MINER_FROZEN", "model": v3.MODEL, "reasoning": v3.REASONING, "regions": 13, "calls": 13, "projectchanges": len([c for r in results for c in r["projectchanges"]]), "unresolved_hints": len([h for r in results for h in r["unresolved_hints"]]), "actual_input_output": usage_total(stage_receipts("MINING")), "budget": STAGE_BUDGET, "truth_opened": False, "pair_a_status": "FROZEN_NOT_RUN", "hashes": hashes})
    print("PAIR_B_PROJECTCHANGE_MINER_V3_FROZEN", flush=True)


def verify_pair_b_miner_freeze() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_B_PROJECTCHANGE_MINER_V3_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        if v3.sha256(OUT / relative) != digest:
            raise RuntimeError(f"Pair B Miner freeze drift: {relative}")
    return freeze


async def dedupe_pair_b() -> None:
    verify_pair_b_miner_freeze()
    mined = v3.read_json(OUT / "PAIR_B_PROJECTCHANGE_MINER_RESULTS.json")
    payload = {"pair": "B", "projectchanges": [v3.compact_change(change) for change in mined["projectchanges"]]}
    if usage_total(stage_receipts("MINING")) + FUTURE_DEDUPE_ALLOWANCE > STAGE_BUDGET:
        raise RuntimeError("PAIR_B_4M_TOKEN_SAFETY_STOP")
    raw = await call("DEDUPE", "PAIR_B_DEDUPE", v3.DEDUPE_PROMPT, payload, v3.DEDUPE_SCHEMA, [])
    if raw["pair"] != "B":
        raise RuntimeError("Dedupe pair mismatch")
    final = v3.apply_dedupe("B", mined["projectchanges"], raw)
    v3.write_new(OUT / "PAIR_B_DEDUPE.json", raw)
    final_path = OUT / "PAIR_B_FINAL_PROJECTCHANGES.json"
    v3.write_new(final_path, {"created_at": now(), "pair": "B", "model": v3.MODEL, "reasoning": v3.REASONING, "projectchanges": final, "unresolved_hints": mined["unresolved_hints"]})
    actual = usage_total(stage_receipts("MINING") + stage_receipts("DEDUPE"))
    if actual > STAGE_BUDGET:
        raise RuntimeError("PAIR_B_4M_TOKEN_SAFETY_STOP_AFTER_DEDUPE")
    hashes = {name: v3.sha256(OUT / name) for name in ("PAIR_B_PROJECTCHANGE_MINER_V3_FREEZE.json", "PAIR_B_DEDUPE.json", "PAIR_B_FINAL_PROJECTCHANGES.json")}
    for path in sorted((OUT / "dedupe_raw" / "PAIR_B_DEDUPE").glob("*")):
        if path.is_file():
            hashes[str(path.relative_to(OUT))] = v3.sha256(path)
    v3.write_new(OUT / "PAIR_B_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json", {"frozen_at": now(), "status": "PAIR_B_FINAL_FROZEN", "model": v3.MODEL, "reasoning": v3.REASONING, "mined": len(mined["projectchanges"]), "final": len(final), "unresolved_hints": len(mined["unresolved_hints"]), "mining_dedupe_input_output": actual, "budget": STAGE_BUDGET, "truth_opened": False, "semantic_mutation_after_freeze": False, "pair_a_status": "FROZEN_NOT_RUN", "pair_a_truth_opened": False, "hashes": hashes})
    print("PAIR_B_AI_FIRST_PROJECTCHANGE_V3_RESULT_FROZEN", flush=True)


def verify_final_freeze() -> dict[str, Any]:
    freeze = v3.read_json(OUT / "PAIR_B_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json")
    for relative, digest in freeze["hashes"].items():
        if v3.sha256(OUT / relative) != digest:
            raise RuntimeError(f"Pair B final freeze drift: {relative}")
    return freeze


def compact_change(change: dict[str, Any]) -> dict[str, Any]:
    keys = ("projectchange_id", "engineering_subject", "scope", "locations", "change_summary", "old_state", "new_state", "changed_parameters", "old_pages", "new_pages", "modalities", "confidence", "why_one_event", "dedupe_lineage", "dedupe_reason")
    return {key: change[key] for key in keys if key in change} | {"evidence_items": [{key: item[key] for key in ("side", "physical_page", "block_id", "block_type", "relevant_fragment", "evidence_role")} for item in change["evidence_items"]]}


def compact_source_audit() -> list[dict[str, Any]]:
    rows = v3.read_json(PAIR_B_SOURCE_AUDIT)["rows"]
    keys = ("change_id", "audit_status", "engineering_subject", "change_summary", "reason", "rationale", "old_pages", "new_pages")
    return [{key: row.get(key) for key in keys} for row in rows]


def evaluation_payload() -> tuple[dict[str, Any], list[str], list[str]]:
    final = v3.read_json(OUT / "PAIR_B_FINAL_PROJECTCHANGES.json")["projectchanges"]
    truth = v3.read_json(PAIR_B_TRUTH)
    proven = [item for item in truth["findings"] if item.get("verdict") == "PROVEN"]
    if len(proven) != 10:
        raise RuntimeError(f"Expected PROVEN10, got {len(proven)}")
    keys = ("finding_id", "engineering_subject", "old_state", "new_state", "old_physical_pdf_pages", "new_physical_pdf_pages", "evidence_types", "minimal_sufficient_evidence", "not_proven")
    payload = {"pair": "B", "frozen_v3_projectchanges": [compact_change(item) for item in final], "source_first_truth": [{key: item.get(key) for key in keys} for item in proven], "f13": v3.read_json(PAIR_B_F13), "prior_same_version_source_audit": compact_source_audit()}
    return payload, [item["projectchange_id"] for item in final], [item["finding_id"] for item in proven]


def validate_evaluation(value: dict[str, Any], projectchange_ids: list[str], truth_ids: list[str]) -> None:
    if value["pair"] != "B":
        raise RuntimeError("Evaluation pair mismatch")
    got_changes = [row["projectchange_id"] for row in value["projectchange_rows"]]
    got_truth = [row["truth_id"] for row in value["truth_rows"]]
    if len(got_changes) != len(set(got_changes)) or set(got_changes) != set(projectchange_ids):
        raise RuntimeError("Evaluation is not exact ProjectChange partition")
    if len(got_truth) != len(set(got_truth)) or set(got_truth) != set(truth_ids):
        raise RuntimeError("Evaluation is not exact PROVEN10 partition")
    matched = {item for row in value["truth_rows"] for item in row["matched_projectchange_ids"]}
    if not matched <= set(projectchange_ids):
        raise RuntimeError("Evaluation refers to unknown ProjectChange")


def summarize(rows: list[dict[str, Any]], field: str, labels: list[str]) -> dict[str, int]:
    return {label: sum(row[field] == label for row in rows) for label in labels}


def usage_report() -> dict[str, Any]:
    def section(receipts: list[dict[str, Any]]) -> dict[str, Any]:
        usages = [usage for receipt in receipts for usage in receipt.get("usage", [])]
        return {
            "calls": len(usages),
            "input_tokens": sum(item.get("input_tokens") or 0 for item in usages),
            "cached_input_tokens": sum(item.get("cached_input_tokens") or 0 for item in usages),
            "output_tokens": sum(item.get("output_tokens") or 0 for item in usages),
            "reasoning_output_tokens": sum(item.get("reasoning_output_tokens") or 0 for item in usages),
            "input_output": sum((item.get("input_tokens") or 0) + (item.get("output_tokens") or 0) for item in usages),
        }
    mining = section(stage_receipts("MINING"))
    dedupe = section(stage_receipts("DEDUPE"))
    evaluation = section(stage_receipts("EVALUATION"))
    total = {key: mining[key] + dedupe[key] for key in mining}
    return {"accounting_note": "cached_input_tokens are a subset of input_tokens and are not added again", "pair_b_mining": mining, "pair_b_dedupe": dedupe, "pair_b_mining_dedupe_total": total, "pair_b_evaluation_separate": evaluation, "pair_a_projected_remaining": 7_884_077}


def write_results(evaluation: dict[str, Any]) -> None:
    quality = summarize(evaluation["projectchange_rows"], "quality", ["CORRECT", "PARTIAL", "FALSE", "INSUFFICIENT"])
    grouping = summarize(evaluation["projectchange_rows"], "grouping_quality", ["GOOD_HUMAN_LEVEL", "TOO_ATOMIC", "OVER_MERGED", "DUPLICATE"])
    proven = summarize(evaluation["truth_rows"], "outcome", ["STRONG", "PARTIAL", "MISSED"])
    final = v3.read_json(OUT / "PAIR_B_FINAL_PROJECTCHANGES.json")
    usage = usage_report()
    gate = (
        proven["MISSED"] == 0
        and evaluation["f13_status"] == "PASS"
        and quality["FALSE"] <= 2
        and grouping["DUPLICATE"] < 22
        and grouping["GOOD_HUMAN_LEVEL"] > 27
        and len(final["projectchanges"]) <= 77
        and usage["pair_b_mining_dedupe_total"]["input_output"] <= STAGE_BUDGET
    )
    recommendation = "PAIR_B_GATE_PASS_RUN_PAIR_A" if gate else "PAIR_B_GATE_FAIL_STOP_V3"
    comparison = {
        "pair_b_v2": {"final_projectchanges": 77, "GOOD_HUMAN_LEVEL": 27, "TOO_ATOMIC": 5, "OVER_MERGED": 2, "DUPLICATE": 43, "PROVEN10": {"STRONG": 8, "PARTIAL": 2, "MISSED": 0}, "F13": "PASS"},
        "pair_b_v3": {"final_projectchanges": len(final["projectchanges"]), **grouping, "PROVEN10": proven, "F13": evaluation["f13_status"]},
        "recommendation": recommendation,
        "gate_rule": {"missed_max": 0, "f13": "PASS", "false_max": 2, "duplicates_less_than": 22, "good_human_level_greater_than": 27, "final_projectchanges_max": 77, "mining_dedupe_budget": STAGE_BUDGET},
    }
    v3.write_new(OUT / "PAIR_B_EVALUATION.json", evaluation)
    v3.write_new(OUT / "PAIR_B_GROUPING_QUALITY.json", {"counts": grouping, "rows": [{"projectchange_id": row["projectchange_id"], "grouping_quality": row["grouping_quality"], "rationale": row["grouping_rationale"]} for row in evaluation["projectchange_rows"]]})
    v3.write_new(OUT / "PAIR_B_TOKEN_USAGE.json", usage)
    v3.write_new(OUT / "PAIR_B_V2_VS_V3.json", comparison)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "V2", "V3"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    summary_rows = [
        ("Final ProjectChanges", 77, len(final["projectchanges"])),
        ("Unresolved hints", "", len(final["unresolved_hints"])),
        *[(label, comparison["pair_b_v2"].get(label, ""), grouping[label]) for label in grouping],
        *[(f"PROVEN10 {label}", comparison["pair_b_v2"]["PROVEN10"][label], proven[label]) for label in proven],
        ("F13", "PASS", evaluation["f13_status"]),
        ("Recommendation", "", recommendation),
    ]
    for row in summary_rows:
        summary.append(row)
    changes = workbook.create_sheet("ProjectChanges")
    changes.append(["projectchange_id", "quality", "grouping_quality", "rationale", "grouping_rationale", "source_refs"])
    for row in evaluation["projectchange_rows"]:
        changes.append([row["projectchange_id"], row["quality"], row["grouping_quality"], row["rationale"], row["grouping_rationale"], "\n".join(row["source_refs"])])
    truth = workbook.create_sheet("PROVEN10")
    truth.append(["truth_id", "outcome", "matched_projectchange_ids", "rationale"])
    for row in evaluation["truth_rows"]:
        truth.append([row["truth_id"], row["outcome"], ", ".join(row["matched_projectchange_ids"]), row["rationale"]])
    usage_sheet = workbook.create_sheet("Token usage")
    usage_sheet.append(["Stage", "calls", "input", "cached", "output", "reasoning", "input+output"])
    for key in ("pair_b_mining", "pair_b_dedupe", "pair_b_mining_dedupe_total", "pair_b_evaluation_separate"):
        row = usage[key]
        usage_sheet.append([key, row["calls"], row["input_tokens"], row["cached_input_tokens"], row["output_tokens"], row["reasoning_output_tokens"], row["input_output"]])
    workbook.save(OUT / "PAIR_B_RESULTS.xlsx")
    report = f"""# ProjectChange V3 — staged Pair B report

STATUS: COMPLETE_PAIR_B_ONLY

MODEL: {v3.MODEL} {v3.REASONING}

Pair A: FROZEN_NOT_RUN; Miner calls 0; REAL15 NOT OPENED.

Pair B semantic regions: 13
Final ProjectChanges: {len(final['projectchanges'])}
Unresolved hints: {len(final['unresolved_hints'])}

PROVEN10: STRONG {proven['STRONG']}; PARTIAL {proven['PARTIAL']}; MISSED {proven['MISSED']}.
F13: {evaluation['f13_status']}.

Quality: CORRECT {quality['CORRECT']}; PARTIAL {quality['PARTIAL']}; FALSE {quality['FALSE']}; INSUFFICIENT {quality['INSUFFICIENT']}.
Grouping: GOOD_HUMAN_LEVEL {grouping['GOOD_HUMAN_LEVEL']}; TOO_ATOMIC {grouping['TOO_ATOMIC']}; OVER_MERGED {grouping['OVER_MERGED']}; DUPLICATE {grouping['DUPLICATE']}.

Pair B mining+dedupe tokens: {usage['pair_b_mining_dedupe_total']['input_output']} / {STAGE_BUDGET}.
Pair A projected remaining: 7,884,077.

FINAL RECOMMENDATION: {recommendation}

PRODUCTION: UNCHANGED
VALIDATION: NOT OPENED
FINAL HOLDOUT: NOT OPENED
"""
    v3.write_new(OUT / "PAIR_B_FINAL_REPORT.md", report)


async def evaluate_pair_b() -> None:
    freeze = verify_final_freeze()
    if not freeze["truth_opened"] is False or freeze["pair_a_truth_opened"] is not False:
        raise RuntimeError("Invalid pre-evaluation truth state")
    payload, projectchange_ids, truth_ids = evaluation_payload()
    hashes = {"pair_b_truth": v3.sha256(PAIR_B_TRUTH), "pair_b_f13": v3.sha256(PAIR_B_F13), "pair_b_source_audit": v3.sha256(PAIR_B_SOURCE_AUDIT), "pair_b_result_freeze": v3.sha256(OUT / "PAIR_B_AI_FIRST_PROJECTCHANGE_V3_RESULT_FREEZE.json")}
    v3.write_new(OUT / "PAIR_B_EVALUATION_INPUT_FREEZE.json", {"opened_at": now(), "pair_a_truth_opened": False, "pair_b_truth_opened_after_final_freeze": True, "hashes": hashes})
    value = await call("EVALUATION", "PAIR_B_SOURCE_FIRST_EVALUATION", EVALUATION_PROMPT, payload, EVAL_SCHEMA, [])
    validate_evaluation(value, projectchange_ids, truth_ids)
    write_results(value)
    v3.write_new(OUT / "PAIR_B_EVALUATION_FREEZE.json", {"frozen_at": now(), "status": "PAIR_B_EVALUATION_COMPLETE_STOP", "pair_a_status": "FROZEN_NOT_RUN", "pair_a_truth_opened": False, "no_post_truth_tuning": True, "hashes": {name: v3.sha256(OUT / name) for name in ("PAIR_B_EVALUATION.json", "PAIR_B_GROUPING_QUALITY.json", "PAIR_B_TOKEN_USAGE.json", "PAIR_B_V2_VS_V3.json", "PAIR_B_RESULTS.xlsx", "PAIR_B_FINAL_REPORT.md")}})
    print((OUT / "PAIR_B_FINAL_REPORT.md").read_text(encoding="utf-8"), flush=True)


def run(action: str) -> None:
    try:
        asyncio.run({"mine": mine_pair_b, "dedupe": dedupe_pair_b, "evaluate": evaluate_pair_b}[action]())
    except BaseException as exc:
        v3.write_new(OUT / f"PAIR_B_STOP_{action}_{int(time.time())}.json", {"at": now(), "status": "STOPPED_NO_AUTOMATIC_RETRY", "action": action, "error": str(exc), "error_type": type(exc).__name__, "pair_a_status": "FROZEN_NOT_RUN", "pair_a_truth_opened": False})
        raise


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["freeze", "mine", "dedupe", "evaluate"])
    action = parser.parse_args().action
    if action == "freeze":
        freeze_architecture()
    else:
        run(action)


if __name__ == "__main__":
    main()
